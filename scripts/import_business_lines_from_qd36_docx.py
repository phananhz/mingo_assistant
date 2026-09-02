from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path

from docx import Document


EXPECTED_COUNTS = {1: 22, 2: 87, 3: 259, 4: 495, 5: 743}
OUTPUT_PATH = Path("assets/business_lines/vsic_2025.json")
DEFAULT_SOURCE_PATH = Path("Templates04/36_2025_QD-TTg_675010.docx")
SOURCE_TITLE = "Hệ thống ngành kinh tế Việt Nam 2025 - Quyết định 36/2025/QĐ-TTg"


@dataclass
class Node:
    code: str
    name: str
    detail: str = ""
    children: list["Node"] = field(default_factory=list)


def main() -> int:
    args = _parse_args()
    docx_path = args.source_docx
    if not docx_path.exists():
        raise FileNotFoundError(f"Không tìm thấy file nguồn: {docx_path}")

    document = Document(str(docx_path))
    catalog_table = _find_catalog_table(document)
    catalog_rows = _catalog_rows(catalog_table)
    nodes = _catalog_tree(catalog_rows)
    details_by_code = _details_by_code(document)
    _apply_details(nodes, details_by_code)
    display_rows = _level4_display_rows(catalog_rows, details_by_code)

    flat_rows = _flatten(nodes)
    counts = _counts_by_level(flat_rows)
    _assert_counts(counts)
    _assert_code_shapes(flat_rows)
    _assert_samples(flat_rows)

    payload = {
        "source_title": SOURCE_TITLE,
        "source_file": str(docx_path),
        "expected_counts": {"4": EXPECTED_COUNTS[4]},
        "actual_counts": {"4": len(display_rows)},
        "display_rows": display_rows,
        "items": [
            asdict(Node(code=str(row["levels"][3]), name=str(row["name"]), detail=str(row["detail"])))
            for row in display_rows
        ],
    }

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"saved {OUTPUT_PATH}")
    print("counts " + ", ".join(f"cap {level}: {counts[level]}" for level in sorted(counts)))
    print(f"total {sum(counts.values())}")
    return 0


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import VSIC 2025 business lines from Quyết định 36/2025 DOCX.")
    parser.add_argument(
        "source_docx",
        nargs="?",
        type=Path,
        default=DEFAULT_SOURCE_PATH,
        help=f"Path to 36/2025/QĐ-TTg DOCX. Default: {DEFAULT_SOURCE_PATH}",
    )
    return parser.parse_args()


def _find_catalog_table(document: Document):
    for table in document.tables:
        if not table.rows:
            continue
        header = [_clean_cell(cell.text) for cell in table.rows[0].cells]
        if header[:6] == ["Cấp 1", "Cấp 2", "Cấp 3", "Cấp 4", "Cấp 5", "Tên ngành"]:
            return table
    raise RuntimeError("Không tìm thấy bảng danh mục ngành kinh tế 6 cột trong file DOCX.")


def _catalog_rows(table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row_number, row in enumerate(table.rows[1:], start=2):
        cells = [_clean_cell(cell.text) for cell in row.cells]
        if _is_header_row(cells) or _is_total_row(cells):
            continue
        if not any(cells):
            continue
        if not cells[5]:
            raise RuntimeError(f"Dòng {row_number} không phải dòng header/tổng nhưng thiếu tên ngành: {cells[:6]}")
        rows.append(cells[:6])
    return rows


def _catalog_tree(catalog_rows: list[list[str]]) -> list[Node]:
    roots: list[Node] = []
    last_by_level: dict[int, Node] = {}
    seen: set[tuple[int, str]] = set()

    for row_number, cells in enumerate(catalog_rows, start=2):
        name = cells[5] if len(cells) > 5 else ""
        codes = _valid_codes(cells[:5])
        invalid_codes = _invalid_codes(cells[:5])
        if invalid_codes:
            invalid = ", ".join(invalid_codes)
            raise RuntimeError(f"Dòng {row_number} có mã ngành sai định dạng: {invalid}")
        if not codes:
            continue
        if not name:
            raise RuntimeError(f"Dòng {row_number} có mã ngành nhưng thiếu tên ngành: {codes}")

        for level, code in codes:
            key = (level, code)
            if key in seen:
                raise RuntimeError(f"Mã ngành bị trùng cấp {level}: {code}")
            node = Node(code=code, name=name)
            if level == 1:
                roots.append(node)
            else:
                parent = last_by_level.get(level - 1)
                if parent is None:
                    raise RuntimeError(f"Dòng {row_number} thiếu mã cha cho cấp {level}: {code} - {name}")
                parent.children.append(node)
            seen.add(key)
            last_by_level[level] = node
            for deeper_level in range(level + 1, 6):
                last_by_level.pop(deeper_level, None)

    return roots


def _level4_display_rows(catalog_rows: list[list[str]], details_by_code: dict[str, str]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    seen: set[str] = set()
    for cells in catalog_rows:
        code = cells[3] if len(cells) > 3 else ""
        if not _code_pattern(4).fullmatch(code):
            continue
        if code in seen:
            raise RuntimeError(f"Duplicate level 4 industry code: {code}")
        seen.add(code)
        rows.append(
            {
                "levels": ["", "", "", code, ""],
                "name": cells[5],
                "detail": details_by_code.get(code, ""),
            }
        )
    if len(rows) != EXPECTED_COUNTS[4]:
        raise RuntimeError(f"Expected {EXPECTED_COUNTS[4]} level 4 rows, found {len(rows)}.")
    return rows


def _display_rows(catalog_rows: list[list[str]], details_by_code: dict[str, str]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for cells in catalog_rows:
        levels = cells[:5]
        codes = [code for code in levels if code]
        detail = ""
        for code in reversed(codes):
            detail = details_by_code.get(code, "")
            if detail:
                break
        rows.append({"levels": levels, "name": cells[5], "detail": detail})
    return rows


def _is_header_row(cells: list[str]) -> bool:
    return cells[:6] == ["Cấp 1", "Cấp 2", "Cấp 3", "Cấp 4", "Cấp 5", "Tên ngành"]


def _is_total_row(cells: list[str]) -> bool:
    return cells[:6] == ["22", "87", "259", "495", "743", ""]


def _valid_codes(cells: list[str]) -> list[tuple[int, str]]:
    return [
        (level, code)
        for level, code in enumerate(cells, start=1)
        if code and _code_pattern(level).fullmatch(code)
    ]


def _invalid_codes(cells: list[str]) -> list[str]:
    return [
        code
        for level, code in enumerate(cells, start=1)
        if code and not _code_pattern(level).fullmatch(code)
    ]


def _code_pattern(level: int) -> re.Pattern[str]:
    patterns = {
        1: re.compile(r"^[A-V]$"),
        2: re.compile(r"^[0-9]{2}$"),
        3: re.compile(r"^[0-9]{3}$"),
        4: re.compile(r"^[0-9]{4}$"),
        5: re.compile(r"^[0-9]{5}$"),
    }
    return patterns[level]


def _details_by_code(document: Document) -> dict[str, str]:
    details: dict[str, list[str]] = {}
    current_codes: list[str] = []

    for paragraph in document.paragraphs:
        text = _clean_cell(paragraph.text)
        if not text:
            continue
        heading_codes = _heading_codes(text)
        if heading_codes:
            current_codes = heading_codes
            for code in current_codes:
                details.setdefault(code, [])
            continue
        for code in current_codes:
            details[code].append(text)

    return {code: " ".join(parts).strip() for code, parts in details.items() if parts}


def _heading_codes(text: str) -> list[str]:
    match = re.match(r"^([A-V]|[0-9]{2,5}(?:\s*[-–]\s*[0-9]{2,5})*):\s+.+$", text)
    if not match:
        return []
    codes = re.findall(r"[A-V]|[0-9]{2,5}", match.group(1))
    valid_codes = [code for code in codes if _level_from_code(code) in EXPECTED_COUNTS]
    return valid_codes if len(valid_codes) == len(codes) else []


def _apply_details(nodes: list[Node], details_by_code: dict[str, str]) -> None:
    for node in nodes:
        node.detail = details_by_code.get(node.code, "")
        _apply_details(node.children, details_by_code)


def _flatten(nodes: list[Node]) -> list[tuple[int, str, str]]:
    rows: list[tuple[int, str, str]] = []
    for node in nodes:
        rows.extend(_flatten_node(node))
    return rows


def _flatten_node(node: Node) -> list[tuple[int, str, str]]:
    rows = [(_level_from_code(node.code), node.code, node.name)]
    for child in node.children:
        rows.extend(_flatten_node(child))
    return rows


def _level_from_code(code: str) -> int:
    if re.fullmatch(r"[A-V]", code):
        return 1
    if re.fullmatch(r"[0-9]{2,5}", code):
        return len(code)
    return 0


def _counts_by_level(rows: list[tuple[int, str, str]]) -> dict[int, int]:
    counts = {level: 0 for level in EXPECTED_COUNTS}
    for level, _, _ in rows:
        counts[level] += 1
    return counts


def _assert_counts(counts: dict[int, int]) -> None:
    if counts != EXPECTED_COUNTS:
        expected = ", ".join(f"cấp {key}: {value}" for key, value in EXPECTED_COUNTS.items())
        actual = ", ".join(f"cấp {key}: {value}" for key, value in counts.items())
        raise RuntimeError(f"Số lượng mã ngành không khớp. Expected {expected}. Actual {actual}.")


def _assert_code_shapes(rows: list[tuple[int, str, str]]) -> None:
    seen: set[tuple[int, str]] = set()
    for level, code, name in rows:
        if level not in EXPECTED_COUNTS:
            raise RuntimeError(f"Không xác định được cấp mã ngành: {code} - {name}")
        if not _code_pattern(level).fullmatch(code):
            raise RuntimeError(f"Mã ngành sai định dạng cấp {level}: {code} - {name}")
        key = (level, code)
        if key in seen:
            raise RuntimeError(f"Mã ngành bị trùng cấp {level}: {code}")
        seen.add(key)


def _assert_samples(rows: list[tuple[int, str, str]]) -> None:
    expected = {
        "01110": "Trồng lúa",
        "03221": "Nuôi cá",
        "03222": "Nuôi tôm",
        "61201": "Hoạt động bán lại dịch vụ viễn thông",
        "61202": "Dịch vụ trung gian cho hoạt động viễn thông",
        "61900": "Hoạt động viễn thông khác",
        "99000": "Hoạt động của các tổ chức và cơ quan quốc tế",
    }
    actual = {code: name for _, code, name in rows}
    for code, name in expected.items():
        if actual.get(code) != name:
            raise RuntimeError(f"Mã mẫu không khớp: {code}. Expected {name}. Actual {actual.get(code)}")


def _clean_cell(value: str) -> str:
    return " ".join(value.replace("\xa0", " ").split())


if __name__ == "__main__":
    sys.exit(main())
