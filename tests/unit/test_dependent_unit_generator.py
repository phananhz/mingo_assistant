from __future__ import annotations

import hashlib
import json
import tempfile
import zipfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

from mingo_assistant.services.documents.dependent_unit_generator import DependentUnitTemplateGenerator


class DependentUnitTemplateGeneratorTest(unittest.TestCase):
    def test_fills_only_the_manifest_approved_template(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template_dir = root / "templates"
            export_root = root / "exports"
            template_dir.mkdir()
            export_root.mkdir()

            editable = template_dir / "1. request.docx"
            document = Document()
            highlighted = document.add_paragraph("Location: ").add_run("SAMPLE")
            highlighted.font.highlight_color = WD_COLOR_INDEX.YELLOW
            document.save(editable)

            passthroughs = []
            for filename in ("2. authorization.docx", "signing-guide.docx"):
                passthrough = template_dir / filename
                document = Document()
                document.add_paragraph(f"Template {filename}")
                document.save(passthrough)
                passthroughs.append(passthrough)

            manifest = {
                "templates": {
                    editable.name: {
                        "mode": "fill_highlighted",
                        "highlighted_fields": {"0": "location_business_name"},
                    }
                }
            }
            (template_dir / "template_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")

            output_dir = root / "out"
            result = DependentUnitTemplateGenerator(template_dir, export_root).generate(
                {"location_business_name": "LOCATION A"}, "LOCATION A", output_dir
            )

            self.assertEqual(3, len(result.generated))
            generated_document = Document(output_dir / editable.name)
            self.assertIn("LOCATION A", generated_document.paragraphs[0].text)
            self.assertEqual(
                WD_COLOR_INDEX.YELLOW,
                generated_document.paragraphs[0].runs[-1].font.highlight_color,
            )
            for source in passthroughs:
                copied = output_dir / source.name
                self.assertEqual(
                    hashlib.sha256(source.read_bytes()).digest(),
                    hashlib.sha256(copied.read_bytes()).digest(),
                )

    def test_runtime_templates_preserve_yellow_fill_and_passthrough_files(self) -> None:
        project_root = Path(__file__).resolve().parents[2]
        template_dir = project_root / "assets" / "templates" / "word" / "dependent_units" / "ddkd_create"
        editable = next(template_dir.glob("1.*.docx"))
        authorization_template = next(template_dir.glob("2.*.docx"))
        guide_template = next(template_dir.glob("H*.docx"))
        values = {
            "business_registration_office": "Phòng Đăng ký kinh doanh Hà Nội",
            "parent_entity_name": "CÔNG TY TNHH MINGO",
            "parent_entity_registration_number": "0101234567",
            "parent_entity_name_prefix": "CÔNG TY TNHH",
            "parent_entity_name_suffix": "MINGO",
            "location_business_name": "ĐỊA ĐIỂM KINH DOANH MINGO",
            "location_address_detail": "Số 1 Phố Mẫu",
            "location_ward": "Phường Mẫu",
            "location_province": "Thành phố Hà Nội",
            "location_manager_full_name": "NGUYỄN VĂN A",
            "location_manager_date_of_birth": "01/01/1990",
            "location_manager_gender": "Nam",
            "location_manager_id_number": "001090000001",
            "location_manager_address_detail": "Số 2 Phố Mẫu",
            "location_manager_ward": "Phường Mẫu",
            "location_manager_province": "Thành phố Hà Nội",
            "location_manager_nationality": "Việt Nam",
            "legal_representative_full_name": "LEGAL REP",
            "legal_representative_id_number": "REP-ID",
            "legal_representative_position": "Giám đốc",
            "company_name": "COMPANY A",
            "company_address": "1 COMPANY STREET, Ward A, Hà Nội",
            "authorized_person_full_name": "AUTHORIZED PERSON",
            "authorized_person_id_number": "AUTH-ID",
            "authorized_person_contact_address": "2 AUTH STREET",
            "authorized_person_phone": "0900000000",
            "authorized_person_email": "authorized@example.test",
            "legal_representative_issue_date": "01/01/2020",
            "legal_representative_issue_place": "Hà Nội",
            "document_place": "Hà Nội",
            "document_day": "10",
            "document_month": "09",
            "document_year": "2026",
            "business_line_1_number": "1",
            "business_line_1_name": "Lập trình máy vi tính",
            "business_line_1_detail": "Lập trình máy vi tính",
            "business_line_1_code": "6201",
            "business_line_1_is_main": "X",
            "tax_accounting_dependent_x": "X",
            "tax_labor_count": "1",
            "tax_bot_project_x": "",
            "tax_vat_deduction_x": "X",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            output_dir = Path(temp_dir) / "output"
            result = DependentUnitTemplateGenerator(template_dir, Path(temp_dir)).generate(
                values,
                "Địa điểm kinh doanh Mingo",
                output_dir,
            )
            self.assertEqual(3, len(result.generated))
            generated_document = Document(output_dir / editable.name)
            paragraphs = list(generated_document.paragraphs)
            for table in generated_document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
            matching_runs = [
                run
                for paragraph in paragraphs
                for run in paragraph.runs
                if "ĐỊA ĐIỂM KINH DOANH MINGO" in run.text
            ]
            self.assertTrue(matching_runs)
            self.assertTrue(all(run.font.highlight_color == WD_COLOR_INDEX.YELLOW for run in matching_runs))
            with zipfile.ZipFile(output_dir / editable.name) as archive:
                self.assertIsNone(archive.testzip())
                self.assertTrue(archive.read("word/document.xml").startswith(b"<?xml"))
            authorization = Document(output_dir / authorization_template.name)
            authorization_text = "\n".join(paragraph.text for paragraph in authorization.paragraphs)
            for table in authorization.tables:
                for row in table.rows:
                    for cell in row.cells:
                        authorization_text += "\n" + "\n".join(paragraph.text for paragraph in cell.paragraphs)
            self.assertIn("LEGAL REP", authorization_text)
            self.assertIn("COMPANY A", authorization_text)
            self.assertIn("Thông báo lập địa điểm kinh doanh", authorization_text)
            guide = Document(output_dir / guide_template.name)
            guide_text = "\n".join(paragraph.text for paragraph in guide.paragraphs)
            for table in guide.tables:
                for row in table.rows:
                    for cell in row.cells:
                        guide_text += "\n" + "\n".join(paragraph.text for paragraph in cell.paragraphs)
            self.assertEqual(2, guide_text.count("LEGAL REP"))

    def test_empty_location_fields_keep_dotted_placeholders(self) -> None:
        project_root = Path(__file__).resolve().parents[2]
        template_dir = project_root / "assets" / "templates" / "word" / "dependent_units" / "ddkd_create"
        with tempfile.TemporaryDirectory() as temp_dir:
            output_dir = Path(temp_dir) / "output"
            DependentUnitTemplateGenerator(template_dir, Path(temp_dir)).generate(
                {"location_business_name": "LOCATION"},
                "location",
                output_dir,
            )
            document = Document(output_dir / next(template_dir.glob("1.*.docx")).name)
            paragraphs = list(document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
            highlighted = [
                run
                for paragraph in paragraphs
                for run in paragraph.runs
                if run.font.highlight_color == WD_COLOR_INDEX.YELLOW
            ]
            for index in (4, 5, 10, 11, 13, 15):
                self.assertTrue(highlighted[index].text.strip(), index)
                self.assertTrue(
                    set(highlighted[index].text.strip()) <= set(".…/"),
                    (index, highlighted[index].text),
                )
    def test_runtime_template_writes_only_actual_industry_rows_and_foreign_highlights(self) -> None:
        project_root = Path(__file__).resolve().parents[2]
        template_dir = project_root / "assets" / "templates" / "word" / "dependent_units" / "ddkd_create"
        editable = next(template_dir.glob("1.*.docx"))
        values = {
            "location_manager_person_type": "Ng\u01b0\u1eddi n\u01b0\u1edbc ngo\u00e0i",
            "location_manager_ethnicity": "Ethnicity A",
            "location_manager_passport_number": "P1234567",
            "location_manager_passport_issue_date": "01/02/2020",
            "location_manager_passport_issue_place": "Place A",
            "location_manager_foreign_address_detail": "Foreign address A",
            "location_manager_foreign_country": "Country A",
            "business_line_1_name": "LINE ONE",
            "business_line_1_code": "0111",
            "business_line_1_is_main": "X",
            "business_line_2_name": "LINE TWO",
            "business_line_2_code": "0150",
            "tax_accounting_method": "H\u1ea1ch to\u00e1n ph\u1ee5 thu\u1ed9c",
            "tax_bot_project": "",
            "tax_vat_method": "Kh\u1ea5u tr\u1eeb",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            output_dir = Path(temp_dir) / "output"
            generator = DependentUnitTemplateGenerator(template_dir, Path(temp_dir))
            generator.generate(values, "location", output_dir)
            document = Document(output_dir / editable.name)
            run_texts = [run.text for paragraph in document.paragraphs for run in paragraph.runs]
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        run_texts.extend(run.text for paragraph in cell.paragraphs for run in paragraph.runs)
            for expected in ("Ethnicity A", "P1234567", "01/02/2020", "Place A", "Foreign address A", "Country A"):
                self.assertIn(expected, run_texts)

            foreign_permanent_address = next(
                paragraph for paragraph in generator._iter_paragraphs(document)
                if generator._fold(paragraph.text).strip().startswith("noi thuong tru")
            )
            self.assertEqual(foreign_permanent_address.text, "\u004e\u01a1i th\u01b0\u1eddng tr\u00fa: " + "." * 60)
            self.assertEqual(foreign_permanent_address.runs[-1].font.size.pt, 13)
            industry_table = next(
                table for table in document.tables
                if generator._business_line_header_index(table) is not None
            )
            header_index = generator._business_line_header_index(industry_table)
            self.assertIsNotNone(header_index)
            self.assertEqual(header_index + 1 + 2, len(industry_table.rows))
            self.assertTrue(all(row.cells[0].text.strip() for row in industry_table.rows[header_index + 1 :]))

            symbol_chars = []
            for paragraph in generator._iter_paragraphs(document):
                for run in paragraph.runs:
                    symbol_chars.extend(symbol.get(qn("w:char"), "").upper() for symbol in run._r.findall(qn("w:sym")))
            self.assertGreaterEqual(symbol_chars.count("F053"), 3)
            self.assertIn("F063", symbol_chars)

            tax_table = next(
                table for table in document.tables
                if "thong tin dang ky thue" in generator._fold(" ".join(cell.text for row in table.rows for cell in row.cells))
            )
            tax_runs = [run for paragraph in generator._table_paragraphs([tax_table]) for run in paragraph.runs]
            placeholder_runs = [
                run for run in tax_runs
                if run.text == "." * 32 or run.text == "\u2026" * 6 + "/" + "\u2026" * 6 + "/" + "\u2026" * 6
            ]
            self.assertNotIn(".", [run.text for run in tax_runs])
            self.assertGreaterEqual(len(placeholder_runs), 2)
            self.assertTrue(all(run.font.size and run.font.size.pt == 13 for run in placeholder_runs))
            tax_symbols = [symbol for run in tax_runs for symbol in run._r.findall(qn("w:sym"))]
            self.assertTrue(any(
                symbol.get(qn("w:font")) == "Wingdings 2" and symbol.get(qn("w:char")) == "F053"
                for symbol in tax_symbols
            ))
            self.assertTrue(any(
                symbol.get(qn("w:font")) == "Webdings" and symbol.get(qn("w:char"), "").upper() == "F063"
                for symbol in tax_symbols
            ))