from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from mingo_assistant.domains.household_business import HouseholdBusinessCase, HouseholdMember
from mingo_assistant.services.documents.household_business_generator import (
    AUTHORIZATION_TEMPLATE,
    AUTHORIZATION_SOURCE_TEMPLATE,
    REQUEST_TEMPLATE,
    REQUEST_SOURCE_TEMPLATE,
    HouseholdBusinessTemplateGenerator,
)


ROOT = Path(__file__).resolve().parents[2]
TEMPLATES = ROOT / "assets" / "templates" / "word" / "household_business"


def sample_values() -> dict[str, str]:
    return {
        "household_registration_office": "Cơ quan đăng ký kinh doanh cấp xã",
        "document_place": "Thành phố Hồ Chí Minh",
        "document_day": "17",
        "document_month": "10",
        "document_year": "2026",
        "household_owner_full_name": "Nguyễn Văn Kiểm Thử",
        "household_owner_date_of_birth": "01/02/1990",
        "household_owner_gender": "Nam",
        "household_owner_id_number": "012345678901",
        "household_owner_issue_date": "03/04/2021",
        "household_owner_issue_place": "Bộ Công an",
        "household_owner_contact_address": "01 Đường Mẫu, Phường Mẫu, Thành phố Hồ Chí Minh",
        "household_business_name": "Hộ kinh doanh kiểm thử",
        "household_address_detail": "02 Đường Kinh Doanh",
        "household_ward": "Phường Mẫu",
        "household_province": "Thành phố Hồ Chí Minh",
        "household_phone": "0900000000",
        "household_email": "hkd@example.com",
        "household_capital": "40000000",
        "household_labor_count": "2",
        "household_subject_type": "individual",
        "household_line_1_name": "Bán buôn tổng hợp",
        "household_line_1_detail": "Chi tiết bán buôn theo hồ sơ kiểm thử",
        "household_line_1_code": "4690",
        "household_line_1_is_main": "X",
        "household_line_2_name": "Bán lẻ mỹ phẩm",
        "household_line_2_detail": "Chi tiết bán lẻ mỹ phẩm",
        "household_line_2_code": "4772",
        "household_line_2_is_main": "",
        "authorized_person_full_name": "Trần Thị Ủy Quyền",
        "authorized_person_id_number": "031199001726",
        "authorized_person_issue_date": "09/12/2024",
        "authorized_person_issue_place": "Bộ Công an",
        "authorized_person_contact_address": "03 Đường Đại Diện",
        "authorized_person_phone": "0988888888",
        "authorized_person_email": "uyquyen@example.com",
    }


def non_yellow_text(document: Document) -> list[str]:
    values: list[str] = []

    def collect(container) -> None:
        for paragraph in container.paragraphs:
            values.extend(
                run.text
                for run in paragraph.runs
                if run.text and run.font.highlight_color is None
            )
        for table in container.tables:
            seen: set[int] = set()
            for row in table.rows:
                for cell in row.cells:
                    marker = id(cell._tc)
                    if marker in seen:
                        continue
                    seen.add(marker)
                    collect(cell)

    collect(document)
    return values


class HouseholdBusinessTest(unittest.TestCase):
    def test_typed_adapter_preserves_unknown_keys_and_replaces_member_rows(self) -> None:
        case = HouseholdBusinessCase(
            business_name="Hộ kinh doanh A",
            owner_name="Nguyễn Văn A",
            members=(HouseholdMember(full_name="Thành viên mới"),),
        )
        values = case.to_form_values({
            "future_key": "giữ nguyên",
            "household_member_8_full_name": "dòng đã xóa",
        })
        self.assertEqual("giữ nguyên", values["future_key"])
        self.assertEqual("Thành viên mới", values["household_member_1_full_name"])
        self.assertNotIn("household_member_8_full_name", values)

    def test_family_subject_is_blocked_when_template_has_no_yellow_member_slots(self) -> None:
        values = sample_values()
        values["household_subject_type"] = "family"
        issues = HouseholdBusinessTemplateGenerator.validation_issues(values)
        self.assertTrue(any("chưa có vùng vàng" in issue for issue in issues))

    def test_generator_fills_reviewed_yellow_slots_and_creates_two_documents(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            out = Path(raw)
            result = HouseholdBusinessTemplateGenerator(TEMPLATES, out).generate(sample_values())
            self.assertFalse(result.unresolved)
            self.assertEqual([REQUEST_TEMPLATE, AUTHORIZATION_TEMPLATE], [item.source_name for item in result.generated])
            request = Document(out / REQUEST_TEMPLATE)
            authorization = Document(out / AUTHORIZATION_TEMPLATE)
            self.assertEqual(
                non_yellow_text(Document(TEMPLATES / REQUEST_SOURCE_TEMPLATE)),
                non_yellow_text(request),
            )
            self.assertEqual(
                non_yellow_text(Document(TEMPLATES / AUTHORIZATION_SOURCE_TEMPLATE)),
                non_yellow_text(authorization),
            )
            self.assertIn("NGUYỄN VĂN KIỂM THỬ", request.paragraphs[5].text.upper())
            self.assertIn("HỘ KINH DOANH KIỂM THỬ", request.paragraphs[13].text)
            self.assertIn("40.000.000 VNĐ (Bốn mươi triệu đồng)", request.paragraphs[25].text)
            self.assertEqual(3, len(request.tables[1].rows))
            self.assertEqual("X", request.tables[1].cell(1, 3).text.strip())
            self.assertEqual(
                "Bán buôn tổng hợp\nChi tiết bán buôn theo hồ sơ kiểm thử",
                request.tables[1].cell(1, 1).text,
            )
            self.assertEqual(
                WD_ALIGN_PARAGRAPH.LEFT,
                request.tables[1].cell(1, 1).paragraphs[0].alignment,
            )
            main_runs = [
                run
                for run in request.tables[1].cell(1, 3).paragraphs[0].runs
                if run.text == "X"
            ]
            self.assertTrue(main_runs)
            self.assertTrue(all(run.bold for run in main_runs))
            self.assertEqual(
                "Bán lẻ mỹ phẩm\nChi tiết bán lẻ mỹ phẩm",
                request.tables[1].cell(2, 1).text,
            )
            self.assertIn("NGUYỄN VĂN KIỂM THỬ", request.tables[5].cell(0, 1).text.upper())
            self.assertIn("TRẦN THỊ ỦY QUYỀN", authorization.tables[0].cell(0, 0).text.upper())
            self.assertIn("NGUYỄN VĂN KIỂM THỬ", authorization.tables[0].cell(0, 1).text.upper())
            self.assertIn("03 Đường Đại Diện", authorization.paragraphs[13].text)
            self.assertIn("Cơ quan đăng ký kinh doanh cấp xã", authorization.paragraphs[16].text)
            self.assertIn("HỘ KINH DOANH KIỂM THỬ", authorization.paragraphs[16].text.upper())
            self.assertIn("Thành phố Hồ Chí Minh", authorization.paragraphs[24].text)
            self.assertIn("17", authorization.paragraphs[24].text)

    def test_template_hash_guard_rejects_an_unreviewed_revision(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source_names = (REQUEST_SOURCE_TEMPLATE, AUTHORIZATION_SOURCE_TEMPLATE)
            for name in source_names:
                (folder / name).write_bytes((TEMPLATES / name).read_bytes())
            with (folder / REQUEST_SOURCE_TEMPLATE).open("ab") as stream:
                stream.write(b"changed")
            issues = HouseholdBusinessTemplateGenerator(folder, folder).validate_templates()
            self.assertIn(REQUEST_TEMPLATE, issues)


if __name__ == "__main__":
    unittest.main()
