import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from mingo_assistant.services.documents.business_registration_generator import DocxTemplateGenerator
from mingo_assistant.services.normalization.address_parser import AddressParser
from mingo_assistant.repositories.authorized_person_repository import AuthorizedPersonProfileRepository
from mingo_assistant.ui.main_window import MainWindow
from mingo_assistant.ui.views.business_registration_workspace import BusinessRegistrationWorkspace


class BusinessRegistrationRegressionTest(unittest.TestCase):
    def test_labeled_manual_company_address_is_split_even_when_confirmation_is_needed(self) -> None:
        result = AddressParser().parse("Số 1, Phường 2, Thành phố Đà Nẵng")

        self.assertEqual(result.normalization_status, "needs_confirmation")
        self.assertEqual(result.street_line, "Số 1")
        self.assertEqual(result.ward, "Phường 2")
        self.assertEqual(result.province, "Thành phố Đà Nẵng")

    def test_power_of_attorney_replaces_authorized_person_address(self) -> None:
        document = Document()
        document.add_paragraph("BÊN NHẬN ỦY QUYỀN (Bên B)")
        address_paragraph = document.add_paragraph("Địa chỉ: địa chỉ mẫu")
        document.add_paragraph("ĐIỀU 1. NỘI DUNG ỦY QUYỀN")

        generator = DocxTemplateGenerator.__new__(DocxTemplateGenerator)
        generator._apply_power_of_attorney_authorized_address(
            document,
            "4. Giấy ủy quyền.docx",
            {"authorized_person_contact_address": "Số 2, Phường 3, Thành phố Đà Nẵng"},
        )

        self.assertEqual(address_paragraph.text, "Địa chỉ: Số 2, Phường 3, Thành phố Đà Nẵng")

    def test_beneficial_owner_address_uses_required_font_style(self) -> None:
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "Số 1, Phường 2, Thành phố Đà Nẵng"

        DocxTemplateGenerator._format_beneficial_owner_address_cell(cell, cell.text)

        run = cell.paragraphs[0].runs[0]
        self.assertFalse(run.italic)
        self.assertEqual(run.font.size, Pt(11.5))
        self.assertEqual(run.font.name, "Times New Roman")

    def test_beneficial_owner_place_date_uses_required_font_style(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Hà Nội, ngày 06 tháng 09 năm 2026")

        generator = DocxTemplateGenerator.__new__(DocxTemplateGenerator)
        generator._format_beneficial_owner_place_date(
            document,
            "3. danh sách chủ sở hữu hưởng lợi.docx",
            {
                "document_place": "Hà Nội",
                "document_day": "06",
                "document_month": "09",
                "document_year": "2026",
            },
        )

        run = paragraph.runs[0]
        self.assertTrue(run.italic)
        self.assertEqual(run.font.size, Pt(11.5))
        self.assertEqual(run.font.name, "Times New Roman")

    def test_real_saved_authorized_profile_with_legacy_builtin_id_is_preserved(self) -> None:
        window = MainWindow.__new__(MainWindow)
        project_root = Path(__file__).resolve().parents[2]
        window.authorized_person_repository = AuthorizedPersonProfileRepository(
            project_root / "data" / "temp" / "not-used.json",
            project_root / "assets" / "defaults" / "authorized_person_profiles.json",
        )
        saved_profile = {
            "id": "builtin-authorized-person-3",
            "authorized_person_full_name": "NGƯỜI DÙNG THẬT",
            "authorized_person_issue_place": "Cơ quan cấp",
            "authorized_person_contact_address": "Địa chỉ đã lưu",
        }

        payload = MainWindow._merge_authorized_person_profiles(
            window, saved_profile["id"], [saved_profile]
        )
        profile = next(
            item for item in payload["profiles"] if item["id"] == saved_profile["id"]
        )

        self.assertEqual(profile, saved_profile)

    def test_shipped_authorized_profiles_do_not_include_sample_people(self) -> None:
        project_root = Path(__file__).resolve().parents[2]
        repository = AuthorizedPersonProfileRepository(
            project_root / "data" / "temp" / "not-used.json",
            project_root / "assets" / "defaults" / "authorized_person_profiles.json",
        )
        profiles = repository.default_profiles()

        self.assertEqual(profiles, [])


    def test_cccd_sync_matches_only_the_same_identity_number(self) -> None:
        identities = {"012345678901"}

        self.assertTrue(BusinessRegistrationWorkspace._same_person_identity("012345678901", identities))
        self.assertTrue(BusinessRegistrationWorkspace._same_person_identity("CCCD 012345678901", identities))
        self.assertFalse(BusinessRegistrationWorkspace._same_person_identity("012345678902", identities))

    def test_case_person_data_is_not_overridden_by_a_global_profile(self) -> None:
        workspace = BusinessRegistrationWorkspace.__new__(BusinessRegistrationWorkspace)
        # Old process-wide data must never replace data belonging to this case.
        workspace.personal_profiles = [
            {
                "legal_representative_full_name": "NGUYEN VAN TEST",
                "legal_representative_id_number": "012345678901",
                "legal_representative_permanent_address": "\u0110\u1ecba ch\u1ec9 h\u1ed3 s\u01a1 kh\u00e1c",
            }
        ]
        person = type("Person", (), {"data": {
            "legal_representative_full_name": "NGUYEN VAN TEST",
            "legal_representative_id_number": "012345678901",
            "legal_representative_permanent_address": "\u0110\u1ecba ch\u1ec9 c\u1ee7a h\u1ed3 s\u01a1 n\u00e0y",
        }})()

        data = workspace._person_data_for_fill(person)

        self.assertEqual(data["legal_representative_permanent_address"], "\u0110\u1ecba ch\u1ec9 c\u1ee7a h\u1ed3 s\u01a1 n\u00e0y")
    def test_empty_company_email_replaces_template_sample_email(self) -> None:
        document = Document()
        paragraph = document.add_paragraph(
            "Th\u01b0 \u0111i\u1ec7n t\u1eed (n\u1ebfu c\u00f3): 4@gmail.com\tWebsite (n\u1ebfu c\u00f3): example.com"
        )
        generator = DocxTemplateGenerator.__new__(DocxTemplateGenerator)

        generator._apply_company_contact_fields(
            document,
            "1. Gi\u1ea5y \u0111\u1ec1 ngh\u1ecb.docx",
            {"company_email": "", "company_website": ""},
        )

        self.assertNotIn("4@gmail.com", paragraph.text)
        self.assertIn("..................", paragraph.text)
    def test_company_email_is_not_replaced_inside_its_new_value(self) -> None:
        document = Document()
        paragraph = document.add_paragraph(
            "Th\u01b0 \u0111i\u1ec7n t\u1eed (n\u1ebfu c\u00f3): 4@gmail.com\tWebsite (n\u1ebfu c\u00f3): example.com"
        )
        generator = DocxTemplateGenerator.__new__(DocxTemplateGenerator)

        generator._apply_company_contact_fields(
            document,
            "1. Gi\u1ea5y \u0111\u1ec1 ngh\u1ecb.docx",
            {"company_email": "234@gmail.com", "company_website": ""},
        )

        self.assertIn("234@gmail.com", paragraph.text)
        self.assertNotIn("23234@gmail.com", paragraph.text)

    def test_case_load_clears_authorized_contact_copied_to_different_applicant(self) -> None:
        values = {
            "applicant_full_name": "TRAN THI TEST",
            "applicant_id_number": "012345678902",
            "applicant_phone": "0900000000",
            "applicant_email": "test@example.invalid",
            "authorized_person_full_name": "NGUYEN VAN TEST",
            "authorized_person_id_number": "012345678901",
            "authorized_person_phone": "0900000000",
            "authorized_person_email": "test@example.invalid",
        }

        BusinessRegistrationWorkspace._clear_cross_person_authorized_contacts(values)

        self.assertEqual(values["applicant_phone"], "")
        self.assertEqual(values["applicant_email"], "")

    def test_case_load_keeps_contacts_when_applicant_is_authorized_person(self) -> None:
        values = {
            "applicant_id_number": "012345678901",
            "applicant_phone": "0900000000",
            "applicant_email": "test@example.invalid",
            "authorized_person_id_number": "012345678901",
            "authorized_person_phone": "0900000000",
            "authorized_person_email": "test@example.invalid",
        }

        BusinessRegistrationWorkspace._clear_cross_person_authorized_contacts(values)

        self.assertEqual(values["applicant_phone"], "0900000000")
        self.assertEqual(values["applicant_email"], "test@example.invalid")

    def test_document_generation_keeps_business_lines_beyond_ten(self) -> None:
        generator = DocxTemplateGenerator.__new__(DocxTemplateGenerator)
        values = {
            "business_line_1_name": "Ngành đầu tiên",
            "business_line_1_code": "1000",
            "business_line_11_name": "Ngành thứ mười một",
            "business_line_11_code": "8569",
        }

        lines = generator._business_lines(values)

        self.assertEqual([line["code"] for line in lines], ["1000", "8569"])
        self.assertEqual(lines[-1]["name"], "Ngành thứ mười một")

    def test_file1_business_line_stt_is_centered_and_does_not_wrap_after_nine(self) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=4)
        table.rows[0].cells[0].text = "STT"
        table.rows[0].cells[1].text = "Ten nganh"
        table.rows[0].cells[2].text = "Ma nganh"
        table.rows[0].cells[3].text = "Nganh chinh"
        table.rows[1].cells[0].paragraphs[0].paragraph_format.left_indent = Pt(24)
        generator = DocxTemplateGenerator.__new__(DocxTemplateGenerator)
        values = {
            f"business_line_{index}_name": f"Nganh {index}"
            for index in range(1, 11)
        }
        values.update(
            {
                f"business_line_{index}_code": f"{index:04d}"
                for index in range(1, 11)
            }
        )

        generator._replace_business_lines_tables_v2(
            document,
            values,
            keep_file1_business_line_numbers_on_one_line=True,
        )

        stt_cell = table.rows[10].cells[0]
        paragraph = stt_cell.paragraphs[0]
        self.assertEqual(stt_cell.text, "10")
        self.assertEqual(len(paragraph.runs), 1)
        self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertIsNone(paragraph.paragraph_format.left_indent)
        self.assertIn("w:noWrap", stt_cell._tc.xml)

    def test_file1_business_line_name_is_left_aligned_without_template_offsets(self) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=4)
        table.rows[0].cells[0].text = "STT"
        table.rows[0].cells[1].text = "Ten nganh"
        table.rows[0].cells[2].text = "Ma nganh"
        table.rows[0].cells[3].text = "Nganh chinh"
        source_paragraph = table.rows[1].cells[1].paragraphs[0]
        source_paragraph.paragraph_format.left_indent = Pt(36)
        source_paragraph._p.get_or_add_pPr().append(OxmlElement("w:tabs"))
        generator = DocxTemplateGenerator.__new__(DocxTemplateGenerator)

        generator._replace_business_lines_tables_v2(
            document,
            {
                "business_line_1_name": "Trong cay lay cu co chat bot",
                "business_line_1_code": "0113",
                "business_line_1_is_main": "X",
            },
            keep_file1_business_line_numbers_on_one_line=True,
        )

        name_paragraph = table.rows[1].cells[1].paragraphs[0]
        paragraph_properties = name_paragraph._p.get_or_add_pPr()
        self.assertEqual(name_paragraph.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertIsNone(name_paragraph.paragraph_format.left_indent)
        self.assertIsNone(paragraph_properties.find(qn("w:tabs")))

    def test_file1_business_line_values_use_times_new_roman_13_point(self) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=4)
        table.rows[0].cells[0].text = "STT"
        table.rows[0].cells[1].text = "Ten nganh"
        table.rows[0].cells[2].text = "Ma nganh"
        table.rows[0].cells[3].text = "Nganh chinh"
        generator = DocxTemplateGenerator.__new__(DocxTemplateGenerator)

        generator._replace_business_lines_tables_v2(
            document,
            {
                "business_line_1_name": "Trong cay lay cu co chat bot",
                "business_line_1_code": "0113",
                "business_line_1_is_main": "X",
            },
            keep_file1_business_line_numbers_on_one_line=True,
        )

        for cell in table.rows[1].cells:
            text_runs = [run for paragraph in cell.paragraphs for run in paragraph.runs if run.text]
            self.assertTrue(text_runs)
            for run in text_runs:
                self.assertEqual(run.font.name, "Times New Roman")
                self.assertEqual(run.font.size, Pt(13))

    def test_legacy_province_address_is_converted_from_administrative_data(self) -> None:
        result = AddressParser().parse(
            "Th\u00f4n B\u1ed3n Th\u00f4n, Trung An, V\u0169 Th\u01b0, Th\u00e1i B\u00ecnh"
        )

        self.assertEqual(result.normalization_status, "converted")
        self.assertEqual(result.village, "Th\u00f4n B\u1ed3n Th\u00f4n")
        self.assertEqual(result.ward, "Ph\u01b0\u1eddng V\u0169 Ph\u00fac")
        self.assertEqual(result.province, "H\u01b0ng Y\u00ean")

    def test_power_of_attorney_place_date_uses_13_point_font(self) -> None:
        document = Document()
        paragraph = document.add_paragraph(
            "Th\u00e0nh ph\u1ed1 H\u1ed3 Ch\u00ed Minh, ng\u00e0y 06 th\u00e1ng 10 n\u0103m 2026"
        )
        generator = DocxTemplateGenerator.__new__(DocxTemplateGenerator)

        generator._ensure_power_of_attorney_place_date_font(
            document,
            "4. Gi\u1ea5y \u1ee7y quy\u1ec1n.docx",
            {
                "document_place": "Th\u00e0nh ph\u1ed1 H\u1ed3 Ch\u00ed Minh",
                "document_day": "06",
                "document_month": "10",
                "document_year": "2026",
            },
        )

        self.assertEqual(paragraph.runs[0].font.size, Pt(13))
    def test_beneficial_owner_template_never_exports_nationality(self) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=18)
        table.rows[0].cells[0].text = "1"
        table.rows[0].cells[17].text = "18"
        generator = DocxTemplateGenerator.__new__(DocxTemplateGenerator)

        generator._apply_beneficial_owner_template_rules(
            document,
            "3. Danh s\u00e1ch ch\u1ee7 s\u1edf h\u1eefu h\u01b0\u1edfng l\u1ee3i.docx",
            {
                "beneficial_owner_list_number": "1",
                "beneficial_owner_list_full_name": "NGUYEN VAN A",
                "beneficial_owner_list_nationality": "Vi\u1ec7t Nam",
            },
        )

        self.assertEqual(table.rows[1].cells[5].text, "")

    def test_one_member_charter_fills_owner_nationality_and_country_only(self) -> None:
        template_dir = (
            Path(__file__).resolve().parents[2]
            / "assets"
            / "templates"
            / "word"
            / "business_registration"
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / DocxTemplateGenerator.PERSONAL_CHARTER_TEMPLATE
            shutil.copy2(template_dir / output_path.name, output_path)
            generator = DocxTemplateGenerator(template_dir, Path(temp_dir))
            foreign_values = generator._values_for_template(
                DocxTemplateGenerator.PERSONAL_CHARTER_TEMPLATE,
                {"legal_representative_nationality": "Úc"},
            )
            self.assertEqual("Úc", foreign_values["beneficial_owner_nationality"])

            generator._replace_placeholders(
                output_path,
                {
                    "company_name": "CÔNG TY TNHH MẪU",
                    "beneficial_owner_full_name": "NGUYỄN VĂN MẪU",
                    "document_place": "Quảng Ninh",
                    "document_day": "19",
                    "document_month": "10",
                    "document_year": "2026",
                },
            )

            document = Document(output_path)
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            self.assertIn("Quốc tịch: Việt Nam", paragraphs[6])
            self.assertIn("Quốc tịch: Việt Nam", paragraphs[42])
            self.assertIn("Quốc gia: Việt Nam", paragraphs[48])
            all_paragraphs = list(generator._iter_paragraphs(document))
            owner_heading_index = next(
                index
                for index, paragraph in enumerate(all_paragraphs)
                if generator._fold(paragraph.text) == "chu so huu"
            )
            self.assertEqual(
                "Quảng Ninh, ngày 19 tháng 10 năm 2026",
                all_paragraphs[owner_heading_index - 1].text,
            )
if __name__ == "__main__":
    unittest.main()
