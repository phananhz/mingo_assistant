from __future__ import annotations

import json
import re
import shutil
import tempfile
import unicodedata
import zipfile
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from lxml import etree


@dataclass(frozen=True)
class GeneratedDocument:
    source_name: str
    output_path: Path


@dataclass(frozen=True)
class GenerationResult:
    output_dir: Path
    generated: list[GeneratedDocument]
    skipped: list[str]
    unresolved: dict[str, list[str]]


class DocxTemplateGenerator:
    PERSONAL_CHARTER_TEMPLATE = "2. Điều lệ.docx"
    ORGANIZATION_CHARTER_TEMPLATE = "2. Dieu le_CSH la to chuc.docx"
    UPPERCASE_BOLD_FIELDS = {
        "company_name",
        "company_short_name",
        "company_english_name",
        "legal_representative_full_name",
        "director_full_name",
        "applicant_full_name",
        "authorized_person_full_name",
        "beneficial_owner_full_name",
        "organization_owner_name",
    }
    DETAIL_ADDRESS_FIELDS = {
        "company_address_detail",
        "applicant_address_detail",
        "legal_representative_address_detail",
        "organization_owner_address_detail",
    }
    WARD_FIELDS = {
        "company_ward",
        "applicant_ward",
        "legal_representative_ward",
        "organization_owner_ward",
    }
    PROVINCE_FIELDS = {
        "company_province",
        "applicant_province",
        "legal_representative_province",
        "organization_owner_province",
    }
    PROVINCE_AFTER_MERGER = {
        "ha_giang": "Tuyên Quang",
        "tuyen_quang": "Tuyên Quang",
        "yen_bai": "Lào Cai",
        "lao_cai": "Lào Cai",
        "bac_kan": "Thái Nguyên",
        "thai_nguyen": "Thái Nguyên",
        "vinh_phuc": "Phú Thọ",
        "hoa_binh": "Phú Thọ",
        "phu_tho": "Phú Thọ",
        "bac_giang": "Bắc Ninh",
        "bac_ninh": "Bắc Ninh",
        "thai_binh": "Hưng Yên",
        "hung_yen": "Hưng Yên",
        "hai_duong": "Thành phố Hải Phòng",
        "hai_phong": "Thành phố Hải Phòng",
        "ha_nam": "Ninh Bình",
        "nam_dinh": "Ninh Bình",
        "ninh_binh": "Ninh Bình",
        "quang_binh": "Quảng Trị",
        "quang_tri": "Quảng Trị",
        "quang_nam": "Thành phố Đà Nẵng",
        "da_nang": "Thành phố Đà Nẵng",
        "kon_tum": "Quảng Ngãi",
        "quang_ngai": "Quảng Ngãi",
        "binh_dinh": "Gia Lai",
        "gia_lai": "Gia Lai",
        "ninh_thuan": "Khánh Hòa",
        "khanh_hoa": "Khánh Hòa",
        "dak_nong": "Lâm Đồng",
        "binh_thuan": "Lâm Đồng",
        "lam_dong": "Lâm Đồng",
        "phu_yen": "Đắk Lắk",
        "dak_lak": "Đắk Lắk",
        "ba_ria_vung_tau": "Thành phố Hồ Chí Minh",
        "binh_duong": "Thành phố Hồ Chí Minh",
        "ho_chi_minh": "Thành phố Hồ Chí Minh",
        "tp_ho_chi_minh": "Thành phố Hồ Chí Minh",
        "tphcm": "Thành phố Hồ Chí Minh",
        "binh_phuoc": "Đồng Nai",
        "dong_nai": "Đồng Nai",
        "long_an": "Tây Ninh",
        "tay_ninh": "Tây Ninh",
        "soc_trang": "Thành phố Cần Thơ",
        "hau_giang": "Thành phố Cần Thơ",
        "can_tho": "Thành phố Cần Thơ",
        "ben_tre": "Vĩnh Long",
        "tra_vinh": "Vĩnh Long",
        "vinh_long": "Vĩnh Long",
        "tien_giang": "Đồng Tháp",
        "dong_thap": "Đồng Tháp",
        "bac_lieu": "Cà Mau",
        "ca_mau": "Cà Mau",
        "kien_giang": "An Giang",
        "an_giang": "An Giang",
    }

    def __init__(self, template_dir: Path, export_root: Path) -> None:
        self.template_dir = template_dir
        self.export_root = export_root
        self._manifest_cache: dict[str, object] | None = None

    def generate(self, values: dict[str, str], case_name: str, output_dir: Path | None = None) -> GenerationResult:
        if output_dir is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            folder_name = f"{self._safe_filename(case_name or 'ho_so_thanh_lap')}_{timestamp}"
            output_dir = self.export_root / folder_name
        output_dir.mkdir(parents=True, exist_ok=True)

        generated: list[GeneratedDocument] = []
        skipped: list[str] = []
        unresolved: dict[str, list[str]] = {}

        for template in self._selected_templates(values):
            if template.suffix.lower() == ".docx":
                template_output_path = output_dir / template.name
                output_path = output_dir / self._output_filename(template.name, values)
                shutil.copy2(template, template_output_path)
                # Placeholder manifests are keyed by the original template name.
                self._replace_placeholders(template_output_path, values)
                if output_path != template_output_path:
                    template_output_path.replace(output_path)
                remaining = self.unresolved_placeholders(output_path)
                if remaining:
                    unresolved[template.name] = remaining
                generated.append(GeneratedDocument(template.name, output_path))
            elif template.suffix.lower() == ".doc":
                skipped.append(f"{template.name} chưa được điền tự động vì đang là .doc")

        return GenerationResult(output_dir=output_dir, generated=generated, skipped=skipped, unresolved=unresolved)

    def _selected_templates(self, values: dict[str, str]) -> list[Path]:
        """Return the five dossier templates, selecting the charter by owner type."""
        owner_is_organization = self._fold(str(values.get("owner_type", "") or "")) == "to chuc"
        selected: list[Path] = []
        for template in sorted(self.template_dir.iterdir()):
            if template.name.startswith("~$"):
                continue
            if template.name == self.PERSONAL_CHARTER_TEMPLATE and owner_is_organization:
                continue
            if template.name == self.ORGANIZATION_CHARTER_TEMPLATE and not owner_is_organization:
                continue
            selected.append(template)
        return selected

    def _output_filename(self, template_name: str, values: dict[str, str]) -> str:
        def without_accents(value: str) -> str:
            text = unicodedata.normalize("NFD", value)
            text = "".join(character for character in text if unicodedata.category(character) != "Mn")
            return text.replace("Đ", "D").replace("đ", "d")

        def safe_component(value: str, fallback: str) -> str:
            text = without_accents(value)
            text = re.sub(r'[<>:"/\\|?*]', " ", text)
            text = re.sub(r"\s+", " ", text).strip(" .")
            return text[:80] or fallback

        raw_company_name = str(values.get("company_name", "") or "")
        company_name_without_legal_form = re.sub(
            r"^\s*(?:(?:công\s*ty|cty)\s+)?(?:tnhh(?:\s+(?:một\s+thành\s+viên|mtv|hai\s+thành\s+viên\s+trở\s+lên))?|cổ\s*phần|cp|hợp\s*danh|tư\s*nhân)\s+",
            "",
            raw_company_name,
            flags=re.IGNORECASE,
        )
        company_words = safe_component(company_name_without_legal_form, "Ho so thanh lap").split()
        company_name = "".join(word[:1].upper() + word[1:].lower() for word in company_words)
        procedure_names = {
            "1.": "Giay de nghi",
            "2.": "Dieu le",
            "3.": "Danh sach chu so huu huong loi",
            "4.": "Giay uy quyen",
            "5.": "Huong dan ky",
        }
        procedure_name = next(
            (name for prefix, name in procedure_names.items() if template_name.startswith(prefix)),
            safe_component(re.sub(r"^\d+\.\s*", "", Path(template_name).stem), "Tai lieu"),
        )
        return f"{company_name}_{procedure_name}.docx"
    def validate_templates(self, values: dict[str, str]) -> dict[str, list[str]]:
        report: dict[str, list[str]] = {}
        with tempfile.TemporaryDirectory(dir=self.export_root if self.export_root.exists() else None) as temp_dir:
            temp_path = Path(temp_dir)
            for template in self._selected_templates(values):
                if template.suffix.lower() != ".docx":
                    continue
                candidate = temp_path / template.name
                shutil.copy2(template, candidate)
                self._replace_placeholders(candidate, values)
                unresolved = self.unresolved_placeholders(candidate)
                if unresolved:
                    report[template.name] = sorted(set(unresolved))
        return report

    def unresolved_placeholders(self, docx_path: Path) -> list[str]:
        placeholders: list[str] = []
        with zipfile.ZipFile(docx_path, "r") as package:
            for item in package.infolist():
                if not item.filename.startswith("word/") or not item.filename.endswith(".xml"):
                    continue
                data = package.read(item.filename)
                if b"[" not in data or b"]" not in data:
                    continue
                placeholders.extend(self._extract_placeholders_from_xml(data))
        return sorted(set(placeholders))

    def _replace_placeholders(self, docx_path: Path, values: dict[str, str]) -> None:
        self._replace_bookmarked_fields_in_package(docx_path, values)
        self._replace_split_placeholders_in_package(docx_path, values)
        document = Document(docx_path)
        uses_highlighted_fields = self._replace_highlighted_runs(document, docx_path.name, values)
        self._apply_organization_owner_highlighted_fields(document, docx_path.name, values)
        self._apply_organization_charter_owner_fields(document, docx_path.name, values)
        self._apply_charter_article_6_capital(document, docx_path.name, values)
        self._apply_organization_owner_representative_table(document, docx_path.name, values)
        self._replace_manifest_plain_text(document, docx_path.name, values)
        self._apply_charter_legal_representative_gender(document, docx_path.name, values)
        self._apply_beneficial_owner_template_rules(document, docx_path.name, values)
        self._apply_power_of_attorney_authorized_address(document, docx_path.name, values)
        self._apply_company_contact_fields(document, docx_path.name, values)
        self._remove_file1_owner_name_spacer(document, docx_path.name, values)
        self._apply_document_place_date(document, docx_path.name, values)
        self._format_beneficial_owner_place_date(document, docx_path.name, values)
        self._apply_organization_charter_document_date(document, docx_path.name, values)
        self._normalize_legal_representative_signature(document, docx_path.name, values)
        if not uses_highlighted_fields:
            self._replace_case_type_text(document, values)
        self._replace_business_lines_tables_v2(
            document,
            values,
            keep_file1_business_line_numbers_on_one_line=docx_path.name.startswith("1."),
        )
        self._replace_asset_contribution_tables(document, values)
        self._normalize_file1_financial_tables(document, docx_path.name)
        if not uses_highlighted_fields:
            self._fix_registration_office_lines(document, values)
            self._fix_contact_lines(document, values)
            self._apply_bold_uppercase_formatting(document, values)
            self._fix_signature_font_sizes(document, docx_path.name, values)
            self._normalize_power_of_attorney_layout(document, docx_path.name, values)
        self._apply_template_checkbox_rules(document, docx_path.name, values)
        self._normalize_document_fonts(document)
        self._ensure_power_of_attorney_address_font(document, docx_path.name, values)
        self._ensure_power_of_attorney_place_date_font(document, docx_path.name, values)
        document.save(docx_path)

    def _replace_case_type_text(self, document: Document, values: dict[str, str]) -> None:
        try:
            member_count = int(values.get("company_member_count", "1") or "1")
        except ValueError:
            member_count = 1
        if member_count <= 1:
            return
        replacements = {
            "CÔNG TY TNHH MỘT THÀNH VIÊN": "CÔNG TY TNHH HAI THÀNH VIÊN TRỞ LÊN",
            "Công ty TNHH Một thành viên": "Công ty TNHH Hai thành viên trở lên",
            "Công ty trách nhiệm hữu hạn một thành viên": "Công ty trách nhiệm hữu hạn hai thành viên trở lên",
            "công ty trách nhiệm hữu hạn một thành viên": "công ty trách nhiệm hữu hạn hai thành viên trở lên",
            "trách nhiệm hữu hạn một thành viên": "trách nhiệm hữu hạn hai thành viên trở lên",
            "một thành viên": "hai thành viên trở lên",
        }
        replacements.update(
            {
                "C\u00d4NG TY TNHH M\u1ed8T TH\u00c0NH VI\u00caN": "C\u00d4NG TY TNHH HAI TH\u00c0NH VI\u00caN TR\u1ede L\u00caN",
                "C\u00f4ng ty TNHH m\u1ed9t th\u00e0nh vi\u00ean": "C\u00f4ng ty TNHH Hai th\u00e0nh vi\u00ean tr\u1edf l\u00ean",
                "C\u00f4ng ty tr\u00e1ch nhi\u1ec7m h\u1eefu h\u1ea1n m\u1ed9t th\u00e0nh vi\u00ean": "C\u00f4ng ty tr\u00e1ch nhi\u1ec7m h\u1eefu h\u1ea1n hai th\u00e0nh vi\u00ean tr\u1edf l\u00ean",
                "c\u00f4ng ty tr\u00e1ch nhi\u1ec7m h\u1eefu h\u1ea1n m\u1ed9t th\u00e0nh vi\u00ean": "c\u00f4ng ty tr\u00e1ch nhi\u1ec7m h\u1eefu h\u1ea1n hai th\u00e0nh vi\u00ean tr\u1edf l\u00ean",
                "tr\u00e1ch nhi\u1ec7m h\u1eefu h\u1ea1n m\u1ed9t th\u00e0nh vi\u00ean": "tr\u00e1ch nhi\u1ec7m h\u1eefu h\u1ea1n hai th\u00e0nh vi\u00ean tr\u1edf l\u00ean",
                "m\u1ed9t th\u00e0nh vi\u00ean": "hai th\u00e0nh vi\u00ean tr\u1edf l\u00ean",
            }
        )
        self._replace_plain_text_in_container(document, replacements)
        for section in document.sections:
            self._replace_plain_text_in_container(section.header, replacements)
            self._replace_plain_text_in_container(section.footer, replacements)

    def _replace_plain_text_in_container(self, container, replacements: dict[str, str]) -> None:
        element = getattr(container, "_element", None)
        if element is None:
            element = getattr(container, "element", None)
        if element is None:
            return
        for paragraph_element in element.xpath(".//w:p"):
            paragraph = Paragraph(paragraph_element, container)
            original = paragraph.text
            replaced = original
            for source, target in replacements.items():
                replaced = replaced.replace(source, target)
            if replaced == original:
                continue
            for run in paragraph.runs:
                run.text = ""
            if paragraph.runs:
                paragraph.runs[0].text = replaced
            else:
                paragraph.add_run(replaced)

    def _template_manifest(self) -> dict[str, object]:
        if self._manifest_cache is not None:
            return self._manifest_cache
        manifest_path = self.template_dir / "template_manifest.json"
        if not manifest_path.exists():
            self._manifest_cache = {}
            return self._manifest_cache
        with manifest_path.open("r", encoding="utf-8-sig") as file:
            loaded = json.load(file)
        self._manifest_cache = loaded if isinstance(loaded, dict) else {}
        return self._manifest_cache

    def _template_part_manifest(self, template_name: str, part_name: str) -> dict[str, object]:
        manifest = self._template_manifest()
        templates = manifest.get("templates", {})
        if not isinstance(templates, dict):
            return {}
        template_manifest = templates.get(template_name, {})
        if not isinstance(template_manifest, dict):
            return {}
        part_manifest = template_manifest.get(part_name, {})
        return part_manifest if isinstance(part_manifest, dict) else {}

    def _replace_highlighted_runs(self, document: Document, template_name: str, values: dict[str, str]) -> bool:
        part_manifest = self._template_part_manifest(template_name, "word/document.xml")
        field_map = part_manifest.get("highlighted_fields", {})
        if not isinstance(field_map, dict) or not field_map:
            return False

        highlighted_index = 0
        for paragraph in self._iter_paragraphs(document):
            for run_index, run in enumerate(paragraph.runs):
                if run.font.highlight_color != WD_COLOR_INDEX.YELLOW:
                    continue
                target = field_map.get(str(highlighted_index))
                if target is not None:
                    replacement = "" if self._clear_file1_individual_owner_field(
                        template_name, str(target), values
                    ) else self._value_for_template_target(str(target), values)
                    if str(target) == "company_english_name":
                        prefix = "".join(previous_run.text for previous_run in paragraph.runs[:run_index])
                        suffix = "".join(next_run.text for next_run in paragraph.runs[run_index + 1 :])
                        replacement = self._replacement_with_boundary_spacing(replacement, prefix, suffix)
                    if run._r.findall(qn("w:sym")):
                        self._set_symbol_checkbox_value(run, replacement)
                    else:
                        run.text = replacement
                highlighted_index += 1
        return True

    def _clear_file1_individual_owner_field(
        self, template_name: str, target: str, values: dict[str, str]
    ) -> bool:
        """Clear only section 5(a)'s personal sample values for an organization owner."""
        return (
            template_name.startswith("1.")
            and target.startswith("beneficial_owner_")
            and self._fold(str(values.get("owner_type", "") or "")) == "to chuc"
        )
    def _set_symbol_checkbox_value(self, run, value: str) -> None:
        """Preserve a Word symbol checkbox while selecting or clearing it."""
        checkbox_code = "F053" if str(value or "").strip() else "F052"
        for symbol in run._r.findall(qn("w:sym")):
            symbol.set(qn("w:char"), checkbox_code)

    def _replace_manifest_plain_text(self, document: Document, template_name: str, values: dict[str, str]) -> None:
        part_manifest = self._template_part_manifest(template_name, "word/document.xml")
        plain_replacements = part_manifest.get("plain_replacements", [])
        if not isinstance(plain_replacements, list):
            return

        replacements: dict[str, str] = {}
        for item in plain_replacements:
            if not isinstance(item, dict):
                continue
            source = str(item.get("text", ""))
            target = str(item.get("field", ""))
            if not source or not target:
                continue
            replacements[source] = self._value_for_template_target(target, values)
        if not replacements:
            return
        self._replace_plain_text_in_container(document, replacements)

    def _apply_charter_legal_representative_gender(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        """Fill the unhighlighted gender sample with stable spacing after a long name."""
        if not template_name.startswith("2."):
            return
        gender = str(values.get("legal_representative_gender", "") or "").strip()
        for paragraph in self._iter_paragraphs(document):
            if "ho va ten nguoi dai dien theo phap luat" not in self._fold(paragraph.text):
                continue
            runs = list(paragraph.runs)
            label_index = next(
                (index for index, run in enumerate(runs) if "gioi tinh" in self._fold(run.text)),
                None,
            )
            if label_index is None:
                continue
            value_index = label_index + 1
            if label_index and "\t" in runs[label_index - 1].text:
                runs[label_index - 1].text = "     "
            if value_index < len(runs):
                runs[value_index].text = gender
                for run in runs[value_index + 1 :]:
                    run.text = ""
            else:
                paragraph.add_run(gender)
            return

    def _apply_power_of_attorney_authorized_address(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        """Fill Bên B's address even though this template has no highlighted run there."""
        if not template_name.startswith("4. "):
            return
        address = str(values.get("authorized_person_contact_address", "") or "").strip()
        if not address:
            return
        in_authorized_party = False
        for paragraph in self._iter_paragraphs(document):
            folded = self._fold(paragraph.text)
            if "ben nhan uy quyen" in folded:
                in_authorized_party = True
                continue
            if not in_authorized_party:
                continue
            if "dieu 1" in folded:
                return
            if folded.startswith("dia chi"):
                self._replace_paragraph_text(paragraph, f"Địa chỉ: {address}")
                return

    def _apply_beneficial_owner_template_rules(self, document: Document, template_name: str, values: dict[str, str]) -> None:
        """Fill the existing beneficial-owner row while preserving the template table geometry."""
        if not template_name.startswith("3."):
            return
        owner_columns = [str(values.get(key, "") or "").strip() for key in ("beneficial_owner_list_number", "beneficial_owner_list_full_name", "beneficial_owner_list_date_of_birth", "beneficial_owner_list_gender", "beneficial_owner_list_identity", "beneficial_owner_list_nationality", "beneficial_owner_list_ethnicity", "beneficial_owner_list_address")]
        owner_columns[4] = self._beneficial_owner_identity_number(owner_columns[4])
        # Column 6 is intentionally left blank in this dossier template.
        # Ignore legacy saved values and any CCCD/profile-derived nationality.
        owner_columns[5] = ""
        column_values = [str(values.get(key, "") or "").strip() for key in ("beneficial_owner_list_capital_direct", "beneficial_owner_list_capital_indirect", "beneficial_owner_list_capital_total", "beneficial_owner_list_voting_direct", "beneficial_owner_list_voting_indirect", "beneficial_owner_list_voting_total", "beneficial_owner_list_general_partner", "beneficial_owner_list_related_group", "beneficial_owner_list_control", "beneficial_owner_list_manager")]
        column_values[:6] = [self._beneficial_owner_percentage(value) for value in column_values[:6]]
        for table in self._iter_tables(document):
            numbering_row_index = next((index for index, row in enumerate(table.rows) if len(row.cells) >= 18 and row.cells[0].text.strip() == "1" and row.cells[17].text.strip() == "18"), None)
            if numbering_row_index is None:
                continue
            for row in table.rows[numbering_row_index + 1:]:
                if len(row.cells) < 18:
                    continue
                for column, (cell, value) in enumerate(zip(row.cells[:8], owner_columns, strict=True), start=1):
                    self._replace_table_cell_text(cell, value)
                    if column == 8:
                        self._format_beneficial_owner_address_cell(cell, value)
                for column, (cell, value) in enumerate(zip(row.cells[8:18], column_values, strict=True), start=9):
                    self._replace_table_cell_text(cell, value)
                    self._format_beneficial_owner_numeric_cell(cell, value)

    @staticmethod
    def _beneficial_owner_identity_number(value: str) -> str:
        match = re.search(r"(?:CCCD\s*số|Số\s*định\s*danh\s*cá\s*nhân)\s*:?\s*([0-9A-Za-z-]+)", value, re.IGNORECASE)
        if match:
            return f"Số định danh cá nhân: {match.group(1)}"
        fallback = re.search(r"[0-9]{6,}", value)
        return f"Số định danh cá nhân: {fallback.group(0)}" if fallback else value

    @staticmethod
    def _beneficial_owner_percentage(value: str) -> str:
        compact = re.sub(r"\s+", "", str(value or "").strip())
        if not compact:
            return ""
        return compact if compact.endswith("%") else f"{compact}%"

    @staticmethod
    def _format_beneficial_owner_numeric_cell(cell, value: str) -> None:
        """Use the compact centered style of column 9 for columns 9 onward."""
        paragraphs = list(cell.paragraphs)
        if not paragraphs:
            paragraphs = [cell.add_paragraph()]
        first = paragraphs[0]
        for paragraph in paragraphs[1:]:
            element = paragraph._element
            element.getparent().remove(element)
        first.alignment = WD_ALIGN_PARAGRAPH.CENTER
        first.paragraph_format.left_indent = Inches(0)
        first.paragraph_format.right_indent = Inches(0)
        first.paragraph_format.first_line_indent = Inches(0)
        first.paragraph_format.space_before = Pt(0)
        first.paragraph_format.space_after = Pt(0)
        first.paragraph_format.line_spacing = 1
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in first.runs:
            if run.text:
                run.font.size = Pt(10)
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW if str(value or "").strip() else None

    @staticmethod
    def _format_beneficial_owner_address_cell(cell, value: str) -> None:
        """Use the required upright Times New Roman 11.5pt address style."""
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                if not run.text:
                    continue
                run.font.name = "Times New Roman"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                run.font.size = Pt(11.5)
                run.italic = False
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW if str(value or "").strip() else None

    def _replace_table_cell_text(self, cell, text: str) -> None:
        paragraphs = list(cell.paragraphs)
        if not paragraphs:
            paragraphs = [cell.add_paragraph()]
        for paragraph in paragraphs:
            for run in paragraph.runs:
                run.text = ""
        first = paragraphs[0]
        if first.runs:
            first.runs[0].text = text
        else:
            first.add_run(text)
    def _apply_organization_owner_highlighted_fields(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        """Fill the organization-owner section without changing its labels."""
        if not template_name.startswith("1.") or self._fold(str(values.get("owner_type", "") or "")) != "to chuc":
            return

        in_organization_owner_section = False
        for paragraph in self._iter_paragraphs(document):
            context = self._fold(paragraph.text)
            if "doi voi chu so huu la to chuc" in context:
                in_organization_owner_section = True
                continue
            if in_organization_owner_section and context.startswith("6."):
                break
            if not in_organization_owner_section:
                continue

            target = ""
            if "ten to chuc" in context:
                target = "organization_owner_name"
            elif "ma so doanh nghiep" in context or "so quyet dinh thanh lap" in context:
                target = "organization_owner_registration_number"
            elif "ngay cap" in context and "noi cap" in context:
                self._replace_paragraph_with_highlighted_values(
                    paragraph,
                    (
                        ("Ngày cấp: ", False),
                        (self._value_for_template_target("organization_owner_issue_date", values), True),
                        ("\tNơi cấp: ", False),
                        (self._value_for_template_target("organization_owner_issue_place", values), True),
                    ),
                )
                continue
            elif "so nha/phong" in context:
                target = "organization_owner_address_detail"
            elif "xa/phuong" in context or "dac khu" in context:
                target = "organization_owner_ward"
            elif "tinh/thanh pho" in context:
                target = "organization_owner_province"
            elif context.startswith("quoc gia"):
                target = "organization_owner_country"
            elif "dien thoai" in context and "so fax" in context:
                self._replace_paragraph_with_highlighted_values(
                    paragraph,
                    (
                        ("Điện thoại (nếu có): ", False),
                        (self._value_for_template_target("organization_owner_phone", values), True),
                        ("\tSố fax (nếu có): ", False),
                        (self._value_for_template_target("organization_owner_fax", values), True),
                    ),
                )
                continue
            elif "thu dien tu" in context and "website" in context:
                self._replace_paragraph_with_highlighted_values(
                    paragraph,
                    (
                        ("Thư điện tử (nếu có): ", False),
                        (self._value_for_template_target("organization_owner_email", values), True),
                        ("\tWebsite (nếu có): ", False),
                        (self._value_for_template_target("organization_owner_website", values), True),
                    ),
                )
                continue
            if not target:
                continue

            label = paragraph.text.split(":", 1)[0].rstrip() + ": "
            self._replace_paragraph_with_highlighted_values(
                paragraph,
                ((label, False), (self._value_for_template_target(target, values), True)),
            )

    def _apply_charter_article_6_capital(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        """Keep Article 6 capital wording exact in both charter variants."""
        if template_name not in {self.PERSONAL_CHARTER_TEMPLATE, self.ORGANIZATION_CHARTER_TEMPLATE}:
            return
        number = self._with_vnd(values.get("charter_capital_number", ""))
        words = str(values.get("charter_capital_text", "") or "").strip()
        if words and not re.search(r"\sđồng$", words, re.IGNORECASE):
            words = f"{words} đồng"
        if not (number and words):
            return
        for paragraph in self._iter_paragraphs(document):
            if "von dieu le cua cong ty la" not in self._fold(paragraph.text):
                continue
            self._replace_paragraph_with_highlighted_values(
                paragraph,
                (
                    ("1. Vốn điều lệ của công ty là: ", False),
                    (f"{number} (Ghi bằng chữ: {words})", True),
                    (", trong đó bao gồm:", False),
                ),
            )
            return

    def _apply_organization_charter_owner_fields(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        """Fill the turquoise organization and capital-representative fields in charter 2."""
        if template_name != self.ORGANIZATION_CHARTER_TEMPLATE:
            return

        targets = (
            "organization_owner_name",
            "organization_owner_registration_number",
            "organization_owner_issue_date",
            "organization_owner_issue_place",
            "organization_owner_address_detail",
            "organization_owner_ward",
            "organization_owner_province",
            "organization_owner_country",
            "organization_owner_phone",
            "organization_owner_fax",
            "organization_owner_email",
            "organization_owner_website",
            "organization_owner_representative_full_name",
            "organization_owner_representative_date_of_birth",
            "organization_owner_representative_id_number",
            "organization_owner_representative_contact_address",
            "organization_owner_name",
            "organization_owner_representative_full_name",
        )
        highlighted_runs = [
            run
            for paragraph in self._iter_paragraphs(document)
            for run in paragraph.runs
            if run.font.highlight_color == WD_COLOR_INDEX.TURQUOISE
        ]
        for run, target in zip(highlighted_runs, targets, strict=False):
            run.text = self._value_for_template_target(target, values)
    def _apply_organization_owner_representative_table(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        """Fill the organization owner's 13-column representative row in template 1."""
        if not template_name.startswith("1.") or self._fold(str(values.get("owner_type", "") or "")) != "to chuc":
            return

        field_keys = (
            "organization_owner_representative_number",
            "organization_owner_representative_full_name",
            "organization_owner_representative_date_of_birth",
            "organization_owner_representative_gender",
            "organization_owner_representative_identity",
            "organization_owner_representative_nationality",
            "organization_owner_representative_ethnicity",
            "organization_owner_representative_contact_address",
            "organization_owner_representative_authorized_capital_value",
            "organization_owner_representative_authorized_capital_ratio",
            "organization_owner_representative_authorized_capital_date",
            "organization_owner_representative_signature",
            "organization_owner_representative_notes",
        )
        for table in self._iter_tables(document):
            if len(table.rows) < 4 or len(table.rows[-1].cells) != len(field_keys):
                continue
            header = self._fold(" ".join(cell.text for row in table.rows[:2] for cell in row.cells))
            if "ten nguoi dai dien theo phap luat" not in header or "von duoc uy quyen" not in header:
                continue
            for cell, field_key in zip(table.rows[-1].cells, field_keys):
                self._set_cell_text_with_highlight(
                    cell, self._value_for_template_target(field_key, values)
                )
            return

    def _replace_paragraph_with_highlighted_values(
        self, paragraph, segments: tuple[tuple[str, bool], ...]
    ) -> None:
        p_element = paragraph._p
        for child in list(p_element):
            if child.tag != qn("w:pPr"):
                p_element.remove(child)
        for text, highlight in segments:
            if not text:
                continue
            run = paragraph.add_run(str(text))
            if highlight:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def _set_cell_text_with_highlight(self, cell, value: str) -> None:
        self._set_cell_text(cell, value)
        if not str(value or "").strip():
            return
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                if run.text:
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    def _apply_company_contact_fields(self, document: Document, template_name: str, values: dict[str, str]) -> None:
        if not template_name.startswith("1."):
            return
        email = str(values.get("company_email", "") or "").strip()
        website = str(values.get("company_website", "") or "").strip()
        in_organization_owner_section = False
        for paragraph in self._iter_paragraphs(document):
            folded = self._fold(paragraph.text)
            if "doi voi chu so huu la to chuc" in folded:
                in_organization_owner_section = True
            elif in_organization_owner_section and folded.startswith("6."):
                in_organization_owner_section = False
            if "thu dien tu" not in folded or "website" not in folded:
                continue
            if in_organization_owner_section:
                continue
            self._replace_paragraph_with_highlighted_values(
                paragraph,
                (
                    ("Thư điện tử (nếu có): ", False),
                    (email or "..................", bool(email)),
                    ("\tWebsite (nếu có): ", False),
                    (website or "..................", bool(website)),
                ),
            )
            try:
                paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(4.1))
            except Exception:
                pass

    def _remove_file1_owner_name_spacer(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        """Remove the section 5(a) spacer only for an individual owner."""
        if (
            not template_name.startswith("1.")
            or self._fold(str(values.get("owner_type", "") or "")) != "ca nhan"
        ):
            return
        paragraphs = list(self._iter_paragraphs(document))
        for index, paragraph in enumerate(paragraphs[:-1]):
            if "ho, chu dem va ten chu so huu" not in self._fold(paragraph.text):
                continue
            spacer = paragraphs[index + 1]
            if spacer.text.strip():
                return
            spacer._element.getparent().remove(spacer._element)
            return
    def _apply_template_checkbox_rules(self, document: Document, template_name: str, values: dict[str, str]) -> None:
        """Update the selectable checkboxes already present in file 1."""
        if not template_name.startswith("1."):
            return
        paragraphs = list(self._iter_paragraphs(document))
        accounting_blank = self._checkbox_target_cell(
            self._find_paragraph_by_folded_text(paragraphs, "hach toan phu thuoc"),
            1,
        )
        vat_blank = self._checkbox_target_cell(
            self._find_paragraph_by_folded_text(paragraphs, "truc tiep tren gtgt"),
            1,
        )
        period_blank = self._checkbox_target_cell(
            self._find_paragraph_by_folded_text(paragraphs, "03 thang mot lan"),
        )
        pending_bot_choice = False
        bot_blank = None
        for index, paragraph in enumerate(paragraphs):
            folded = self._fold(paragraph.text)
            if "giay chung nhan quyen su dung dat" in folded and "quoc phong" in folded:
                label = paragraph.text.split("Có", 1)[0].rstrip(" :\t") + ":"
                self._set_checkbox_options(
                    paragraph,
                    label,
                    (("Có", bool(values.get("company_land_use_certificate", "").strip())), ("Không", not bool(values.get("company_land_use_certificate", "").strip()))),
                    new_line=True,
                )
                continue
            if "co hien thi thong tin ve gia tri tuong duong" in folded:
                if index + 1 >= len(paragraphs):
                    continue
                choices = paragraphs[index + 1]
                choices_folded = self._fold(choices.text)
                if "co" in choices_folded and "khong" in choices_folded:
                    show_foreign_currency = bool(str(values.get("foreign_currency_equivalent", "") or "").strip())
                    self._set_checkbox_options(
                        choices,
                        "",
                        (("Có", show_foreign_currency), ("Không", not show_foreign_currency)),
                    )
                continue
            if "hoat dong theo du an bot" in folded:
                pending_bot_choice = True
                continue
            if pending_bot_choice and folded == "co":
                bot_blank = self._checkbox_target_cell(paragraph)
                self._set_checkbox_cell(
                    bot_blank,
                    bool(values.get("tax_bot_project", "").strip()),
                    bot_blank,
                )
                continue
            if pending_bot_choice and folded == "khong":
                self._set_checkbox_cell(
                    self._checkbox_target_cell(paragraph),
                    not bool(values.get("tax_bot_project", "").strip()),
                    bot_blank,
                )
                pending_bot_choice = False
                continue
            if folded == "hach toan doc lap":
                self._set_checkbox_cell(
                    self._checkbox_target_cell(paragraph, 1),
                    self._fold(str(values.get("tax_accounting_method", ""))) == self._fold("Hạch toán độc lập"),
                    accounting_blank,
                )
                continue
            if folded == "co bao cao tai chinh hop nhat":
                target = self._checkbox_target_cell(paragraph)
                self._set_checkbox_cell(
                    target,
                    bool(values.get("tax_consolidated_financial_statement", "").strip()),
                    target,
                )
                continue
            if folded == "hach toan phu thuoc":
                self._set_checkbox_cell(
                    self._checkbox_target_cell(paragraph, 1),
                    self._fold(str(values.get("tax_accounting_method", ""))) == self._fold("Hạch toán phụ thuộc"),
                    accounting_blank,
                )
                continue
            if "ap dung tu ngay" in folded and "den ngay" in folded:
                start = str(values.get("tax_fiscal_year_start", "01/01") or "01/01").strip()
                end = str(values.get("tax_fiscal_year_end", "31/12") or "31/12").strip()
                self._replace_paragraph_text(paragraph, f"Áp dụng từ ngày {start} đến ngày {end}")
                continue
            if "tong so lao dong" in folded:
                labor_count = str(values.get("tax_labor_count", "") or "").strip()
                if labor_count:
                    self._replace_paragraph_text(paragraph, f"Tổng số lao động (dự kiến): {labor_count}")
                continue
            vat_labels = {
                "khau tru": "Khấu trừ",
                "truc tiep tren gtgt": "Trực tiếp trên GTGT",
                "truc tiep tren doanh so": "Trực tiếp trên doanh số",
                "khong phai nop thue gtgt": "Không phải nộp thuế GTGT",
            }
            if folded in vat_labels:
                label = vat_labels[folded]
                self._set_checkbox_cell(
                    self._checkbox_target_cell(paragraph, 1),
                    self._fold(str(values.get("tax_vat_method", ""))) == self._fold(label),
                    vat_blank,
                    align_left=True,
                )
                continue
            period_labels = {
                "hang thang": "Hàng tháng",
                "03 thang mot lan": "03 tháng một lần",
                "06 thang mot lan": "06 tháng một lần",
            }
            if folded in period_labels:
                label = period_labels[folded]
                self._set_checkbox_cell(
                    self._checkbox_target_cell(paragraph),
                    self._fold(str(values.get("tax_declaration_period", ""))) == self._fold(label),
                    period_blank,
                )

    def _find_paragraph_by_folded_text(self, paragraphs: list[Paragraph], text: str) -> Paragraph | None:
        return next((paragraph for paragraph in paragraphs if self._fold(paragraph.text) == text), None)

    def _checkbox_target_cell(self, paragraph: Paragraph | None, offset: int = 0):
        if paragraph is None or not hasattr(paragraph._parent, "_tc"):
            return None
        cell = paragraph._parent
        row = cell._tc.getparent()
        cells = list(row.tc_lst)
        try:
            return cells[cells.index(cell._tc) + offset]
        except (IndexError, ValueError):
            return None

    def _set_checkbox_cell(self, target_cell, checked: bool, blank_source_cell, align_left: bool = False) -> None:
        """Toggle an existing checkbox without recreating its label or table layout."""
        if target_cell is None:
            return
        if align_left:
            self._align_checkbox_cell_left(target_cell)
        symbols = list(target_cell.findall(".//" + qn("w:sym")))
        drawings = list(target_cell.findall(".//" + qn("w:drawing")))
        if checked:
            if symbols:
                self._set_symbol_checkbox_value_from_element(symbols[0], True)
                for symbol in symbols[1:]:
                    symbol.getparent().remove(symbol)
                for drawing in drawings:
                    drawing.getparent().remove(drawing)
                return
            if drawings:
                drawing = drawings[0]
                run = drawing.getparent()
                position = run.index(drawing)
                run.remove(drawing)
                symbol = OxmlElement("w:sym")
                symbol.set(qn("w:font"), "Wingdings 2")
                symbol.set(qn("w:char"), "F053")
                run.insert(position, symbol)
                for extra_drawing in drawings[1:]:
                    extra_drawing.getparent().remove(extra_drawing)
            return
        if drawings:
            for symbol in symbols:
                symbol.getparent().remove(symbol)
            return
        if not symbols or blank_source_cell is None:
            return
        blank_drawings = blank_source_cell.findall(".//" + qn("w:drawing"))
        if not blank_drawings:
            return
        symbol = symbols[0]
        run = symbol.getparent()
        position = run.index(symbol)
        run.remove(symbol)
        run.insert(position, deepcopy(blank_drawings[0]))
        for extra_symbol in symbols[1:]:
            extra_symbol.getparent().remove(extra_symbol)

    def _align_checkbox_cell_left(self, target_cell) -> None:
        paragraph = target_cell.find(qn("w:p"))
        if paragraph is None:
            return
        for text in list(paragraph.findall(".//" + qn("w:t"))):
            if not (text.text or "").strip():
                text.getparent().remove(text)
        paragraph_properties = paragraph.get_or_add_pPr()
        alignment = paragraph_properties.find(qn("w:jc"))
        if alignment is None:
            alignment = OxmlElement("w:jc")
            paragraph_properties.append(alignment)
        alignment.set(qn("w:val"), "left")
    def _set_symbol_checkbox_value_from_element(self, symbol, checked: bool) -> None:
        symbol.set(qn("w:font"), "Wingdings 2")
        symbol.set(qn("w:char"), "F053" if checked else "F052")

    def _set_checkbox_options(self, paragraph, prefix: str, options: tuple[tuple[str, bool], ...], new_line: bool = False) -> None:
        """Render plain inline checkboxes after their labels on a single baseline."""
        p_element = paragraph._p
        for child in list(p_element):
            if child.tag != qn("w:pPr"):
                p_element.remove(child)
        if prefix:
            paragraph.add_run(prefix)
        if new_line:
            paragraph.add_run().add_break()
        for index, (label, checked) in enumerate(options):
            if index:
                paragraph.add_run("        ")
            text_run = paragraph.add_run(f"{label} ")
            text_run.font.name = "Times New Roman"
            text_run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
            text_run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
            checkbox = paragraph.add_run("☒" if checked else "☐")
            checkbox.font.name = "Segoe UI Symbol"
            checkbox.font.size = Pt(11)
            checkbox._element.rPr.rFonts.set(qn("w:ascii"), "Segoe UI Symbol")
            checkbox._element.rPr.rFonts.set(qn("w:hAnsi"), "Segoe UI Symbol")

    def _normalize_document_fonts(self, document: Document) -> None:
        """Use Times New Roman consistently while preserving Wingdings checkboxes."""
        for paragraph in self._iter_paragraphs(document):
            for run in paragraph.runs:
                if run._r.findall(qn("w:sym")) or "☒" in run.text or "☐" in run.text:
                    continue
                run.font.name = "Times New Roman"
                r_pr = run._r.get_or_add_rPr()
                fonts = r_pr.find(qn("w:rFonts"))
                if fonts is None:
                    fonts = OxmlElement("w:rFonts")
                    r_pr.insert(0, fonts)
                for font_name in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                    fonts.set(qn(font_name), "Times New Roman")
    def _apply_organization_charter_document_date(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        """The organization charter has a split date line; rewrite it as one stable value."""
        if template_name != self.ORGANIZATION_CHARTER_TEMPLATE:
            return
        place = str(values.get("document_place", "") or "").strip()
        day = str(values.get("document_day", "") or "").strip()
        month = str(values.get("document_month", "") or "").strip()
        year = str(values.get("document_year", "") or "").strip()
        if not (place and day and month and year):
            return
        if day.isdigit():
            day = day.zfill(2)
        if month.isdigit():
            month = month.zfill(2)
        date_line = f"{place}, ngày {day} tháng {month} năm {year}"
        paragraphs = list(self._iter_paragraphs(document))
        for index, paragraph in enumerate(paragraphs[:-1]):
            if self._fold(paragraphs[index + 1].text) != "chu so huu":
                continue
            self._replace_paragraph_with_highlighted_values(paragraph, ((date_line, True),))
            return

    def _ensure_power_of_attorney_address_font(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        if not template_name.startswith("4. "):
            return
        addresses = {
            str(values.get(key, "") or "").strip()
            for key in (
                "legal_representative_permanent_address",
                "legal_representative_contact_address",
                "authorized_person_contact_address",
            )
        } - {""}
        for paragraph in self._iter_paragraphs(document):
            if not any(address in paragraph.text for address in addresses):
                continue
            for run in paragraph.runs:
                if run.text:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(13)
                    r_pr = run._r.get_or_add_rPr()
                    fonts = r_pr.rFonts
                    for font_name in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                        fonts.set(qn(font_name), "Times New Roman")

    def _ensure_power_of_attorney_place_date_font(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        """Keep the date line in the power-of-attorney template at 13pt."""
        if not template_name.startswith("4. "):
            return
        place = str(values.get("document_place", "") or "").strip()
        day = str(values.get("document_day", "") or "").strip()
        month = str(values.get("document_month", "") or "").strip()
        year = str(values.get("document_year", "") or "").strip()
        if not (place and day and month and year):
            return
        folded_line = self._fold(f"{place}, ngay {day} thang {month} nam {year}")
        for paragraph in self._iter_paragraphs(document):
            if self._fold(paragraph.text) != folded_line:
                continue
            for run in paragraph.runs:
                if not run.text:
                    continue
                run.font.name = "Times New Roman"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                run.font.size = Pt(13)
            return
    def _apply_document_place_date(self, document: Document, template_name: str, values: dict[str, str]) -> None:
        place = str(values.get("document_place", "") or "").strip()
        day = str(values.get("document_day", "") or "").strip()
        month = str(values.get("document_month", "") or "").strip()
        year = str(values.get("document_year", "") or "").strip()
        if not (place and day and month and year):
            return
        date_line = f"{place}, ngày {day} tháng {month} năm {year}"
        for paragraph in self._iter_paragraphs(document):
            folded = self._fold(paragraph.text)
            if not (folded.startswith(("thanh pho ", "tinh ")) and "ngay" in folded and "thang" in folded and "nam" in folded):
                continue
            self._replace_paragraph_text(paragraph, date_line)

    def _format_beneficial_owner_place_date(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        if not template_name.startswith("3."):
            return
        place = str(values.get("document_place", "") or "").strip()
        day = str(values.get("document_day", "") or "").strip()
        month = str(values.get("document_month", "") or "").strip()
        year = str(values.get("document_year", "") or "").strip()
        if not (place and day and month and year):
            return
        folded_line = self._fold(f"{place}, ngay {day} thang {month} nam {year}")
        for paragraph in self._iter_paragraphs(document):
            if self._fold(paragraph.text) != folded_line:
                continue
            for run in paragraph.runs:
                if not run.text:
                    continue
                run.font.name = "Times New Roman"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                run.font.size = Pt(11.5)
                run.italic = True
            return

    def _normalize_legal_representative_signature(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        if not template_name.startswith(("1.", "3.")):
            return
        signer = str(values.get("legal_representative_full_name", "") or "").strip().upper()
        if not signer:
            return
        folded_signer = self._fold(signer)
        for table in self._iter_tables(document):
            for row in table.rows:
                for cell in row.cells:
                    paragraphs = list(cell.paragraphs)
                    signer_indexes = [
                        index for index, paragraph in enumerate(paragraphs)
                        if folded_signer in self._fold(paragraph.text)
                    ]
                    if not signer_indexes:
                        continue
                    signer_index = signer_indexes[-1]
                    signature_context = " ".join(
                        paragraph.text for paragraph in paragraphs[max(0, signer_index - 8) : signer_index]
                    )
                    if "ky va ghi ho ten" not in self._fold(signature_context):
                        continue
                    tr_pr = row._tr.get_or_add_trPr()
                    for flag in tr_pr.findall(qn("w:cantSplit")):
                        tr_pr.remove(flag)
                    for paragraph in paragraphs:
                        paragraph.paragraph_format.keep_together = False
                        paragraph.paragraph_format.keep_with_next = False
                    signer_paragraph = paragraphs[signer_index]
                    signer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    signer_paragraph.paragraph_format.space_before = Pt(0)
                    signer_paragraph.paragraph_format.space_after = Pt(0)
                    for run in signer_paragraph.runs:
                        if run.text:
                            run.bold = True
                            run.italic = False
                            run.font.size = Pt(13)

    def _value_for_template_target(self, target: str, values: dict[str, str]) -> str:
        if target == "__clear":
            return ""
        if target.endswith("_date_of_birth_day"):
            return self._date_component(values.get(target.removesuffix("_day"), ""), 0)
        if target.endswith("_date_of_birth_month"):
            return self._date_component(values.get(target.removesuffix("_month"), ""), 1)
        if target.endswith("_date_of_birth_year"):
            return self._date_component(values.get(target.removesuffix("_year"), ""), 2)
        if target == "document_year_with_prefix":
            year = str(values.get("document_year", "") or "").strip()
            return f" năm {year}" if year else ""
        if target == "company_province_with_nationality":
            province = self._value_for_template_target("company_province", values)
            return f"{province}, Việt Nam" if province else ""
        if target == "legal_representative_province_with_nationality":
            province = self._value_for_template_target("legal_representative_province", values)
            return f"{province}, Việt Nam" if province else ""
        if target == "charter_capital_article_6_text":
            number = self._with_vnd(values.get("charter_capital_number", ""))
            words = str(values.get("charter_capital_text", "") or "").strip()
            if words and not re.search(r"\sđồng$", words, re.IGNORECASE):
                words = f"{words} đồng"
            if number and words:
                return f"{number} (Ghi bằng chữ: {words})"
            return number
        if target == "organization_owner_issue_place":
            return str(values.get(target, "") or "").split(" - ", 1)[0].strip(" ,;-")
        if target == "charter_capital_words_without_currency":
            return re.sub(r"\s*đồng\s*$", "", str(values.get("charter_capital_text", "") or ""), flags=re.IGNORECASE).strip()
        if target.startswith("organization_owner_representative_contact_address_"):
            address = str(values.get("organization_owner_representative_contact_address", "") or "").strip()
            suffix = target.removeprefix("organization_owner_representative_contact_address_")
            if suffix == "detail":
                return self._address_detail_only(address)
            if suffix == "ward":
                return self._administrative_part_only(address, {"xa", "phuong", "thi_tran"})
            if suffix == "province":
                return self._normalize_province(self._administrative_part_only(address, {"tinh", "thanh_pho"}))
        if target == "organization_owner_representative_id_number":
            identity = str(values.get("organization_owner_representative_identity", "") or "")
            match = re.search(r"(?:CCCD\s*số|Số\s*định\s*danh\s*cá\s*nhân)\s*:?\s*([0-9A-Za-z-]+)", identity, re.IGNORECASE)
            if match:
                return match.group(1)
            fallback = re.search(r"[0-9]{6,}", identity)
            return fallback.group(0) if fallback else ""
        if target == "legal_representative_full_name_title":
            return str(values.get("legal_representative_full_name", "") or "").title()
        if target in {"legal_representative_address_without_province", "beneficial_owner_address_without_province"}:
            source = "legal_representative_permanent_address" if target.startswith("legal_representative_") else "beneficial_owner_address"
            parts = [part.strip() for part in str(values.get(source, "") or "").split(",") if part.strip()]
            return ", ".join(parts[:-1]) if len(parts) > 1 else ""
        if target == "beneficial_owner_province":
            parts = [part.strip() for part in str(values.get("beneficial_owner_address", "") or "").split(",") if part.strip()]
            return parts[-1] if parts else ""
        if target == "tax_vat_method_x":
            return "X" if values.get("tax_vat_method") == "Khấu trừ" else ""
        if target == "tax_declaration_period_x":
            return "X" if values.get("tax_declaration_period") == "Hàng tháng" else ""
        if target == "document_date_only":
            day = str(values.get("document_day", "") or "").strip()
            month = str(values.get("document_month", "") or "").strip()
            year = str(values.get("document_year", "") or "").strip()
            if day and month and year:
                return f"ngày {day} tháng {month} năm {year}."
            return ""
        if target == "document_day_with_trailing_space":
            day = str(values.get("document_day", "") or "").strip()
            return f"{day} " if day else ""
        if target == "charter_capital_number_vnd":
            return self._with_vnd(values.get("charter_capital_number", ""))
        if target == "capital_total_ratio_number":
            return self._capital_total_ratio(values)
        if target == "member_capital_ratio_number":
            ratio = str(values.get("member_capital_ratio", "") or "").strip()
            return ratio.rstrip("% ").strip()
        raw = str(values.get(target, "") or "").strip()
        if target in self.DETAIL_ADDRESS_FIELDS:
            if target.startswith("company_"):
                return self._company_address_part(values, "detail")
            return self._address_detail_only(raw)
        if target in self.WARD_FIELDS:
            if target.startswith("company_"):
                return self._company_address_part(values, "ward")
            return self._administrative_part_only(raw, {"xa", "phuong", "thi_tran"})
        if target in self.PROVINCE_FIELDS:
            if target.startswith("company_"):
                return self._company_address_part(values, "province")
            return self._normalize_province(self._administrative_part_only(raw, {"tinh", "thanh_pho"}))
        if target in self.UPPERCASE_BOLD_FIELDS:
            return raw.upper()
        return raw

    def _date_component(self, value: object, component: int) -> str:
        parts = re.findall(r"\d+", str(value or ""))
        if len(parts) != 3:
            return ""
        if component in (0, 1):
            return parts[component].zfill(2)
        return parts[component]
    def _company_address_part(self, values: dict[str, str], part: str) -> str:
        full = str(values.get("company_address", "") or "").strip()
        detail = str(values.get("company_address_detail", "") or "").strip()
        ward = str(values.get("company_ward", "") or "").strip()
        province = str(values.get("company_province", "") or "").strip()
        has_admin_parts = bool(re.search(r"(?i)\b(xã|phường|thị trấn|đặc khu|tỉnh|thành phố)\b", full))
        source = full if has_admin_parts else detail
        if part == "detail":
            return self._address_detail_only(source or detail)
        if part == "ward":
            parsed = self._administrative_part_only(source, {"xa", "phuong", "thi_tran", "dac_khu"})
            return parsed if parsed != source or source else ward
        parsed = self._normalize_province(self._administrative_part_only(source, {"tinh", "thanh_pho"}))
        return parsed if parsed != source or source else province

    def _capital_total_ratio(self, values: dict[str, str]) -> str:
        ratio = str(values.get("capital_total_ratio", "") or "").strip()
        if ratio:
            return ratio.rstrip("% ").strip()
        member_count = self._member_count(values)
        if member_count <= 1:
            return "100"
        return "100"

    def _member_count(self, values: dict[str, str]) -> int:
        explicit = self._money_number(values.get("member_count", "")) or self._money_number(
            values.get("company_member_count", "")
        )
        if explicit:
            return explicit
        count = 0
        for index in range(1, 21):
            if str(values.get(f"member_{index}_full_name", "") or "").strip():
                count += 1
        return count or 1

    def _with_vnd(self, value: str) -> str:
        text = str(value or "").strip()
        if not text:
            return ""
        folded = self._fold(text)
        if "vnd" in folded or "dong" in folded:
            return text
        return f"{text} VNĐ"

    def _address_detail_only(self, value: str) -> str:
        parts = [part.strip(" ,.;") for part in re.split(r"[,;]", str(value or "")) if part.strip(" ,.;")]
        kept: list[str] = []
        blocked_prefixes = {"xa", "phuong", "thi_tran", "huyen", "quan", "thi_xa", "thanh_pho", "tinh"}
        for part in parts:
            normalized = self._normalize_key(part)
            if any(normalized.startswith(prefix + "_") or normalized == prefix for prefix in blocked_prefixes):
                continue
            kept.append(part)
        return ", ".join(kept) if kept else str(value or "").strip()

    def _administrative_part_only(self, value: str, allowed_prefixes: set[str]) -> str:
        text = str(value or "").strip()
        if not text:
            return ""
        parts = [part.strip(" ,.;") for part in re.split(r"[,;]", text) if part.strip(" ,.;")]
        for part in parts:
            normalized = self._normalize_key(part)
            if any(normalized.startswith(prefix + "_") or normalized == prefix for prefix in allowed_prefixes):
                return part
        return text

    def _normalize_province(self, value: str) -> str:
        text = str(value or "").strip()
        if not text:
            return ""
        normalized = self._normalize_key(text)
        for prefix in ("tinh_", "thanh_pho_"):
            if normalized.startswith(prefix):
                normalized = normalized[len(prefix) :]
                break
        return self.PROVINCE_AFTER_MERGER.get(normalized, text)

    def _replace_business_lines_tables_v2(
        self,
        document: Document,
        values: dict[str, str],
        *,
        keep_file1_business_line_numbers_on_one_line: bool = False,
    ) -> None:
        business_lines = self._business_lines(values)
        for table in self._iter_tables(document):
            if not table.rows:
                continue
            header_index = self._business_line_header_index(table)
            if header_index is None or len(table.rows[header_index].cells) != 4:
                continue
            widths = self._business_line_widths_v2(table)
            self._prepare_business_line_table_v2(table, header_index, widths)
            source_row = deepcopy(table.rows[header_index + 1]._tr) if len(table.rows) > header_index + 1 else None
            for row in list(table.rows[header_index + 1 :]):
                table._tbl.remove(row._tr)
            rows_to_write = business_lines or [{"name": "", "code": "", "is_main": False}]
            for index, line in enumerate(rows_to_write, start=1):
                row_xml = deepcopy(source_row) if source_row is not None else self._new_business_line_row_xml_v2(widths)
                self._strip_business_line_row_content_v2(row_xml)
                table._tbl.append(row_xml)
                self._write_business_line_row_v2(
                    table.rows[-1],
                    index,
                    line,
                    widths,
                    keep_stt_on_one_line=keep_file1_business_line_numbers_on_one_line,
                )

    def _business_line_widths_v2(self, table) -> list[int]:
        widths = [int(column.get(qn("w:w")) or 0) for column in table._tbl.tblGrid.gridCol_lst[:4]]
        if len(widths) != 4 or not all(widths):
            raise ValueError("Bảng ngành nghề phải có geometry hợp lệ cho 4 cột.")
        return widths

    def _prepare_business_line_table_v2(self, table, header_index: int, widths: list[int]) -> None:
        table.autofit = False
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.insert(0, tbl_w)
        tbl_w.set(qn("w:w"), str(sum(widths)))
        tbl_w.set(qn("w:type"), "dxa")
        tbl_layout = tbl_pr.find(qn("w:tblLayout"))
        if tbl_layout is None:
            tbl_layout = OxmlElement("w:tblLayout")
            tbl_pr.append(tbl_layout)
        tbl_layout.set(qn("w:type"), "fixed")
        for column, width in zip(table._tbl.tblGrid.gridCol_lst[:4], widths):
            column.set(qn("w:w"), str(width))

        header = table.rows[header_index]
        self._remove_row_height_v2(header._tr)
        self._unmark_repeating_table_header_v2(header)
        # The heading is legally prescribed in the supplied template. Only body
        # rows are regenerated, so its full wording and layout stay untouched.
        for cell, width in zip(header.cells[:4], widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is not None:
                tc_w.set(qn("w:w"), str(width))
                tc_w.set(qn("w:type"), "dxa")

    def _unmark_repeating_table_header_v2(self, row) -> None:
        """Keep the industry table header on the first page only.

        Word repeats a row on every page when ``w:tblHeader`` is present.
        The legal form's reference layout intentionally continues the table
        without repeating the column headings, so remove the flag explicitly
        even if it exists in the source template.
        """
        tr_pr = row._tr.find(qn("w:trPr"))
        if tr_pr is None:
            return
        for header_flag in tr_pr.findall(qn("w:tblHeader")):
            tr_pr.remove(header_flag)

    def _new_business_line_row_xml_v2(self, widths: list[int]):
        row = OxmlElement("w:tr")
        for width in widths:
            cell = OxmlElement("w:tc")
            tc_pr = OxmlElement("w:tcPr")
            tc_w = OxmlElement("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            tc_pr.append(tc_w)
            cell.append(tc_pr)
            row.append(cell)
        return row

    def _strip_business_line_row_content_v2(self, row_xml) -> None:
        self._remove_row_height_v2(row_xml)
        tr_pr = row_xml.find(qn("w:trPr"))
        if tr_pr is not None:
            for element in list(tr_pr):
                if element.tag in {qn("w:trHeight"), qn("w:cantSplit")}:
                    tr_pr.remove(element)
        # Keep the sample paragraphs in the cloned row. They carry the approved
        # left/right indents and line spacing from the Word template; the writer
        # below replaces only their text.

    def _remove_row_height_v2(self, row_xml) -> None:
        tr_pr = row_xml.find(qn("w:trPr"))
        if tr_pr is not None:
            for element in tr_pr.findall(qn("w:trHeight")):
                tr_pr.remove(element)

    def _write_business_line_row_v2(
        self,
        row,
        index: int,
        line: dict[str, str | bool],
        widths: list[int],
        *,
        keep_stt_on_one_line: bool = False,
    ) -> None:
        self._remove_row_height_v2(row._tr)
        self._write_business_line_cell_v2(
            row.cells[0],
            str(index),
            widths[0],
            "center",
            no_wrap=keep_stt_on_one_line,
        )
        self._write_business_line_cell_v2(row.cells[1], str(line["name"]), widths[1], "left")
        self._write_business_line_cell_v2(row.cells[2], str(line["code"]), widths[2], "center", bold=bool(line["is_main"]))
        self._write_business_line_cell_v2(row.cells[3], "X" if line["is_main"] else "", widths[3], "center", bold=bool(line["is_main"]))

    def _write_business_line_cell_v2(
        self,
        cell,
        value: str,
        width: int,
        alignment: str,
        bold: bool = False,
        *,
        no_wrap: bool = False,
    ) -> None:
        """Replace text while retaining the paragraph spacing from the template row."""
        tc = cell._tc
        tc_pr = tc.get_or_add_tcPr()
        tc_w = tc_pr.find(qn("w:tcW"))
        if tc_w is None:
            tc_w = OxmlElement("w:tcW")
            tc_pr.insert(0, tc_w)
        tc_w.set(qn("w:w"), str(width))
        tc_w.set(qn("w:type"), "dxa")
        v_align = tc_pr.find(qn("w:vAlign"))
        if v_align is None:
            v_align = OxmlElement("w:vAlign")
            tc_pr.append(v_align)
        v_align.set(qn("w:val"), "center")
        if no_wrap:
            no_wrap_element = tc_pr.find(qn("w:noWrap"))
            if no_wrap_element is None:
                tc_pr.append(OxmlElement("w:noWrap"))

        source_paragraph = next((child for child in tc if child.tag == qn("w:p")), None)
        source_p_pr = deepcopy(source_paragraph.find(qn("w:pPr"))) if source_paragraph is not None else None
        source_r_pr = None
        if source_paragraph is not None:
            source_run = source_paragraph.find(qn("w:r"))
            if source_run is not None:
                source_r_pr = deepcopy(source_run.find(qn("w:rPr")))
        for child in list(tc):
            if child.tag != qn("w:tcPr"):
                tc.remove(child)

        paragraph = OxmlElement("w:p")
        p_pr = source_p_pr if source_p_pr is not None else OxmlElement("w:pPr")
        for justification in p_pr.findall(qn("w:jc")):
            p_pr.remove(justification)
        if alignment == "left" or no_wrap:
            # Some legacy file-1 sample rows carry an indentation intended for
            # a single digit. It reduces the usable width enough for 10, 11,
            # ... to split vertically. The STT cell owns its full width.
            for indentation in p_pr.findall(qn("w:ind")):
                p_pr.remove(indentation)
            for tab_stops in p_pr.findall(qn("w:tabs")):
                p_pr.remove(tab_stops)
        justification = OxmlElement("w:jc")
        justification.set(qn("w:val"), alignment)
        p_pr.append(justification)
        paragraph.append(p_pr)
        for line_index, line in enumerate(str(value or "").splitlines() or [""]):
            if line_index:
                break_run = OxmlElement("w:r")
                break_run.append(OxmlElement("w:br"))
                paragraph.append(break_run)
            run = OxmlElement("w:r")
            r_pr = deepcopy(source_r_pr) if source_r_pr is not None else OxmlElement("w:rPr")
            for element in list(r_pr.findall(qn("w:b"))):
                r_pr.remove(element)
            if bold:
                r_pr.append(OxmlElement("w:b"))

            # The registration-industry table is a formal Word table: keep all rows
            # in one font/size regardless of the exemplar row used by the template.
            fonts = r_pr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                r_pr.insert(0, fonts)
            for font_name in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                fonts.set(qn(font_name), "Times New Roman")
            for tag in ("w:sz", "w:szCs"):
                for element in list(r_pr.findall(qn(tag))):
                    r_pr.remove(element)
                size = OxmlElement(tag)
                size.set(qn("w:val"), "24")
                r_pr.append(size)
            run.append(r_pr)
            text_node = OxmlElement("w:t")
            text_node.text = line
            if line[:1].isspace() or line[-1:].isspace():
                text_node.set(qn("xml:space"), "preserve")
            run.append(text_node)
            paragraph.append(run)
        tc.append(paragraph)

    def _replace_business_lines_tables(self, document: Document, values: dict[str, str]) -> None:
        business_lines = self._business_lines(values)

        for table in self._iter_tables(document):
            if not table.rows:
                continue
            header_index = self._business_line_header_index(table)
            if header_index is None:
                continue
            if len(table.rows[header_index].cells) < 4:
                continue

            column_widths = self._normalize_business_line_table(table, header_index)
            template_index = min(header_index + 1, len(table.rows) - 1)
            template_row = deepcopy(table.rows[template_index]._tr)
            for row in list(table.rows[header_index + 1 :]):
                table._tbl.remove(row._tr)

            rows_to_write = business_lines or [{"name": "", "code": "", "is_main": False}]
            for index, line in enumerate(rows_to_write, start=1):
                table._tbl.append(deepcopy(template_row))
                row = table.rows[-1]
                row.height = None
                self._set_business_line_row_widths(row, column_widths)
                self._set_business_line_cell(row.cells[0], str(index), WD_ALIGN_PARAGRAPH.CENTER)
                self._set_business_line_cell(row.cells[1], str(line["name"]), WD_ALIGN_PARAGRAPH.LEFT)
                self._set_business_line_cell(
                    row.cells[2], str(line["code"]), WD_ALIGN_PARAGRAPH.CENTER, bold=bool(line["is_main"])
                )
                self._set_business_line_cell(
                    row.cells[3], "X" if line["is_main"] else "", WD_ALIGN_PARAGRAPH.CENTER, bold=bool(line["is_main"])
                )

    def _normalize_business_line_table(self, table, header_index: int) -> list[int]:
        grid_columns = table._tbl.tblGrid.gridCol_lst
        widths = [int(column.get(qn("w:w")) or 0) for column in grid_columns[:4]]
        if len(widths) != 4 or not all(widths):
            widths = [int(cell.width or 0) for cell in table.rows[header_index].cells[:4]]
        if len(widths) != 4 or not all(widths):
            widths = [900, 5100, 1500, 1900]

        table.autofit = False
        table_width = sum(widths)
        tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            table._tbl.tblPr.insert(0, tbl_w)
        tbl_w.set(qn("w:w"), str(table_width))
        tbl_w.set(qn("w:type"), "dxa")

        for column, width in zip(grid_columns[:4], widths):
            column.set(qn("w:w"), str(width))

        header = table.rows[header_index]
        self._mark_repeating_table_header(header)
        header_texts = ["STT", "Tên ngành", "Mã ngành", "Ngành, nghề\nkinh doanh\nchính"]
        for index, cell in enumerate(header.cells[:4]):
            self._set_business_line_cell(cell, header_texts[index], WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        self._set_business_line_row_widths(header, widths)

        for row in table.rows[header_index + 1 :]:
            row.height = None
            self._set_business_line_row_widths(row, widths)
        return widths

    def _set_business_line_row_widths(self, row, widths: list[int]) -> None:
        for cell, width in zip(row.cells[:4], widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")

    def _mark_repeating_table_header(self, row) -> None:
        # Industry tables must continue without repeating their headings on
        # later pages. Keep this legacy helper defensive for older callers.
        tr_pr = row._tr.get_or_add_trPr()
        for header_flag in tr_pr.findall(qn("w:tblHeader")):
            tr_pr.remove(header_flag)

    def _iter_tables(self, container):
        for table in container.tables:
            yield table
            for row in table.rows:
                for cell in row.cells:
                    yield from self._iter_tables(cell)

    def _business_line_header_index(self, table) -> int | None:
        for index, row in enumerate(table.rows[:3]):
            folded_header = self._fold(" ".join(cell.text for cell in row.cells))
            compact_header = folded_header.replace(" ", "")
            has_name = "ten nganh" in folded_header or "tennganh" in compact_header
            has_code = "ma nganh" in folded_header or "manganh" in compact_header
            if has_name and has_code:
                return index
        return None

    def _replace_asset_contribution_tables(self, document: Document, values: dict[str, str]) -> None:
        for table in self._iter_tables(document):
            if not table.rows:
                continue
            header_index = self._asset_header_index(table)
            if header_index is None or len(table.rows[header_index].cells) < 4:
                continue
            asset_rows, total_row = self._asset_table_rows(table, header_index)
            if not asset_rows:
                continue
            assets = self._asset_template_contributions(values, [row.cells[1].text for row in asset_rows])
            for index, row in enumerate(asset_rows, start=1):
                asset = assets[index - 1]
                self._set_cell_text(row.cells[0], str(index))
                self._set_cell_text(row.cells[1], asset["name"])
                self._set_cell_text(row.cells[2], asset["amount"])
                self._set_cell_text(row.cells[3], asset["ratio"])
            if total_row is not None:
                total_label_left = " ".join(total_row.cells[0].text.split()) or "Tổng số"
                total_label_right = " ".join(total_row.cells[1].text.split()) or "Tổng số"
                self._set_cell_text(total_row.cells[0], total_label_left)
                self._set_cell_text(total_row.cells[1], total_label_right)
                self._set_cell_text(total_row.cells[2], self._asset_total_amount_text(values, assets))
                self._set_cell_text(total_row.cells[3], self._asset_total_ratio_text(values, assets))

    def _normalize_file1_financial_tables(self, document: Document, template_name: str) -> None:
        """Use one stable table system for sections 7 and 8 of file 1."""
        if not template_name.startswith("1. "):
            return

        # Sections 7 and 8 are legal-form tables. Their original column widths,
        # cell margins and page-flow properties are part of the approved layout,
        # so only cell values are replaced by _set_cell_text().
        return

    def _normalize_financial_table_geometry(self, table, widths: list[int]) -> None:
        if len(table.rows[0].cells) != len(widths):
            return
        table.autofit = False
        table_properties = table._tbl.tblPr
        table_width = table_properties.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            table_properties.insert(0, table_width)
        table_width.set(qn("w:w"), str(sum(widths)))
        table_width.set(qn("w:type"), "dxa")

        table_indent = table_properties.find(qn("w:tblInd"))
        if table_indent is not None:
            table_indent.set(qn("w:w"), "0")
            table_indent.set(qn("w:type"), "dxa")
        table_layout = table_properties.find(qn("w:tblLayout"))
        if table_layout is None:
            table_layout = OxmlElement("w:tblLayout")
            table_properties.append(table_layout)
        table_layout.set(qn("w:type"), "fixed")

        grid_columns = table._tbl.tblGrid.gridCol_lst
        if len(grid_columns) != len(widths):
            return
        for column, width in zip(grid_columns, widths):
            column.set(qn("w:w"), str(width))
        for row in table.rows:
            self._remove_row_height_v2(row._tr)
            for cell, width in zip(row.cells, widths):
                cell_properties = cell._tc.get_or_add_tcPr()
                cell_width = cell_properties.find(qn("w:tcW"))
                if cell_width is None:
                    cell_width = OxmlElement("w:tcW")
                    cell_properties.insert(0, cell_width)
                cell_width.set(qn("w:w"), str(width))
                cell_width.set(qn("w:type"), "dxa")

    def _normalize_financial_table_cells(self, table, header_index: int, name_column: int) -> None:
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                self._remove_vertical_cell_padding(cell)
                alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == header_index or column_index != name_column else WD_ALIGN_PARAGRAPH.LEFT
                for paragraph in cell.paragraphs:
                    paragraph.alignment = alignment
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1
                    paragraph.paragraph_format.left_indent = Inches(0)
                    paragraph.paragraph_format.right_indent = Inches(0)
                    paragraph.paragraph_format.first_line_indent = Inches(0)
                    self._set_compact_table_paragraph_spacing(paragraph)

    def _remove_vertical_cell_padding(self, cell) -> None:
        """Match the reference form's tight table cells.

        Word applies a default top/bottom cell margin when the template does
        not define one. That margin remains visible even when paragraph
        before/after spacing is zero, which makes generated tables look
        vertically loose compared with the reference form.
        """
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.find(qn("w:tcMar"))
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for side in ("top", "bottom"):
            margin = tc_mar.find(qn(f"w:{side}"))
            if margin is None:
                margin = OxmlElement(f"w:{side}")
                tc_mar.append(margin)
            margin.set(qn("w:w"), "0")
            margin.set(qn("w:type"), "dxa")

    def _set_compact_table_paragraph_spacing(self, paragraph) -> None:
        """Write deterministic single-line spacing into the paragraph XML."""
        p_pr = paragraph._p.get_or_add_pPr()
        spacing = p_pr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            p_pr.append(spacing)
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "240")
        spacing.set(qn("w:lineRule"), "auto")

    def _asset_header_index(self, table) -> int | None:
        for index, row in enumerate(table.rows[:3]):
            folded_header = self._fold(" ".join(cell.text for cell in row.cells))
            compact_header = folded_header.replace(" ", "")
            has_asset = "tai san gop von" in folded_header or "taisangopvon" in compact_header
            has_ratio = "ty le" in folded_header or "tyle" in compact_header
            if has_asset and has_ratio:
                return index
        return None

    def _asset_contributions(self, values: dict[str, str]) -> list[dict[str, str]]:
        total = self._money_number(values.get("charter_capital_number", ""))
        assets: list[dict[str, str]] = []
        for index in range(1, 11):
            name = str(values.get(f"asset_contribution_{index}_name", "") or "").strip()
            amount = str(values.get(f"asset_contribution_{index}_amount", "") or "").strip()
            if not name and not amount:
                continue
            ratio = str(values.get(f"asset_contribution_{index}_ratio", "") or "").strip()
            if not ratio and total:
                amount_number = self._money_number(amount)
                if amount_number:
                    ratio = f"{round(amount_number * 100 / total, 2):g}"
            assets.append({"name": name or f"Tài sản {index}", "amount": self._with_vnd(amount), "ratio": ratio.rstrip("% ")})
        if assets:
            return assets
        return [
            {
                "name": "Đồng Việt Nam",
                "amount": self._with_vnd(values.get("charter_capital_number", "")),
                "ratio": "100",
            }
        ]

    def _asset_table_rows(self, table, header_index: int):
        body_rows = list(table.rows[header_index + 1 :])
        if not body_rows:
            return [], None
        asset_rows = []
        total_row = None
        for row in body_rows:
            folded = self._fold(" ".join(cell.text for cell in row.cells[:2]))
            if "tong so" in folded:
                total_row = row
                continue
            asset_rows.append(row)
        return asset_rows, total_row

    def _asset_template_contributions(self, values: dict[str, str], default_names: list[str]) -> list[dict[str, str]]:
        total = self._money_number(values.get("charter_capital_number", ""))
        manual_assets = self._has_manual_asset_contributions(values)
        assets: list[dict[str, str]] = []
        for index, default_name in enumerate(default_names, start=1):
            name = str(values.get(f"asset_contribution_{index}_name", "") or "").strip()
            amount = str(values.get(f"asset_contribution_{index}_amount", "") or "").strip()
            ratio = str(values.get(f"asset_contribution_{index}_ratio", "") or "").strip()
            if not name:
                name = " ".join(str(default_name or "").split())
            if not amount:
                amount = str(values.get("charter_capital_number", "") or "").strip() if index == 1 and total and not manual_assets else "0"
            amount_number = self._money_number(amount)
            if not ratio and total:
                if amount_number:
                    ratio = f"{round(amount_number * 100 / total, 2):g}"
                else:
                    ratio = "0"
            elif not ratio:
                ratio = "100" if index == 1 and amount_number else "0"
            assets.append(
                {
                    "name": name or f"Tai san {index}",
                    "amount": self._asset_amount_text(amount),
                    "ratio": ratio.rstrip("% "),
                }
            )
        return assets

    def _has_manual_asset_contributions(self, values: dict[str, str]) -> bool:
        # Zero is the required display value for an omitted asset, not an actual
        # contribution that should suppress the automatic VND allocation.
        for index in range(1, 11):
            if self._money_number(values.get(f"asset_contribution_{index}_amount", "")):
                return True
        return False

    def _asset_amount_text(self, value: str) -> str:
        text = str(value or "").strip()
        if not text:
            return "0"
        text = re.sub(r"\s*(VNĐ|VND|đồng)\s*$", "", text, flags=re.IGNORECASE).strip()
        return text or "0"

    def _asset_total_amount_text(self, values: dict[str, str], assets: list[dict[str, str]]) -> str:
        total = str(values.get("charter_capital_number", "") or "").strip()
        if total:
            return self._asset_amount_text(total)
        summed = sum(self._money_number(asset["amount"]) for asset in assets)
        return self._asset_amount_text(str(summed))

    def _asset_total_ratio_text(self, values: dict[str, str], assets: list[dict[str, str]]) -> str:
        ratio = str(values.get("capital_total_ratio", "") or "").strip().rstrip("% ")
        if ratio:
            return ratio
        summed = sum(float(asset["ratio"] or "0") for asset in assets if str(asset["ratio"]).strip())
        if summed:
            return f"{summed:g}"
        return "100"

    def _money_number(self, value: str) -> int:
        digits = re.sub(r"\D", "", str(value or ""))
        return int(digits) if digits else 0

    @staticmethod
    def _business_line_indices(values: dict[str, str]) -> list[int]:
        indices = {
            int(match.group(1))
            for key in values
            if (match := re.match(r"^business_line_(\d+)_(?:name|detail|code|is_main)$", key))
        }
        return sorted(index for index in indices if index > 0)

    def _business_lines(self, values: dict[str, str]) -> list[dict[str, str | bool]]:
        lines: list[dict[str, str | bool]] = []
        for index in self._business_line_indices(values):
            name = self._normalize_business_line_text(values.get(f"business_line_{index}_name", ""))
            detail = self._normalize_business_line_text(values.get(f"business_line_{index}_detail", ""))
            if detail:
                name = f"{name}\n{detail}" if name else detail
            code = re.sub(r"\D", "", values.get(f"business_line_{index}_code", ""))[:4]
            is_main = bool(values.get(f"business_line_{index}_is_main", "").strip())
            if name or code or is_main:
                lines.append({"name": name, "code": code, "is_main": is_main})
        return lines

    def _normalize_business_line_text(self, value: str) -> str:
        text = "\n".join(part.strip() for part in value.replace("\r", "\n").split("\n") if part.strip())
        text = re.sub(r"\s+(Chi\s*tiết\s*:)", r"\n\1", text, flags=re.IGNORECASE)
        text = re.sub(r"(?i)chi\s*tiết\s*:", "Chi tiết:", text)
        return text

    def _set_cell_text(self, cell, value: str, bold: bool = False) -> None:
        """Set cell content without discarding template paragraph margins/styles."""
        tc = cell._tc
        source_paragraph = next((child for child in tc if child.tag == qn("w:p")), None)
        source_p_pr = deepcopy(source_paragraph.find(qn("w:pPr"))) if source_paragraph is not None else None
        source_r_pr = None
        if source_paragraph is not None:
            source_run = source_paragraph.find(qn("w:r"))
            if source_run is not None:
                source_r_pr = deepcopy(source_run.find(qn("w:rPr")))
        for child in list(tc):
            if child.tag != qn("w:tcPr"):
                tc.remove(child)
        paragraph = OxmlElement("w:p")
        paragraph.append(source_p_pr if source_p_pr is not None else OxmlElement("w:pPr"))
        for line_index, line in enumerate(str(value or "").splitlines() or [""]):
            if line_index:
                break_run = OxmlElement("w:r")
                break_run.append(OxmlElement("w:br"))
                paragraph.append(break_run)
            run = OxmlElement("w:r")
            r_pr = deepcopy(source_r_pr) if source_r_pr is not None else OxmlElement("w:rPr")
            for element in list(r_pr.findall(qn("w:b"))):
                r_pr.remove(element)
            if bold:
                r_pr.append(OxmlElement("w:b"))

            # The registration-industry table is a formal Word table: keep all rows
            # in one font/size regardless of the exemplar row used by the template.
            fonts = r_pr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                r_pr.insert(0, fonts)
            for font_name in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                fonts.set(qn(font_name), "Times New Roman")
            for tag in ("w:sz", "w:szCs"):
                for element in list(r_pr.findall(qn(tag))):
                    r_pr.remove(element)
                size = OxmlElement(tag)
                size.set(qn("w:val"), "24")
                r_pr.append(size)
            run.append(r_pr)
            text_node = OxmlElement("w:t")
            text_node.text = line
            if line[:1].isspace() or line[-1:].isspace():
                text_node.set(qn("xml:space"), "preserve")
            run.append(text_node)
            paragraph.append(run)
        tc.append(paragraph)

    def _set_business_line_cell(self, cell, value: str, alignment, bold: bool = False) -> None:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = alignment
        paragraph.paragraph_format.left_indent = None
        paragraph.paragraph_format.right_indent = None
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        lines = str(value or "").splitlines() or [""]
        for line_index, line in enumerate(lines):
            if line_index:
                paragraph.add_run().add_break()
            run = paragraph.add_run(line)
            run.bold = bold
            run.font.size = Pt(13)

    def _fix_contact_lines(self, document: Document, values: dict[str, str]) -> None:
        phone = str(values.get("authorized_person_phone", "") or "").strip()
        email = str(values.get("authorized_person_email", "") or "").strip()
        if not phone and not email:
            return
        for paragraph in self._iter_paragraphs(document):
            folded = self._fold(paragraph.text)
            if "dien thoai" not in folded or "email" not in folded:
                continue
            self._replace_paragraph_text(paragraph, f"Điện thoại: {phone}\tEmail: {email}".rstrip())
            try:
                paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(3.8))
            except Exception:
                pass

    def _fix_registration_office_lines(self, document: Document, values: dict[str, str]) -> None:
        office = self._registration_office_text(values)
        if not office:
            return
        for paragraph in self._iter_paragraphs(document):
            folded = self._fold(paragraph.text)
            if "kinh gui" not in folded or "phong dang ky kinh doanh" not in folded:
                continue
            self._replace_paragraph_text(paragraph, f"Kính gửi: {office}")
            for run in paragraph.runs:
                run.bold = False

    def _registration_office_text(self, values: dict[str, str]) -> str:
        office = str(values.get("business_registration_office", "") or "").strip()
        if not office:
            return ""
        return re.sub(r"^\s*Kính gửi\s*:\s*", "", office, flags=re.IGNORECASE).strip()

    def _apply_bold_uppercase_formatting(self, document: Document, values: dict[str, str]) -> None:
        terms = self._bold_terms(values)
        if not terms:
            return
        for paragraph in self._iter_paragraphs(document):
            self._bold_terms_in_paragraph(paragraph, terms)

    def _fix_signature_font_sizes(self, document: Document, template_name: str, values: dict[str, str]) -> None:
        if not template_name.startswith("3. "):
            return
        signer = str(values.get("legal_representative_full_name", "") or "").strip().upper()
        if not signer:
            return
        folded_signer = self._fold(signer)
        for table in self._iter_tables(document):
            for row in table.rows:
                for cell in row.cells:
                    folded_cell = self._fold(cell.text)
                    if folded_signer not in folded_cell:
                        continue
                    if "nguoi dai dien" not in folded_cell and "chu tich" not in folded_cell:
                        continue
                    for paragraph in cell.paragraphs:
                        if folded_signer not in self._fold(paragraph.text):
                            continue
                        for run in paragraph.runs:
                            if run.text:
                                run.font.size = Pt(13)

    def _normalize_power_of_attorney_layout(
        self, document: Document, template_name: str, values: dict[str, str]
    ) -> None:
        if not template_name.startswith("4. "):
            return

        for paragraph in self._iter_paragraphs(document):
            if not paragraph.text.strip():
                continue
            paragraph.paragraph_format.keep_together = False
            paragraph.paragraph_format.keep_with_next = False
            for run in paragraph.runs:
                if run.text and (run.font.size is None or run.font.size.pt > 13):
                    run.font.size = Pt(12.5)

        body = document._body._element
        table_seen = False
        children = list(body)
        for index, child in enumerate(children):
            if child.tag.endswith("}tbl"):
                table_seen = True
                continue
            next_child = children[index + 1] if index + 1 < len(children) else None
            if (
                child.tag.endswith("}p")
                and next_child is not None
                and next_child.tag.endswith("}tbl")
                and not "".join(child.itertext()).strip()
            ):
                body.remove(child)
                continue
            if table_seen and child.tag.endswith("}p") and not "".join(child.itertext()).strip():
                body.remove(child)

        signer = str(values.get("legal_representative_full_name", "") or "").strip().upper()
        if not signer:
            return
        for table in self._iter_tables(document):
            if not table.rows:
                continue
            for cell in table.rows[-1].cells:
                if "ben uy quyen" not in self._fold(cell.text):
                    continue
                cell.text = ""
                entries = [("BÊN UỶ QUYỀN", True), ("(ký, ghi rõ họ tên)", False), ("", False), (signer, True)]
                for index, (text, bold) in enumerate(entries):
                    paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1
                    run = paragraph.add_run(text)
                    run.bold = bold
                    run.font.size = Pt(12.5)
                table.rows[-1].height = None
                return

    def _bold_terms(self, values: dict[str, str]) -> list[str]:
        terms: set[str] = set()
        for key in self.UPPERCASE_BOLD_FIELDS:
            value = str(values.get(key, "") or "").strip()
            if value:
                terms.add(value.upper())
        return sorted(terms, key=len, reverse=True)

    def _bold_terms_in_paragraph(self, paragraph, terms: list[str]) -> None:
        text = paragraph.text
        if not text:
            return
        folded_text = self._fold(text)
        if not any(self._fold(term) in folded_text for term in terms):
            return

        matches: list[tuple[int, int]] = []
        occupied = [False] * len(text)
        for term in terms:
            if not term:
                continue
            start = 0
            lowered_text = text.lower()
            lowered_term = term.lower()
            while True:
                index = lowered_text.find(lowered_term, start)
                if index == -1:
                    break
                end = index + len(term)
                if end <= len(text) and not any(occupied[index:end]):
                    matches.append((index, end))
                    for offset in range(index, end):
                        occupied[offset] = True
                start = index + 1
        if not matches:
            return
        matches.sort()

        self._clear_paragraph_runs(paragraph)
        cursor = 0
        for start, end in matches:
            if start > cursor:
                paragraph.add_run(text[cursor:start])
            bold_run = paragraph.add_run(text[start:end].upper())
            bold_run.bold = True
            cursor = end
        if cursor < len(text):
            paragraph.add_run(text[cursor:])

    def _iter_paragraphs(self, container):
        for paragraph in getattr(container, "paragraphs", []):
            yield paragraph
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    yield from self._iter_paragraphs(cell)
        for section in getattr(container, "sections", []):
            yield from self._iter_paragraphs(section.header)
            yield from self._iter_paragraphs(section.footer)

    def _replace_paragraph_text(self, paragraph, text: str) -> None:
        # A template paragraph can hold text in hyperlinks as well as runs.
        # Remove every inline child so sample text cannot survive replacement.
        p_element = paragraph._p
        for child in list(p_element):
            if child.tag != qn("w:pPr"):
                p_element.remove(child)
        paragraph.add_run(text)

    def _clear_paragraph_runs(self, paragraph) -> None:
        for run in paragraph.runs:
            run.text = ""

    def _replace_in_container(self, container, values: dict[str, str]) -> None:
        element = getattr(container, "_element", None)
        if element is None:
            element = getattr(container, "element", None)
        if element is None:
            return
        for paragraph_element in element.xpath(".//w:p"):
            self._replace_in_paragraph(Paragraph(paragraph_element, container), values)

    def _replace_in_paragraph(self, paragraph, values: dict[str, str]) -> None:
        original = paragraph.text
        if "[" not in original or "]" not in original:
            return
        replaced = self._replace_text(original, values)
        if replaced == original:
            return

        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = replaced
        else:
            paragraph.add_run(replaced)

    def _replace_split_placeholders_in_package(self, docx_path: Path, values: dict[str, str]) -> None:
        tmp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx", dir=docx_path.parent)
        tmp_path = Path(tmp_file.name)
        tmp_file.close()

        try:
            with zipfile.ZipFile(docx_path, "r") as source_zip:
                with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as target_zip:
                    for item in source_zip.infolist():
                        data = source_zip.read(item.filename)
                        if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                            data = self._replace_split_placeholders_in_xml(data, values, docx_path.name, item.filename)
                        target_zip.writestr(item, data)
            tmp_path.replace(docx_path)
        except Exception:
            tmp_path.unlink(missing_ok=True)
            raise

    def _replace_bookmarked_fields_in_package(self, docx_path: Path, values: dict[str, str]) -> None:
        tmp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx", dir=docx_path.parent)
        tmp_path = Path(tmp_file.name)
        tmp_file.close()

        changed = False
        try:
            with zipfile.ZipFile(docx_path, "r") as source_zip:
                with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as target_zip:
                    for item in source_zip.infolist():
                        data = source_zip.read(item.filename)
                        if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                            replaced = self._replace_bookmarked_fields_in_xml(
                                data,
                                values,
                                docx_path.name,
                                item.filename,
                            )
                            if replaced != data:
                                changed = True
                            data = replaced
                        target_zip.writestr(item, data)
            if changed:
                tmp_path.replace(docx_path)
            else:
                tmp_path.unlink(missing_ok=True)
        except Exception:
            tmp_path.unlink(missing_ok=True)
            raise

    def _replace_bookmarked_fields_in_xml(
        self,
        data: bytes,
        values: dict[str, str],
        template_name: str,
        part_name: str,
    ) -> bytes:
        if b"MINGO_FIELD_" not in data:
            return data

        root = etree.fromstring(data)
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        w_namespace = namespace["w"]
        part_manifest = self._template_part_manifest(template_name, part_name)
        manifest_fields = {
            str(index): field
            for index, field in dict(part_manifest.get("placeholder_fields", {})).items()
        }
        if not manifest_fields:
            return data

        changed = False
        for bookmark in root.xpath(".//w:bookmarkStart[starts-with(@w:name, 'MINGO_FIELD_')]", namespaces=namespace):
            name = str(bookmark.get(f"{{{w_namespace}}}name") or "")
            match = re.fullmatch(r"MINGO_FIELD_(\d+)", name)
            if not match:
                continue
            target = manifest_fields.get(match.group(1))
            if target is None:
                continue
            end_id = bookmark.get(f"{{{w_namespace}}}id")
            paragraph = bookmark.getparent()
            while paragraph is not None and paragraph.tag != f"{{{w_namespace}}}p":
                paragraph = paragraph.getparent()
            if paragraph is None:
                continue

            children = list(paragraph)
            try:
                start_index = children.index(bookmark)
            except ValueError:
                continue

            end_index = None
            for index in range(start_index + 1, len(children)):
                child = children[index]
                if child.tag == f"{{{w_namespace}}}bookmarkEnd" and child.get(f"{{{w_namespace}}}id") == end_id:
                    end_index = index
                    break
            if end_index is None:
                continue

            prefix = "".join(
                node.text or ""
                for child in children[:start_index]
                for node in child.xpath(".//w:t", namespaces=namespace)
            )
            suffix = "".join(
                node.text or ""
                for child in children[end_index + 1 :]
                for node in child.xpath(".//w:t", namespaces=namespace)
            )
            replacement = self._replacement_with_boundary_spacing(
                self._value_for_template_target(str(target), values),
                prefix,
                suffix,
            )

            text_node = None
            for child in children[start_index + 1 : end_index]:
                for node in child.xpath(".//w:t", namespaces=namespace):
                    if text_node is None:
                        text_node = node
                    else:
                        node.text = ""
            if text_node is None:
                run = etree.Element(f"{{{w_namespace}}}r")
                text_node = etree.SubElement(run, f"{{{w_namespace}}}t")
                paragraph.insert(start_index + 1, run)
            text_node.text = replacement
            changed = True

        if not changed:
            return data
        return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)

    def _replace_split_placeholders_in_xml(
        self,
        data: bytes,
        values: dict[str, str],
        template_name: str = "",
        part_name: str = "",
    ) -> bytes:
        if b"[" not in data or b"]" not in data:
            return data

        root = etree.fromstring(data)
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        part_manifest = self._template_part_manifest(template_name, part_name)
        manifest_fields = {
            str(index): field
            for index, field in dict(part_manifest.get("placeholder_fields", {})).items()
        }
        placeholder_index = 0

        # A placeholder can be split over several runs.  Process one paragraph
        # at a time so text from adjacent fields or table rows can never be
        # joined into a fake placeholder.
        paragraphs = root.xpath(".//w:p", namespaces=namespace)
        for paragraph_index, paragraph in enumerate(paragraphs):
            text_nodes = paragraph.xpath(".//w:t", namespaces=namespace)
            if not text_nodes:
                continue

            segments = []
            chunks = []
            cursor = 0
            for node in text_nodes:
                text = node.text or ""
                start = cursor
                cursor += len(text)
                segments.append((start, cursor, node, text))
                chunks.append(text)

            joined = "".join(chunks)
            matches = list(re.finditer(r"\[([^\[\]]{1,300})\]", joined))
            if not matches:
                continue

            context_parts = [joined]
            context_parts.extend(
                "".join(node.text or "" for node in candidate.xpath(".//w:t", namespaces=namespace))
                for candidate in paragraphs[max(0, paragraph_index - 4) : paragraph_index]
            )
            context_parts.extend(
                "".join(node.text or "" for node in candidate.xpath(".//w:t", namespaces=namespace))
                for candidate in paragraphs[paragraph_index + 1 : paragraph_index + 3]
            )
            table = paragraph.xpath("ancestor::w:tbl[1]", namespaces=namespace)
            if table:
                context_parts.append("".join(node.text or "" for node in table[0].xpath(".//w:t", namespaces=namespace)))
            context = " ".join(context_parts)

            replacements: list[tuple[re.Match[str], str | None]] = []
            for match in matches:
                placeholder = match.group(1).strip()
                manifest_target = manifest_fields.get(str(placeholder_index))
                # Explicit field placeholders added to a template must win over
                # the legacy sequential manifest index.  The latter describes
                # bookmark fields, not the order of visible [field] tokens.
                if placeholder in {"company_english_name", "company_short_name"}:
                    replacement = self._value_for_template_target(placeholder, values)
                elif manifest_target is not None:
                    replacement = self._value_for_template_target(str(manifest_target), values)
                else:
                    replacement = self._value_for_placeholder(placeholder, values, context)
                replacements.append((match, replacement))
                placeholder_index += 1

            if not any(replacement is not None for _, replacement in replacements):
                continue

            pieces: list[str] = []
            last_end = 0
            for match, replacement in replacements:
                pieces.append(joined[last_end : match.start()])
                if replacement is None:
                    pieces.append(match.group(0))
                else:
                    prefix = "".join(pieces)
                    suffix = joined[match.end() :]
                    replacement = self._replacement_with_boundary_spacing(replacement, prefix, suffix)
                    pieces.append(replacement)
                last_end = match.end()
            pieces.append(joined[last_end:])
            replaced = "".join(pieces)

            text_nodes[0].text = replaced
            for node in text_nodes[1:]:
                node.text = ""

        return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)

    def _replacement_with_boundary_spacing(self, replacement: str, prefix: str, suffix: str) -> str:
        text = replacement or ""
        if prefix and text and not prefix[-1].isspace() and prefix[-1] not in "([{":
            text = " " + text
        if suffix and text and not suffix[0].isspace() and suffix[0] not in ".,;:)]}":
            text = text + " "
        return text

    def _extract_placeholders_from_xml(self, data: bytes) -> list[str]:
        try:
            root = etree.fromstring(data)
        except etree.XMLSyntaxError:
            return []
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        placeholders: list[str] = []
        for paragraph in root.xpath(".//w:p", namespaces=namespace):
            text_nodes = paragraph.xpath(".//w:t", namespaces=namespace)
            joined = "".join(node.text or "" for node in text_nodes)
            placeholders.extend(match.group(1).strip() for match in re.finditer(r"\[([^\[\]]{1,300})\]", joined))
        return placeholders

    def _replace_text(self, text: str, values: dict[str, str]) -> str:
        def replacement(match: re.Match[str]) -> str:
            placeholder = match.group(0)
            key_or_sample = match.group(1).strip()
            value = self._value_for_placeholder(key_or_sample, values, text)
            return placeholder if value is None else value

        return re.sub(r"\[([^\[\]]+)\]", replacement, text)

    def _value_for_placeholder(self, placeholder: str, values: dict[str, str], context: str = "") -> str | None:
        normalized = self._normalize_key(placeholder)
        known_direct_keys = {
            "business_registration_office",
            "company_name",
            "company_english_name",
            "company_short_name",
            "company_address",
            "company_address_detail",
            "company_ward",
            "company_province",
            "company_phone",
            "company_email",
            "company_member_count",
            "company_type_label",
            "business_lines",
            "charter_capital_number",
            "charter_capital_text",
            "charter_capital_full_text",
            "cash_capital_text",
            "asset_capital_text",
            "member_capital_amount",
            "member_capital_ratio",
            "capital_zero",
            "legal_representative_full_name",
            "legal_representative_id_number",
            "legal_representative_date_of_birth",
            "legal_representative_gender",
            "legal_representative_nationality",
            "legal_representative_permanent_address",
            "legal_representative_address_detail",
            "legal_representative_ward",
            "legal_representative_province",
            "legal_representative_contact_address",
            "legal_representative_issue_date",
            "legal_representative_expiry_date",
            "legal_representative_position",
            "director_full_name",
            "director_date_of_birth",
            "director_gender",
            "director_id_number",
            "director_phone",
            "director_position",
            "applicant_full_name",
            "applicant_date_of_birth",
            "applicant_gender",
            "applicant_id_number",
            "applicant_address_detail",
            "applicant_ward",
            "applicant_province",
            "authorized_person_full_name",
            "authorized_person_date_of_birth",
            "authorized_person_gender",
            "authorized_person_id_number",
            "authorized_person_contact_address",
            "authorized_person_phone",
            "authorized_person_email",
            "beneficial_owner_full_name",
            "beneficial_owner_date_of_birth",
            "beneficial_owner_gender",
            "beneficial_owner_id_number",
            "beneficial_owner_issue_date",
            "beneficial_owner_address",
            "document_place",
            "document_day",
            "document_month",
            "document_year",
            "document_place_date",
        }
        for index in range(1, 21):
            known_direct_keys.update(
                {
                    f"business_line_{index}_name",
                    f"business_line_{index}_code",
                    f"business_line_{index}_is_main",
                    f"member_{index}_full_name",
                    f"member_{index}_id_number",
                    f"member_{index}_date_of_birth",
                    f"member_{index}_gender",
                    f"member_{index}_permanent_address",
                    f"member_{index}_capital_amount",
                    f"member_{index}_capital_ratio",
                }
            )
        if normalized in known_direct_keys:
            return self._value_for_template_target(normalized, values)

        for key, raw_value in values.items():
            text_value = str(raw_value or "").strip()
            if self._normalize_key(key) == normalized and text_value:
                return self._value_for_template_target(key, values)

        aliases = {
            "phong_dang_ky_kinh_doanh": "business_registration_office",
            "ten_doanh_nghiep": "company_name",
            "ten_cong_ty": "company_name",
            "ten_viet_tat": "company_short_name",
            "dia_chi_tru_so": "company_address",
            "so_dien_thoai_cong_ty": "company_phone",
            "email_cong_ty": "company_email",
            "nganh_nghe_kinh_doanh": "business_lines",
            "von_dieu_le_bang_so": "charter_capital_number_vnd",
            "von_dieu_le_bang_chu": "charter_capital_text",
            "ho_ten": "legal_representative_full_name",
            "ho_va_ten": "legal_representative_full_name",
            "member_full_name": "member_1_full_name",
            "member_id_number": "member_1_id_number",
            "member_date_of_birth": "member_1_date_of_birth",
            "member_permanent_address": "member_1_permanent_address",
            "so_cccd": "legal_representative_id_number",
            "so_dinh_danh_ca_nhan": "legal_representative_id_number",
            "ngay_sinh": "legal_representative_date_of_birth",
            "gioi_tinh": "legal_representative_gender",
            "quoc_tich": "legal_representative_nationality",
            "dia_chi_thuong_tru": "legal_representative_permanent_address",
            "que_quan": "legal_representative_permanent_address",
            "ngay_cap": "legal_representative_issue_date",
            "ngay_het_han": "legal_representative_expiry_date",
            "chuc_danh": "legal_representative_position",
            "noi_ky": "document_place",
            "ngay_ky": "document_day",
            "thang_ky": "document_month",
            "nam_ky": "document_year",
        }
        alias_key = aliases.get(normalized)
        if alias_key:
            return self._value_for_template_target(alias_key, values)

        return self._guess_value_from_sample(placeholder, values, context)

    def _guess_value_from_sample(self, placeholder: str, values: dict[str, str], context: str = "") -> str | None:
        lowered = placeholder.lower()
        folded_placeholder = self._fold(placeholder)
        folded_context = self._fold(context)
        digits = re.sub(r"\D", "", placeholder)

        def value(key: str) -> str:
            return self._value_for_template_target(key, values)

        # The supplied templates use sample values as placeholders.  The same
        # sample name/id is reused in several sections, so resolve it from the
        # label in the same paragraph before falling back to sample mappings.
        if "toi la" in folded_context:
            person_fields = {
                "name": "applicant_full_name",
                "birth": "applicant_date_of_birth",
                "gender": "applicant_gender",
                "id": "applicant_id_number",
                "address": "applicant_address_detail",
                "ward": "applicant_ward",
                "province": "applicant_province",
            }
            if "ngay" in folded_context and "sinh" in folded_context:
                return value(person_fields["birth"])
            if "gioi tinh" in folded_context:
                return value(person_fields["gender"])
            if "so dinh danh" in folded_context:
                return value(person_fields["id"])
            if "xa/phuong" in folded_context or "dac khu" in folded_context:
                return value(person_fields["ward"])
            if "tinh/thanh pho" in folded_context:
                return value(person_fields["province"])
            if "so nha/phong" in folded_context:
                return value(person_fields["address"])
            if "toi la" in folded_context and ("thomas" in folded_placeholder or "ho" in folded_context):
                return value(person_fields["name"])

        if "giam doc" in folded_context or "tong giam doc" in folded_context:
            if "ngay" in folded_context and "sinh" in folded_context:
                return value("director_date_of_birth")
            if "gioi tinh" in folded_context:
                return value("director_gender")
            if "so dinh danh" in folded_context:
                return value("director_id_number")
            if "dien thoai" in folded_context:
                return value("director_phone")
            if "ho" in folded_context or "ten" in folded_context:
                return value("director_full_name")

        if "chu so huu huong loi" in folded_context:
            normalized_sample = self._normalize_key(placeholder)
            if normalized_sample == "27_01_2026":
                return value("beneficial_owner_issue_date")
            if normalized_sample == "11_11_1111":
                return value("beneficial_owner_date_of_birth")
            if normalized_sample == "nu":
                return value("beneficial_owner_gender")
            if normalized_sample == "thomas_editon":
                return value("beneficial_owner_full_name")
            if "ngay cap" in folded_context:
                return value("beneficial_owner_issue_date")
            if "ngay" in folded_context and "sinh" in folded_context:
                return value("beneficial_owner_date_of_birth")
            if "gioi tinh" in folded_context:
                return value("beneficial_owner_gender")
            if "dia chi" in folded_context:
                return value("beneficial_owner_address")
            if "thomas" in folded_placeholder or "ho va ten" in folded_context:
                return value("beneficial_owner_full_name")

        if "chu so huu" in folded_context and "huong loi" not in folded_context:
            if "ngay" in folded_context and "sinh" in folded_context:
                return value("legal_representative_date_of_birth")
            if "gioi tinh" in folded_context:
                return value("legal_representative_gender")
            if "so dinh danh" in folded_context:
                return value("legal_representative_id_number")
            if "so nha" in folded_context or "dia chi" in folded_context:
                return value("legal_representative_permanent_address")
            if "thomas" in folded_placeholder or "ho" in folded_context or "ten" in folded_context:
                return value("legal_representative_full_name")

        if "ben b" in folded_context:
            if "ngay" in folded_context and "sinh" in folded_context:
                return value("authorized_person_date_of_birth")
            if "gioi tinh" in folded_context:
                return value("authorized_person_gender")
            if "so dinh danh" in folded_context:
                return value("authorized_person_id_number")
            if "dia chi" in folded_context:
                return value("authorized_person_contact_address")
            if "dien thoai" in folded_context:
                return value("authorized_person_phone")
            if "email" in folded_context:
                return value("authorized_person_email")
            if "ho" in folded_context or "ten" in folded_context:
                return value("authorized_person_full_name")

        if "uy quyen" in folded_context and "ben b" not in folded_context:
            if "ngay" in folded_context and "sinh" in folded_context:
                return value("legal_representative_date_of_birth")
            if "gioi tinh" in folded_context:
                return value("legal_representative_gender")
            if "so dinh danh" in folded_context:
                return value("legal_representative_id_number")
            if "dia chi" in folded_context:
                return value("legal_representative_permanent_address")
            if "thomas" in folded_placeholder or "ho" in folded_context or "ten" in folded_context:
                return value("legal_representative_full_name")

        if "ben a" in folded_context or "nguoi dai dien theo phap luat" in folded_context:
            if "ngay" in folded_context and "sinh" in folded_context:
                return value("legal_representative_date_of_birth")
            if "gioi tinh" in folded_context:
                return value("legal_representative_gender")
            if "so dinh danh" in folded_context:
                return value("legal_representative_id_number")
            if "dia chi" in folded_context:
                return value("legal_representative_permanent_address")
            if "ho" in folded_context or "ten" in folded_context or "ky" in folded_context:
                return value("legal_representative_full_name")

        company_name_samples = {
            "cong_ty_tnhh_thuong_mai_va_dich_vu_editon",
            "cong_ty_tnhh_thuong_mai_va_dich_vu_edition",
            "cong_ty_tnhh_thuong_mai_va_dich_vu_thomas",
        }
        if folded_placeholder in company_name_samples:
            return value("company_name")
        if "tru so cong ty" in folded_context or "ten cong ty viet bang tieng viet" in folded_context:
            if "so nha" in folded_placeholder or "thon" in folded_placeholder or "xa" in folded_placeholder:
                return value("company_address")
            if "dien thoai" in folded_context:
                return value("company_phone")

        if "von dieu le" in folded_context or "von bang tien mat" in folded_context:
            if "bang chu" in folded_context:
                return value("charter_capital_text")
            if "von bang tien mat" in folded_context:
                return value("cash_capital_text") or value("charter_capital_full_text")
            if "bang so" in folded_context:
                return value("charter_capital_number_vnd")
            return value("charter_capital_full_text") or value("charter_capital_number")

        if "ngay 23 thang 03 nam 2026" in folded_placeholder or (
            "dieu le nay" in folded_context and "thong qua" in folded_context
        ):
            return value("document_place_date")

        if "ngay" in folded_context and "thang" in folded_context and "nam" in folded_context:
            return value("document_place_date")

        direct_sample_map = {
            "0": "capital_zero",
            "100": "member_capital_ratio",
            "ba_ty_dong": "charter_capital_text",
            "company": "company_name",
            "bac": "company_short_name",
            "cong_ty_tnhh_thuong_mai_va_dich_vu_editon": "company_name",
            "cong_ty_tnhh_thuong_mai_va_dich_vu_edition": "company_name",
            "cong_ty_tnhh_thuong_mai_va_dich_vu_thomas": "company_name",
            "so_nha_13_thon_4_xa_hoa_lac_ha_noi": "company_address",
            "3_000_000_000_vnd": "charter_capital_number_vnd",
            "3_000_000_000_dong_ba_ty_dong": "charter_capital_full_text",
            "ngay_23_thang_03_nam_2026": "document_place_date",
            "x": "business_line_1_is_main",
            "nguyen_thi_khanh_haha": "authorized_person_full_name",
            "22_12_1111": "authorized_person_date_of_birth",
            "098765434567": "authorized_person_id_number",
            "thon_lam_xa_hai_ba_tinh_bon_nam": "authorized_person_contact_address",
            "09876545678": "authorized_person_phone",
            "fsgdhfjg_gmail_com": "authorized_person_email",
            "nguyen_thanh_hoa": "legal_representative_full_name",
            "080177011933": "legal_representative_id_number",
            "12_07_1977": "legal_representative_date_of_birth",
            "ap_hau_quoi_hau_my_dong_thap": "legal_representative_permanent_address",
            "thon_hai_xa_ba_tinh_bon_nam": "beneficial_owner_address",
            "09876654311": "company_phone",
            "09876543": "director_phone",
            "so_nha_13": "company_address_detail",
            "tinh_ha_giang": "legal_representative_province",
        }
        mapped_key = direct_sample_map.get(self._normalize_key(placeholder))
        if mapped_key:
            mapped_value = value(mapped_key)
            if self._normalize_key(placeholder) == "100":
                return mapped_value.rstrip("% ")
            return mapped_value

        business_line_map = {
            "ban_buon_gao_lua_my_san_pham_tu_ngu_coc_khac_bot_mychi_tiet_ban_buon_gao": "business_line_1_name",
            "4631": "business_line_1_code",
            "ban_buon_thuc_pham": "business_line_2_name",
            "4632": "business_line_2_code",
            "ban_buon_nong_lam_san_nguyen_lieu_tru_go_tre_nua_va_dong_vat_song_chi_tiet_ban_buon_lua_tam_cam_lua_giong_va_cac_mat_hang_nong_lam_san_nguyen_lieu_tru_go_tre_nua": "business_line_3_name",
            "4620": "business_line_3_code",
            "dai_ly_moi_gioi_dau_gia_hang_hoa_chi_tiet_tru_dau_gia_hang_hoa": "business_line_4_name",
            "dai_ly_moi_gioi_dau_gia_hang_hoachi_tiet_tru_dau_gia_hang_hoa": "business_line_4_name",
            "4610": "business_line_4_code",
            "ban_le_luong_thuc": "business_line_5_name",
            "4721": "business_line_5_code",
            "hoat_dong_dich_vu_ho_tro_khac_lien_quan_den_van_tai_chi_tiet_hoat_dong_cua_cac_dai_ly_ban_ve_may_bay_tru_hoa_long_khi_de_van_chuyen_va_hoat_dong_lien_quan_den_van_tai_duong_hang_khong": "business_line_6_name",
            "hoat_dong_dich_vu_ho_tro_kinh_doanh_khac_con_lai_chua_duoc_phan_vao_dau_chi_tiet_xuat_nhap_khau_cac_mat_hang_doanh_nghiep_kinh_doanh": "business_line_6_name",
            "8299": "business_line_6_code",
            "kho_bai_va_luu_giu_hang_hoa_chi_tiet_luu_giu_hang_hoa_tru_kinh_doanh_kho_bai": "business_line_7_name",
            "5210": "business_line_7_code",
            "5229": "business_line_8_code",
            "hoat_dong_tu_van_quan_ly_kinh_doanh_va_hoat_dong_tu_van_quan_ly_khac_tru_tu_van_tai_chinh_ke_toan_phap_luat": "business_line_8_name",
            "7020": "business_line_9_code",
            "quang_cao_chi_tiet_tru_quang_cao_thuoc_la": "business_line_9_name",
            "7310": "business_line_10_code",
        }
        business_key = business_line_map.get(self._normalize_key(placeholder))
        if business_key:
            return value(business_key)

        if folded_placeholder == "x" and "nganh nghe kinh doanh chinh" in folded_context:
            for index in range(1, 11):
                if values.get(f"business_line_{index}_is_main"):
                    return "X"
            return ""
        if "phong dang ky kinh doanh" in folded_placeholder or "phong dang ky kinh doanh" in folded_context:
            return value("business_registration_office")
        if "von bang tai san" in folded_context or "khong co" in folded_placeholder:
            return value("asset_capital_text")
        if "von bang tien mat" in folded_context:
            return value("cash_capital_text") or value("charter_capital_full_text")
        if "von dieu le" in folded_context:
            if "bang chu" in folded_context:
                return value("charter_capital_text")
            if "bang so" in folded_context or "vnd" in folded_context:
                return value("charter_capital_number_vnd")
            return value("charter_capital_full_text") or value("charter_capital_number")
        if "dong" in folded_placeholder or "vnd" in folded_placeholder:
            return value("charter_capital_number_vnd")
        if folded_context.startswith("[") and "ngay" in folded_placeholder and "thang" in folded_placeholder:
            return value("document_place_date")
        if folded_placeholder.startswith("ngay") and "thang" in folded_placeholder and "nam" in folded_placeholder:
            return value("document_place_date")
        if "noi ky" in folded_context:
            return value("document_place")
        if "ngay [" in folded_context and "thang [" in folded_context and "nam [" in folded_context:
            if len(digits) == 4:
                return value("document_year")
            if placeholder.strip() == str(values.get("document_month", "") or ""):
                return value("document_month")
            return value("document_day")
        if "ngay cap" in folded_context:
            return value("legal_representative_issue_date") or value("beneficial_owner_issue_date")
        if "ngay, thang, nam sinh" in lowered or "ngay thang nam sinh" in folded_context:
            if "22/12/1111" in placeholder:
                return value("authorized_person_date_of_birth")
            return value("legal_representative_date_of_birth") or value("applicant_date_of_birth")
        if "tinh/thanh pho truc thuoc trung uong" in folded_context:
            if "tinh ha giang" in folded_placeholder:
                return value("legal_representative_province") or value("applicant_province")
            return value("company_province") or value("applicant_province")
        if "xa/phuong" in folded_context or "dac khu" in folded_context:
            return value("company_ward") or value("applicant_ward") or value("legal_representative_ward")
        if "so nha/phong" in folded_context:
            if "thôn" in lowered or "thon" in folded_placeholder:
                return value("applicant_address_detail") or value("legal_representative_address_detail")
            return value("company_address_detail") or value("legal_representative_address_detail")
        if "dia chi lien lac" in folded_context:
            return (
                value("authorized_person_contact_address")
                or value("legal_representative_contact_address")
                or value("legal_representative_permanent_address")
            )
        if "dia chi thuong tru" in folded_context:
            return value("legal_representative_permanent_address")
        if "giam doc" in folded_context:
            if "dien thoai" in folded_context:
                return value("director_phone")
            return value("director_full_name") or value("legal_representative_full_name")
        if "so dinh danh ca nhan" in folded_context or re.fullmatch(r"\d{9,12}", digits):
            if "098765434567" in placeholder:
                return value("authorized_person_id_number")
            return value("legal_representative_id_number") or value("applicant_id_number")
        if "gioi tinh" in folded_context or folded_placeholder in {"nam", "nu"}:
            return value("legal_representative_gender") or value("applicant_gender")
        if "quoc gia" in folded_context or "viet nam" in folded_placeholder:
            return value("legal_representative_nationality")
        if "dien thoai" in folded_context:
            return value("company_phone") or value("authorized_person_phone")
        if "email" in folded_context:
            return value("company_email") or value("authorized_person_email")
        if "ten cong ty viet bang tieng viet" in folded_context or "cong ty tnhh" in folded_placeholder:
            return value("company_name")
        if "tru so cong ty" in folded_context:
            return value("company_address")
        if "toi la" in folded_context:
            return value("applicant_full_name") or value("legal_representative_full_name")
        if "ho chu dem va ten" in folded_context or "ho, chu dem va ten" in lowered:
            return value("legal_representative_full_name") or value("applicant_full_name")
        if "thomas" in folded_placeholder:
            if "viet tat" in folded_context or "ten viet tat" in folded_context:
                return value("company_short_name") or value("company_name")
            if "cong ty" in folded_context or "doanh nghiep" in folded_context:
                return value("company_name")
            return value("company_short_name") or value("company_name")
        if re.search(r"\d{1,2}/\d{1,2}/\d{4}", placeholder):
            return value("legal_representative_date_of_birth") or value("document_place_date")
        if any(word in lowered for word in ["tỉnh", "xã", "phường", "thôn", "số nhà"]):
            return value("legal_representative_permanent_address") or value("company_address")
        return None

    def _normalize_key(self, value: str) -> str:
        text = self._fold(value)
        text = re.sub(r"[^a-z0-9]+", "_", text)
        return text.strip("_")

    def _fold(self, value: str) -> str:
        value = value.replace(chr(273), "d").replace(chr(272), "D")
        text = value.strip().lower().replace("đ", "d")
        text = unicodedata.normalize("NFD", text)
        text = "".join(char for char in text if unicodedata.category(char) != "Mn")
        return text.replace("đ", "d")

    def _safe_filename(self, value: str) -> str:
        text = re.sub(r"[^\w\s.-]", "", value, flags=re.UNICODE).strip()
        text = re.sub(r"\s+", "_", text)
        return text[:80] or "ho_so"









