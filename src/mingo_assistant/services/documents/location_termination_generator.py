"""Fill the highlighted fields in the location-termination dossier."""
from __future__ import annotations

from pathlib import Path

from docx import Document

from mingo_assistant.services.documents.dependent_unit_generator import DependentUnitTemplateGenerator


class LocationTerminationTemplateGenerator(DependentUnitTemplateGenerator):
    """Keep every supplied template run and replace only highlighted values."""

    NOTICE_NAME = "1. Thông báo chấm dứt hoạt động.docx"
    AUTHORIZATION_NAME = "2. Giấy ủy quyền.docx"
    SIGNING_GUIDE_NAME = "3. Huong dan ky.docx"

    def _templates(self, values: dict[str, str] | None = None) -> list[Path]:
        return [
            self.template_dir / self.NOTICE_NAME,
            self.template_dir / self.AUTHORIZATION_NAME,
            self.template_dir / self.SIGNING_GUIDE_NAME,
        ]

    def _mode(self, template_name: str) -> str:
        return "fill_highlighted" if template_name in {
            self.NOTICE_NAME,
            self.AUTHORIZATION_NAME,
            self.SIGNING_GUIDE_NAME,
        } else "copy"

    def _output_filename(self, template_name: str, values: dict[str, str], case_name: str) -> str:
        label_by_template = {
            self.NOTICE_NAME: "Thong bao cham dut hoat dong",
            self.AUTHORIZATION_NAME: "Giay uy quyen",
            self.SIGNING_GUIDE_NAME: "Huong dan ky",
        }
        label = label_by_template.get(template_name, "Ho so cham dut")
        return label + "_" + self._safe_filename(case_name) + ".docx"

    @staticmethod
    def _highlighted_runs(paragraph) -> list:
        return [run for run in paragraph.runs if run.font.highlight_color is not None]

    def _replace_highlighted_runs(self, docx_path: Path, template_name: str, values: dict[str, str]) -> None:
        document = Document(docx_path)
        if template_name == self.NOTICE_NAME:
            self._fill_notice(document, values)
            expected_highlights = 13
            template_label = "thông báo chấm dứt hoạt động"
        elif template_name == self.AUTHORIZATION_NAME:
            self._fill_authorization(document, values)
            expected_highlights = 20
            template_label = "giấy ủy quyền"
        elif template_name == self.SIGNING_GUIDE_NAME:
            self._fill_signing_guide(document, values)
            expected_highlights = 2
            template_label = "hướng dẫn ký"
        else:
            raise ValueError(f"Không hỗ trợ mẫu Word: {template_name}")
        actual_highlights = sum(
            len(self._highlighted_runs(paragraph))
            for paragraph in self._iter_paragraphs(document)
        )
        if actual_highlights != expected_highlights:
            raise ValueError(
                f"Mẫu {template_label} có {actual_highlights} vùng tô sáng, dự kiến {expected_highlights}."
            )
        document.save(docx_path)

    def _fill_notice(self, document: Document, values: dict[str, str]) -> None:
        if len(document.tables) < 2:
            raise ValueError("Mẫu thông báo chấm dứt hoạt động đã thay đổi cấu trúc bảng.")

        def fill(paragraph, field: str) -> None:
            runs = self._highlighted_runs(paragraph)
            if len(runs) != 1:
                raise ValueError(f"Không xác định được vùng tô sáng: {paragraph.text[:60]}")
            replacement = str(values.get(field, "") or "").strip()
            if replacement:
                runs[0].text = replacement

        header = document.tables[0].cell(0, 0).paragraphs[0]
        fill(header, "parent_entity_name")
        date = document.tables[0].cell(0, 1).paragraphs[3]
        date_runs = self._highlighted_runs(date)
        if len(date_runs) != 3:
            raise ValueError("Không xác định được ba vùng tô sáng của ngày ký.")
        for run, field in zip(date_runs, ("document_day", "document_month", "document_year")):
            replacement = str(values.get(field, "") or "").strip()
            if replacement:
                run.text = replacement

        fields_by_prefix = {
            "Tên doanh nghiệp (": "parent_entity_name",
            "Mã số doanh nghiệp/Mã số thuế:": "parent_entity_registration_number",
            "1. Tên địa điểm kinh doanh (": "location_business_name",
            "Mã số địa điểm kinh doanh:": "location_business_code",
            "Số nhà/phòng,": "location_address_detail",
            "Xã/Phường/Đặc khu:": "location_ward",
            "Tỉnh/Thành phố trực thuộc trung ương:": "location_province",
            "4. Lý do chấm dứt hoạt động:": "termination_reason",
        }
        for prefix, field in fields_by_prefix.items():
            matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
            if len(matches) != 1:
                raise ValueError(f"Không xác định được dòng mẫu Word: {prefix}")
            fill(matches[0], field)

        signature = document.tables[1].cell(0, 1).paragraphs[8]
        fill(signature, "legal_representative_full_name")

    def _fill_authorization(self, document: Document, values: dict[str, str]) -> None:
        if len(document.tables) != 1:
            raise ValueError("Mẫu giấy ủy quyền đã thay đổi cấu trúc bảng ký.")

        def fill(paragraph, field: str, *, value: str | None = None) -> None:
            runs = self._highlighted_runs(paragraph)
            if len(runs) != 1:
                raise ValueError(f"Không xác định được vùng tô sáng: {paragraph.text[:60]}")
            replacement = str(values.get(field, "") if value is None else value).strip()
            if replacement:
                runs[0].text = replacement

        def paragraph(index: int):
            try:
                return document.paragraphs[index]
            except IndexError as exc:
                raise ValueError("Mẫu giấy ủy quyền không đủ các dòng nội dung.") from exc

        fill(paragraph(5), "legal_representative_full_name")
        fill(paragraph(6), "legal_representative_id_number")
        fill(paragraph(7), "legal_representative_permanent_address")
        legal_company = paragraph(8)
        legal_company_runs = self._highlighted_runs(legal_company)
        if len(legal_company_runs) != 3:
            raise ValueError("Không xác định được ba vùng thông tin doanh nghiệp của Bên A.")
        for run, field in zip(
            legal_company_runs,
            ("parent_entity_name", "parent_entity_registration_number", "legal_representative_position"),
        ):
            replacement = str(values.get(field, "") or "").strip()
            if replacement:
                run.text = replacement

        fill(paragraph(10), "authorized_person_full_name")
        fill(paragraph(11), "authorized_person_id_number")
        authorized_contact = paragraph(13)
        authorized_contact_runs = self._highlighted_runs(authorized_contact)
        if len(authorized_contact_runs) != 2:
            raise ValueError("Không xác định được thông tin liên hệ của Bên B.")
        for run, field in zip(authorized_contact_runs, ("authorized_person_phone", "authorized_person_email")):
            replacement = str(values.get(field, "") or "").strip()
            if replacement:
                run.text = replacement

        office = str(values.get("business_registration_office", "") or "").strip()
        office_prefix = "Phòng Đăng ký kinh doanh"
        office_location = office[len(office_prefix):].strip() if office.startswith(office_prefix) else ""
        # Word collapses a trailing ordinary space at a run boundary. Keep a
        # visible separator between the two independently highlighted slots.
        authority = f"{office_prefix}\u00A0" if office_location else office
        office_line = paragraph(15)
        office_runs = self._highlighted_runs(office_line)
        if len(office_runs) != 2:
            raise ValueError("Không xác định được cơ quan tiếp nhận trong Điều 1.")
        for run, replacement in zip(office_runs, (authority, office_location)):
            if replacement:
                run.text = replacement

        procedure_line = paragraph(16)
        procedure_runs = self._highlighted_runs(procedure_line)
        if len(procedure_runs) != 2:
            raise ValueError("Không xác định được nội dung ủy quyền trong Điều 1.")
        procedure_runs[0].text = "đăng ký chấm dứt hoạt động của địa điểm kinh doanh"
        location_name = str(values.get("location_business_name", "") or "").strip()
        if location_name:
            procedure_runs[1].text = location_name

        fill(paragraph(19), "business_registration_office")

        signing = paragraph(33)
        signing_runs = self._highlighted_runs(signing)
        if len(signing_runs) != 4:
            raise ValueError("Không xác định được bốn vùng tô sáng của ngày ký giấy ủy quyền.")
        for run, field in zip(
            signing_runs,
            ("document_place", "document_day", "document_month", "document_year"),
        ):
            replacement = str(values.get(field, "") or "").strip()
            if replacement:
                run.text = replacement

        signature = document.tables[0].cell(0, 1).paragraphs[5]
        fill(signature, "legal_representative_full_name")

    def _fill_signing_guide(self, document: Document, values: dict[str, str]) -> None:
        if len(document.tables) != 1 or len(document.tables[0].rows) != 3:
            raise ValueError("Mẫu hướng dẫn ký đã thay đổi cấu trúc bảng.")
        # The signing guide is an instruction sheet, so render the signer in
        # normal name case even when legal-form data is stored in ALL CAPS.
        signer = str(values.get("legal_representative_full_name", "") or "").strip().title()
        for row_index in (1, 2):
            paragraph = document.tables[0].cell(row_index, 3).paragraphs[0]
            runs = self._highlighted_runs(paragraph)
            if len(runs) != 1:
                raise ValueError(f"Không xác định được người ký tại dòng {row_index} của mẫu hướng dẫn ký.")
            if signer:
                runs[0].text = signer
