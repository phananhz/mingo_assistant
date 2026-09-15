from __future__ import annotations

import hashlib
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

from mingo_assistant.services.documents.business_registration_variants import (
    TWO_MEMBER_INDIVIDUAL_FILE1,
    TWO_MEMBER_INDIVIDUAL_FILE2,
    TWO_MEMBER_INDIVIDUAL_FILE3,
    TWO_MEMBER_INDIVIDUAL_FILE4,
    TWO_MEMBER_INDIVIDUAL_TEMPLATE_RELATIVE_DIR,
)
from mingo_assistant.services.documents.docx_generator import (
    TwoMemberIndividualFile1DocxTemplateGenerator,
)
from mingo_assistant.ui.views.business_registration.schema import FIELD_GROUPS


class TwoMemberIndividualFile1Test(unittest.TestCase):
    def setUp(self) -> None:
        self.project_root = Path(__file__).resolve().parents[2]
        self.template_dir = (
            self.project_root
            / "assets"
            / "templates"
            / "word"
            / "business_registration"
            / TWO_MEMBER_INDIVIDUAL_TEMPLATE_RELATIVE_DIR
        )

    def test_generator_uses_file1_and_file2_and_preserves_source_templates(self) -> None:
        source_template = self.template_dir / TWO_MEMBER_INDIVIDUAL_FILE1
        source_hash = hashlib.sha256(source_template.read_bytes()).hexdigest()
        file2_template = self.template_dir / TWO_MEMBER_INDIVIDUAL_FILE2
        file2_hash = hashlib.sha256(file2_template.read_bytes()).hexdigest()
        values = {
            "company_name": "CÔNG TY TNHH MẪU HAI THÀNH VIÊN",
            "applicant_full_name": "NGUYỄN VĂN MẪU",
            "applicant_date_of_birth": "01/02/1990",
            "applicant_gender": "Nam",
            "applicant_id_number": "001090000001",
            "legal_representative_full_name": "NGUYỄN VĂN MẪU",
            "legal_representative_date_of_birth": "01/02/1990",
            "legal_representative_gender": "Nam",
            "legal_representative_id_number": "001090000001",
            "legal_representative_position": "Giám đốc",
            "document_place": "Hà Nội",
            "document_day": "14",
            "document_month": "09",
            "document_year": "2026",
            "business_registration_office": "Phòng Đăng ký kinh doanh Hà Nội",
            "legal_representative_nationality": "Việt Nam",
            "company_address_detail": "Số 1 phố Mẫu",
            "company_ward": "Phường Ba Đình",
            "company_province": "Thành phố Hà Nội",
            "charter_capital_number_vnd": "100.000.000 đồng",
            "charter_capital_text": "Một trăm triệu đồng",
            "capital_source_1_amount": "0",
            "capital_source_1_ratio": "0%",
            "capital_source_2_amount": "100.000.000",
            "capital_source_2_ratio": "100%",
            "capital_source_3_amount": "0",
            "capital_source_3_ratio": "0%",
            "capital_source_4_amount": "0",
            "capital_source_4_ratio": "0%",
            "capital_source_total_amount": "100.000.000",
            "capital_source_total_ratio": "100%",
            "director_full_name": "NGUYỄN VĂN MẪU",
            "director_date_of_birth": "01/02/1990",
            "director_id_number": "001090000001",
            "director_phone": "0999999999",
            "tax_labor_count": "2",
            "director_gender": "Nam",
            "company_phone": "0901000001",
            "company_fax": "0241234567",
            "company_email": "contact@example.vn",
            "company_website": "example.vn",
            "company_english_name": "SAMPLE TWO MEMBER COMPANY LIMITED",
            "company_short_name": "SAMPLE 2TV",
            "company_address": "Số 1 phố Mẫu, Phường Ba Đình, Thành phố Hà Nội",
            "charter_capital_number": "100.000.000",
            "member_count": "2",
            "member_1_full_name": "NGUYỄN VĂN MẪU",
            "member_1_date_of_birth": "01/02/1990",
            "member_1_gender": "Nam",
            "member_1_id_number": "001090000001",
            "member_1_permanent_address": "Số 1 phố Mẫu, Hà Nội",
            "member_1_capital_amount": "50000000",
            "member_1_capital_ratio": "50%",
            "member_1_capital_asset_type": "Quy\u1ec1n s\u1eed d\u1ee5ng tài s\u1ea3n",
            "member_1_capital_deadline": "Hoàn thành trong 30 ngày",
            "member_2_full_name": "TRẦN THỊ BÌNH",
            "member_2_date_of_birth": "02/03/1991",
            "member_2_gender": "Nữ",
            "member_2_id_number": "001091000002",
            "member_2_permanent_address": "Số 2 phố Mẫu, Hà Nội",
            "member_2_capital_amount": "50.000.000",
            "member_2_capital_ratio": "50%",
            "beneficial_owner_list_count": "2",
            "beneficial_owner_list_1_number": "1",
            "beneficial_owner_list_1_full_name": "MEMBER ONE",
            "beneficial_owner_list_1_date_of_birth": "01/01/1990",
            "beneficial_owner_list_1_gender": "Nam",
            "beneficial_owner_list_1_identity": "001090000001",
            "beneficial_owner_list_1_address": "Address One",
            "beneficial_owner_list_1_capital_direct": "80",
            "beneficial_owner_list_1_capital_total": "80",
            "beneficial_owner_list_2_number": "2",
            "beneficial_owner_list_2_full_name": "MEMBER TWO",
            "beneficial_owner_list_2_date_of_birth": "02/02/1991",
            "beneficial_owner_list_2_gender": "N?",
            "beneficial_owner_list_2_identity": "001091000002",
            "beneficial_owner_list_2_address": "Address Two",
            "beneficial_owner_list_2_capital_direct": "20",
            "beneficial_owner_list_2_capital_total": "20",
            "business_line_1_name": "Trồng ngô và cây lương thực có hạt khác",
            "business_line_1_code": "0112",
            "business_line_1_is_main": "X",
            "business_line_2_name": "Trồng cây thuốc lá",
            "business_line_2_code": "0115",
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            generator = TwoMemberIndividualFile1DocxTemplateGenerator(
                self.template_dir, Path(temp_dir)
            )
            result = generator.generate(values, "ho_so", Path(temp_dir))

            self.assertEqual(
                [TWO_MEMBER_INDIVIDUAL_FILE1, TWO_MEMBER_INDIVIDUAL_FILE2, TWO_MEMBER_INDIVIDUAL_FILE3, TWO_MEMBER_INDIVIDUAL_FILE4, "4. Giấy ủy quyền.docx", "6. Huong dan ky.docx"],
                [item.source_name for item in result.generated],
            )
            self.assertEqual(
                [TWO_MEMBER_INDIVIDUAL_FILE1, TWO_MEMBER_INDIVIDUAL_FILE2, TWO_MEMBER_INDIVIDUAL_FILE3, TWO_MEMBER_INDIVIDUAL_FILE4, "5. Giay uy quyen.docx", "6. Huong dan ky.docx"],
                [item.output_path.name for item in result.generated],
            )
            output_path = result.generated[0].output_path
            output_by_source = {item.source_name: item.output_path for item in result.generated}
            self.assertTrue(output_path.is_file())
            generated_text = "\n".join(paragraph.text for paragraph in Document(output_path).paragraphs)
            self.assertIn("CÔNG TY TRÁCH NHIỆM HỮU HẠN HAI THÀNH VIÊN TRỞ LÊN", generated_text)
            self.assertIn("NGUYỄN VĂN MẪU", generated_text)

            generated_document = Document(output_path)
            all_text = "\n".join(
                paragraph.text
                for paragraph in TwoMemberIndividualFile1DocxTemplateGenerator(self.template_dir, Path(temp_dir))._iter_paragraphs(generated_document)
            )
            self.assertIn("Gi\u1edbi t\u00ednh: Nam", all_text)
            self.assertIn("\u0110i\u1ec7n tho\u1ea1i: 0901000001", all_text)
            self.assertIn("S\u1ed1 fax (n\u1ebfu c\u00f3): 0241234567", all_text)
            self.assertIn("Th\u01b0 \u0111i\u1ec7n t\u1eed (n\u1ebfu c\u00f3): contact@example.vn", all_text)
            self.assertGreaterEqual(all_text.count("\u0110i\u1ec7n tho\u1ea1i: 0901000001"), 2)
            self.assertGreaterEqual(all_text.count("S\u1ed1 fax (n\u1ebfu c\u00f3): 0241234567"), 2)
            self.assertGreaterEqual(all_text.count("Th\u01b0 \u0111i\u1ec7n t\u1eed (n\u1ebfu c\u00f3): contact@example.vn"), 2)

            file1 = Document(output_by_source[TWO_MEMBER_INDIVIDUAL_FILE1])
            industry_table = next(
                table
                for table in file1.tables
                if TwoMemberIndividualFile1DocxTemplateGenerator(self.template_dir, Path(temp_dir))._business_line_header_index(table) is not None
            )
            header_index = TwoMemberIndividualFile1DocxTemplateGenerator(self.template_dir, Path(temp_dir))._business_line_header_index(industry_table)
            self.assertEqual("1", industry_table.rows[header_index + 1].cells[0].text)
            self.assertEqual("2", industry_table.rows[header_index + 2].cells[0].text)
            self.assertNotIn("numPr", industry_table.rows[header_index + 1].cells[0]._tc.xml)
            self.assertNotIn("numPr", industry_table.rows[header_index + 2].cells[0]._tc.xml)

            member_list = Document(output_by_source[TWO_MEMBER_INDIVIDUAL_FILE3])
            self.assertEqual(5, len(member_list.tables[0].rows))
            self.assertEqual("NGUY\u1ec4N V\u0102N M\u1eaaU", member_list.tables[0].rows[3].cells[1].text)
            self.assertEqual("TR\u1ea6N TH\u1eca B\u00ccNH", member_list.tables[0].rows[4].cells[1].text)
            self.assertEqual("", member_list.tables[0].rows[3].cells[5].text)
            self.assertEqual("", member_list.tables[0].rows[3].cells[12].text)
            self.assertIn("50.000.000", member_list.tables[0].rows[3].cells[8].text)

            charter = Document(output_by_source[TWO_MEMBER_INDIVIDUAL_FILE2])
            charter_text = "\n".join(paragraph.text for paragraph in TwoMemberIndividualFile1DocxTemplateGenerator(self.template_dir, Path(temp_dir))._iter_paragraphs(charter))
            self.assertIn("NGUY\u1ec4N V\u0102N M\u1eaaU", charter_text)
            self.assertIn("TR\u1ea6N TH\u1eca B\u00ccNH", charter_text)
            self.assertIn("SAMPLE TWO MEMBER COMPANY LIMITED", charter_text)
            self.assertIn("0901000001", charter_text)
            self.assertIn("contact@example.vn", charter_text)
            self.assertNotIn("L\u00ca TH\u1eca T", charter_text)
            charter_paragraphs = list(
                TwoMemberIndividualFile1DocxTemplateGenerator(
                    self.template_dir, Path(temp_dir)
                )._iter_paragraphs(charter)
            )
            body_heading_index = next(
                index
                for index, paragraph in enumerate(charter_paragraphs)
                if "co trach nhiem sau day"
                in generator._fold(paragraph.text)
            )
            self.assertNotEqual(
                "NGUYỄN VĂN MẪU",
                charter_paragraphs[body_heading_index + 1].text.strip(),
            )
            signature_heading_index = next(
                index
                for index, paragraph in enumerate(charter_paragraphs)
                if generator._fold(paragraph.text)
                == "nguoi dai dien theo phap luat cua doanh nghiep"
            )
            signature_name = next(
                paragraph
                for paragraph in charter_paragraphs[signature_heading_index + 1 :]
                if paragraph.text.strip()
            )
            self.assertEqual("NGUYỄN VĂN MẪU", signature_name.text.strip())
            self.assertTrue(
                all(
                    run.font.highlight_color == WD_COLOR_INDEX.YELLOW
                    for run in signature_name.runs
                    if run.text
                )
            )

            member_signature_table = next(
                table
                for table in charter.tables
                if len(table.rows) >= 1
                and len(table.rows[0].cells) == 2
                and "thanh vien" in generator._fold(" ".join(cell.text for cell in table.rows[0].cells))
            )
            self.assertEqual("Thành viên", member_signature_table.rows[0].cells[0].paragraphs[0].text)
            member_signature_name = next(
                paragraph
                for paragraph in reversed(member_signature_table.rows[0].cells[0].paragraphs)
                if paragraph.text.strip()
            )
            self.assertEqual("NGUYỄN VĂN MẪU", member_signature_name.text.strip())
            self.assertTrue(
                all(
                    run.font.highlight_color == WD_COLOR_INDEX.YELLOW and run.bold is True
                    for run in member_signature_name.runs
                    if run.text
                )
            )

            self.assertEqual(3, len(charter.tables[0].rows))
            self.assertEqual(4, len(charter.tables[2].rows))

            beneficial_owner = Document(output_by_source[TWO_MEMBER_INDIVIDUAL_FILE4])
            self.assertEqual(6, len(beneficial_owner.tables[0].rows))
            self.assertEqual("MEMBER ONE", beneficial_owner.tables[0].rows[4].cells[1].text.strip())
            self.assertEqual("MEMBER TWO", beneficial_owner.tables[0].rows[5].cells[1].text.strip())
            self.assertEqual("80%", beneficial_owner.tables[0].rows[4].cells[8].text.strip())
            self.assertIn("NGUY", beneficial_owner.tables[1].rows[0].cells[0].text)
            template_file4 = Document(self.template_dir / TWO_MEMBER_INDIVIDUAL_FILE4)
            for column in (0, 1, 2, 3, 4, 7, 8, 10):
                source_cell = template_file4.tables[0].rows[4].cells[column]
                output_cell = beneficial_owner.tables[0].rows[4].cells[column]
                source_p = next(p for p in source_cell.paragraphs if p.text.strip())
                output_p = next(p for p in output_cell.paragraphs if p.text.strip())
                source_run = next(r for r in source_p.runs if r.text.strip())
                output_run = next(r for r in output_p.runs if r.text.strip())
                self.assertEqual(source_run.font.size, output_run.font.size)
                self.assertEqual(source_p.alignment, output_p.alignment)
            guide = Document(output_by_source["6. Huong dan ky.docx"])
            names = [r.text for p in generator._iter_paragraphs(guide)
                     for r in p.runs if r.text and r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
            self.assertGreaterEqual(len(names), 5)
            self.assertTrue(all(name == "Nguyễn Văn Mẫu" for name in names))
            self.assertEqual("50", charter.tables[2].rows[2].cells[3].text)
            self.assertEqual("50", charter.tables[2].rows[3].cells[3].text)
            self.assertEqual("Quy\u1ec1n s\u1eed d\u1ee5ng tài s\u1ea3n", charter.tables[2].rows[2].cells[4].text)
            self.assertEqual("Hoàn thành trong 30 ngày", charter.tables[2].rows[2].cells[5].text)

            template_file2 = Document(self.template_dir / TWO_MEMBER_INDIVIDUAL_FILE2)
            for column in range(6):
                template_run = next(
                    run
                    for paragraph in template_file2.tables[0].rows[1].cells[column].paragraphs
                    for run in paragraph.runs
                    if run.text
                )
                generated_run = next(
                    run
                    for paragraph in charter.tables[0].rows[1].cells[column].paragraphs
                    for run in paragraph.runs
                    if run.text
                )
                self.assertEqual(
                    template_run._element.rPr.find(qn("w:sz")).get(qn("w:val")),
                    generated_run._element.rPr.find(qn("w:sz")).get(qn("w:val")),
                )
        self.assertEqual(source_hash, hashlib.sha256(source_template.read_bytes()).hexdigest())
        self.assertEqual(file2_hash, hashlib.sha256(file2_template.read_bytes()).hexdigest())

    def test_applicant_identity_label_is_correct(self) -> None:
        applicant_group = next(fields for title, fields in FIELD_GROUPS if "người đăng ký" in title)
        labels = dict(applicant_group)
        self.assertEqual("Số định danh cá nhân", labels["applicant_id_number"])
