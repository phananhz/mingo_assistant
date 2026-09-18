from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from mingo_assistant.services.documents.dependent_unit_generator import DependentUnitTemplateGenerator


def _paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        yield from _table_paragraphs(table)


def _table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _table_paragraphs(nested)


class DdkdUpdateTemplatesTest(unittest.TestCase):
    def setUp(self) -> None:
        self.project_root = Path(__file__).resolve().parents[2]
        self.template_dir = self.project_root / "assets" / "templates" / "word" / "dependent_units" / "ddkd_update"
        self.values = {
            "legal_representative_full_name": "NGUYEN VAN A",
            "legal_representative_id_number": "001090000001",
            "legal_representative_position": "Director",
            "parent_entity_name": "COMPANY SAMPLE",
            "parent_entity_registration_number": "0312345678",
            "company_address": "1 Sample Street, Ha Noi",
            "authorized_person_full_name": "TRAN THI B",
            "authorized_person_id_number": "001090000002",
            "authorized_person_contact_address": "2 Sample Street, Ha Noi",
            "authorized_person_phone": "0900000002",
            "authorized_person_email": "authorized@example.test",
            "business_registration_office": "Business Registration Office Ha Noi",
            "document_place": "Ha Noi",
            "document_day": "10",
            "document_month": "09",
            "document_year": "2026",
            "location_business_name": "LOCATION SAMPLE",
            "location_business_code": "00001",
        }

    def _generate(self, values: dict[str, str]):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        return DependentUnitTemplateGenerator(self.template_dir, Path(temporary.name)).generate(
            values, "Sample location", Path(temporary.name) / "output"
        )

    def test_authorization_and_signing_guide_use_the_legal_representative(self) -> None:
        result = self._generate(dict(self.values))
        self.assertEqual(3, len(result.generated))
        self.assertEqual({}, result.unresolved)
        by_source = {item.source_name: item.output_path for item in result.generated}
        authorization = Document(by_source["Giay uy quyen.docx"])
        yellow = [
            run.text for paragraph in _paragraphs(authorization) for run in paragraph.runs
            if run.font.highlight_color == WD_COLOR_INDEX.YELLOW
        ]
        self.assertEqual("NGUYEN VAN A", yellow[0])
        self.assertEqual("TRAN THI B", yellow[6])
        self.assertEqual("NGUYEN VAN A", yellow[-1])
        green = [
            run.text for paragraph in _paragraphs(authorization) for run in paragraph.runs
            if run.font.highlight_color == WD_COLOR_INDEX.BRIGHT_GREEN
        ]
        self.assertEqual(1, len(green))
        guide_text = "\n".join(
            paragraph.text for paragraph in _paragraphs(Document(by_source["Huong dan ky.docx"]))
        )
        self.assertEqual(2, guide_text.count("Nguyen Van A"))

    def test_notification_writes_selected_main_business_line_x(self) -> None:
        values = dict(self.values)
        values.update({
            "business_line_1_name": "LINE ONE",
            "business_line_1_code": "6201",
            "business_line_1_is_main": "X",
        })
        generator = DependentUnitTemplateGenerator(self.template_dir, Path(tempfile.gettempdir()))
        result = self._generate(values)
        notification = next(item.output_path for item in result.generated if item.source_name == "Thong bao thay doi.docx")
        document = Document(notification)
        industry_table = next(
            table for table in document.tables if generator._business_line_header_index(table) is not None
        )
        header_index = generator._business_line_header_index(industry_table)
        self.assertIsNotNone(header_index)
        self.assertEqual("1", industry_table.rows[header_index + 1].cells[0].text.strip())
        self.assertEqual("LINE ONE", industry_table.rows[header_index + 1].cells[1].text.strip())
        self.assertEqual("X", industry_table.rows[header_index + 1].cells[3].text.strip())

    def test_notification_address_contact_pairs_clear_dots_and_keep_their_template_lines(self) -> None:
        values = dict(self.values)
        values.update({
            "new_location_phone": "0901000001",
            "new_location_fax": "0281000001",
            "new_location_email": "location@example.test",
            "new_location_website": "example.test",
        })
        result = self._generate(values)
        notification = next(item.output_path for item in result.generated if item.source_name == "Thong bao thay doi.docx")
        paragraphs = list(_paragraphs(Document(notification)))
        phone_line = next(paragraph for paragraph in paragraphs if "0901000001" in paragraph.text)
        email_line = next(paragraph for paragraph in paragraphs if "location@example.test" in paragraph.text)
        self.assertIn("S\u1ed1 fax", phone_line.text)
        self.assertIn("0281000001", phone_line.text)
        self.assertNotIn("...", phone_line.text)
        self.assertNotIn('w:leader="dot"', phone_line._p.xml)
        self.assertIn("Website", email_line.text)
        self.assertIn("example.test", email_line.text)
        self.assertNotIn("...", email_line.text)
        self.assertNotIn('w:leader="dot"', email_line._p.xml)

    def test_notification_preserves_unselected_change_sections(self) -> None:
        values = dict(self.values)
        values.update({
            "change_name_selected": "X",
            "new_location_business_name": "LOCATION NEW NAME",
            "new_location_business_english_name": "LOCATION NEW ENGLISH",
            "new_location_business_short_name": "LNN",
            "change_address_selected": "X",
            "new_location_address_detail": "12 Sample Street",
            "new_location_ward": "Sample Ward",
            "new_location_province": "Ha Noi",
            "new_location_phone": "0901000001",
            "new_location_fax": "0281000001",
            "new_location_email": "location@example.test",
            "new_location_website": "example.test",
            "change_tax_selected": "X",
            "tax_accountant_full_name": "ACCOUNTANT SAMPLE",
            "tax_accountant_date_of_birth": "01/01/1990",
            "tax_notification_address_detail": "Tax Street",
            "tax_notification_ward": "Tax Ward",
            "tax_notification_province": "Ha Noi",
            "tax_notification_phone": "0902000001",
            "tax_notification_email": "tax@example.test",
            "tax_accounting_independent_x": "X",
            "tax_financial_year_start": "01/01",
            "tax_financial_year_end": "31/12",
            "tax_labor_count": "12",
        })
        generator = DependentUnitTemplateGenerator(self.template_dir, Path(tempfile.gettempdir()))
        result = self._generate(values)
        notification = next(item.output_path for item in result.generated if item.source_name == "Thong bao thay doi.docx")
        document = Document(notification)
        source_document = Document(self.template_dir / "Thong bao thay doi.docx")
        source_industry_table = next(
            table for table in source_document.tables if generator._business_line_header_index(table) is not None
        )
        output_industry_table = next(
            table for table in document.tables if generator._business_line_header_index(table) is not None
        )
        self.assertEqual(len(source_industry_table.rows), len(output_industry_table.rows))
        folded = generator._fold("\n".join(paragraph.text for paragraph in _paragraphs(document)))
        self.assertIn("dang ky thay doi dia chi dia diem kinh doanh", folded)
        self.assertIn("dang ky thay doi ten dia diem kinh doanh", folded)
        self.assertIn("dang ky thay doi nganh, nghe kinh doanh/noi dung hoat dong", folded)
        self.assertIn("dang ky thay doi nguoi dung dau dia diem kinh doanh", folded)
        self.assertIn("dang ky thay doi thong tin dang ky thue", folded)
        self.assertIn("bo sung, cap nhat thong tin dang ky hoat dong dia diem kinh doanh", folded)
        self.assertIn("de nghi hieu dinh thong tin dang ky hoat dong", folded)
        self.assertIn("........................................................", folded)
        self.assertIn("12 sample street", folded)
        self.assertIn("location new name", folded)
        self.assertIn("0901000001", folded)
        self.assertIn("location@example.test", folded)
        self.assertIn("accountant sample", folded)
        self.assertIn("tax@example.test", folded)
        self.assertIn("tong so lao dong: 12", folded)
        checkbox_runs = [
            run.text for paragraph in _paragraphs(document) for run in paragraph.runs
            if run.font.highlight_color == WD_COLOR_INDEX.YELLOW and run.text.strip() == "X"
        ]
        self.assertEqual([], checkbox_runs)
        self.assertEqual("Thong bao thay doi_", notification.name[:19])


    def test_notification_selected_manager_uses_template_format_without_dotted_filler(self) -> None:
        values = dict(self.values)
        values.update({
            "change_name_selected": "X",
            "new_location_business_name": "LOCATION AFTER CHANGE",
            "new_location_business_english_name": "LOCATION ENGLISH",
            "new_location_business_short_name": "LAC",
            "change_manager_selected": "X",
            "location_manager_full_name": "LE HUU DUC DUY",
            "location_manager_date_of_birth": "22/03/2000",
            "location_manager_gender": "Nam",
            "location_manager_id_number": "09876543",
            "location_manager_address_detail": "Thon 1",
            "location_manager_ward": "Xa 2",
            "location_manager_province": "Lam Dong",
            "location_manager_nationality": "Viet Nam",
            "location_manager_phone": "0900000000",
            "location_manager_email": "manager@example.test",
        })
        result = self._generate(values)
        notification = next(item.output_path for item in result.generated if item.source_name == "Thong bao thay doi.docx")
        document = Document(notification)
        paragraphs = list(_paragraphs(document))
        name_paragraph = next(paragraph for paragraph in paragraphs if "LOCATION AFTER CHANGE" in paragraph.text)
        english_name_paragraph = next(paragraph for paragraph in paragraphs if "LOCATION ENGLISH" in paragraph.text)
        short_name_paragraph = next(paragraph for paragraph in paragraphs if "LAC" in paragraph.text)
        for paragraph in (name_paragraph, english_name_paragraph, short_name_paragraph):
            self.assertNotIn("...", paragraph.text)
            self.assertNotIn("\t", paragraph.text)
        manager_date = next(paragraph for paragraph in paragraphs if "22/03/2000" in paragraph.text)
        self.assertNotIn("\t/\t/\t", manager_date.text)
        manager_name = next(paragraph for paragraph in paragraphs if "LE HUU DUC DUY" in paragraph.text)
        self.assertNotIn("...", manager_name.text)
        manager_name_run = next(run for run in manager_name.runs if "LE HUU DUC DUY" in run.text)
        self.assertEqual(12.5, manager_name_run.font.size.pt)
        manager_email = next(paragraph for paragraph in paragraphs if "manager@example.test" in paragraph.text)
        self.assertIn("Th\u01b0 \u0111i\u1ec7n t\u1eed (n\u1ebfu c\u00f3): manager@example.test", manager_email.text)
        self.assertNotIn("?", manager_email.text)
        label_run = next(run for run in manager_email.runs if "Th\u01b0 \u0111i\u1ec7n t\u1eed" in run.text)
        self.assertIsNotNone(label_run.font.size)
        self.assertEqual(12.5, label_run.font.size.pt)
        folded = DependentUnitTemplateGenerator(self.template_dir, Path(tempfile.gettempdir()))._fold(
            "\n".join(paragraph.text for paragraph in paragraphs)
        )
        self.assertIn("dang ky thay doi dia chi dia diem kinh doanh", folded)
        self.assertIn("dang ky thay doi thong tin dang ky thue", folded)

    def test_notification_manifest_matches_actual_yellow_highlight_count(self) -> None:
        manifest = json.loads((self.template_dir / "template_manifest.json").read_text(encoding="utf-8"))
        config = manifest["templates"]["Thong bao thay doi.docx"]
        document = Document(self.template_dir / "Thong bao thay doi.docx")
        highlighted = [run for paragraph in _paragraphs(document) for run in paragraph.runs if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]
        self.assertEqual(config["expected_highlight_count"], len(highlighted))
        self.assertEqual(config["expected_highlight_count"], len(config["highlighted_fields"]))

    def test_notification_preserves_empty_selected_field_placeholder(self) -> None:
        values = dict(self.values)
        values.update({"change_manager_selected": "X", "location_manager_full_name": "LE HUU DUC DUY"})
        result = self._generate(values)
        notification = next(item.output_path for item in result.generated if item.source_name == "Thong bao thay doi.docx")
        source_runs = [run.text for paragraph in _paragraphs(Document(self.template_dir / "Thong bao thay doi.docx")) for run in paragraph.runs if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]
        output_runs = [run.text for paragraph in _paragraphs(Document(notification)) for run in paragraph.runs if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]
        self.assertEqual("LE HUU DUC DUY", output_runs[35])
        self.assertEqual(source_runs[43], output_runs[43])

    def test_notification_writes_multiple_business_lines_with_one_main_x(self) -> None:
        values = dict(self.values)
        values.update({"change_business_lines_selected": "X", "business_line_1_name": "LINE ONE", "business_line_1_code": "6201", "business_line_2_name": "LINE TWO", "business_line_2_code": "6202", "business_line_2_is_main": "X"})
        generator = DependentUnitTemplateGenerator(self.template_dir, Path(tempfile.gettempdir()))
        result = self._generate(values)
        notification = next(item.output_path for item in result.generated if item.source_name == "Thong bao thay doi.docx")
        document = Document(notification)
        industry_table = next(table for table in document.tables if generator._business_line_header_index(table) is not None)
        header_index = generator._business_line_header_index(industry_table)
        self.assertIsNotNone(header_index)
        rows = industry_table.rows[header_index + 1:]
        self.assertEqual(["1", "2"], [row.cells[0].text.strip() for row in rows])
        self.assertEqual(["LINE ONE", "LINE TWO"], [row.cells[1].text.strip() for row in rows])
        self.assertEqual(["6201", "6202"], [row.cells[2].text.strip() for row in rows])
        self.assertEqual(["", "X"], [row.cells[3].text.strip() for row in rows])
        stt_paragraph = rows[0].cells[0].paragraphs[0]
        self.assertEqual(1, stt_paragraph.paragraph_format.alignment)
        self.assertEqual(12.5, stt_paragraph.runs[0].font.size.pt)

    def test_notification_fills_tax_person_fields_and_clears_their_template_dots(self) -> None:
        values = dict(self.values)
        values.update({
            "tax_accountant_full_name": "NGUYEN VAN A",
            "tax_accountant_date_of_birth": "11/11/2011",
            "tax_accountant_gender": "Nam",
            "tax_accountant_id_number": "012345678901",
            "tax_accountant_phone": "0912345678",
        })
        result = self._generate(values)
        notification = next(item.output_path for item in result.generated if item.source_name == "Thong bao thay doi.docx")
        paragraphs = list(_paragraphs(Document(notification)))
        for label, value in (
            ("Giới tính", "Nam"),
            ("Số định danh cá nhân", "012345678901"),
            ("Điện thoại", "0912345678"),
        ):
            paragraph = next(item for item in paragraphs if label in item.text and value in item.text)
            self.assertNotIn("...", paragraph.text)
            value_run = next(run for run in paragraph.runs if value in run.text)
            self.assertEqual(12.5, value_run.font.size.pt)

    def test_notification_writes_unlimited_multiline_supplemental_content_and_removes_spares(self) -> None:
        values = dict(self.values)
        values.update({
            "supplemental_registration_content": "BO SUNG NOI DUNG MOT\nBO SUNG NOI DUNG HAI\nBO SUNG NOI DUNG BA\nBO SUNG NOI DUNG BON",
        })
        result = self._generate(values)
        notification = next(item.output_path for item in result.generated if item.source_name == "Thong bao thay doi.docx")
        paragraphs = list(_paragraphs(Document(notification)))
        text = "\n".join(paragraph.text for paragraph in paragraphs)
        for line in (
            "BO SUNG NOI DUNG MOT",
            "BO SUNG NOI DUNG HAI",
            "BO SUNG NOI DUNG BA",
            "BO SUNG NOI DUNG BON",
        ):
            self.assertIn(line, text)
        supplemental_paragraph = next(paragraph for paragraph in paragraphs if "BO SUNG NOI DUNG MOT" in paragraph.text)
        self.assertIn("\n", supplemental_paragraph.text)
        self.assertNotIn("....", supplemental_paragraph.text)

    def test_notification_fills_entered_values_without_change_checkboxes(self) -> None:
        values = dict(self.values)
        values.update({
            "new_location_business_name": "LOCATION WITHOUT CHECKBOX",
            "location_manager_full_name": "MANAGER WITHOUT CHECKBOX",
            "location_manager_date_of_birth": "22/03/2000",
            "tax_notification_email": "tax-without-checkbox@example.test",
            "supplemental_registration_content_1": "SUPPLEMENT WITHOUT CHECKBOX",
            "business_line_1_name": "LINE WITHOUT CHECKBOX",
            "business_line_1_code": "0112",
            "business_line_1_is_main": "X",
        })
        result = self._generate(values)
        notification = next(item.output_path for item in result.generated if item.source_name == "Thong bao thay doi.docx")
        document = Document(notification)
        paragraphs = list(_paragraphs(document))
        text = "\n".join(paragraph.text for paragraph in paragraphs)
        for value in (
            "LOCATION WITHOUT CHECKBOX",
            "MANAGER WITHOUT CHECKBOX",
            "22/03/2000",
            "tax-without-checkbox@example.test",
            "SUPPLEMENT WITHOUT CHECKBOX",
        ):
            self.assertIn(value, text)
            run = next(run for paragraph in paragraphs for run in paragraph.runs if value in run.text)
            self.assertEqual(WD_COLOR_INDEX.YELLOW, run.font.highlight_color)
            self.assertNotIn("...", next(paragraph for paragraph in paragraphs if value in paragraph.text).text)
        generator = DependentUnitTemplateGenerator(self.template_dir, Path(tempfile.gettempdir()))
        industry_table = next(table for table in document.tables if generator._business_line_header_index(table) is not None)
        header_index = generator._business_line_header_index(industry_table)
        self.assertEqual("1", industry_table.rows[header_index + 1].cells[0].text.strip())
        self.assertEqual("LINE WITHOUT CHECKBOX", industry_table.rows[header_index + 1].cells[1].text.strip())
if __name__ == "__main__":
    unittest.main()
