from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from PySide6.QtWidgets import QApplication, QGroupBox, QPushButton, QTextEdit
from docx import Document

from mingo_assistant.models.business_registration import OcrPerson
from mingo_assistant.services.documents.location_termination_generator import LocationTerminationTemplateGenerator
from mingo_assistant.services.organization_documents.location_certificate_service import LocationCertificateService
from mingo_assistant.services.organization_documents.organization_document_service import OrganizationDocumentResult
from mingo_assistant.repositories.case_repository import CaseRepository
from mingo_assistant.ui.dialogs.dependent_unit_modal import UNIT_TYPE_SPECS
from mingo_assistant.ui.views.location_business_termination_workspace import (
    CERTIFICATE_FIELDS,
    LOCATION_TERMINATION_CASE_TYPE,
    TERMINATION_FIELD_GROUPS,
    LocationBusinessTerminationWorkspace,
)


RAPID_OCR_TEXT = """GIAY CHU'NG NHAN DANG KY DIA DIEM KINH DOANH
Ma s6 dia diém kinh doanh: 00001
1. Ten dia diém kinh doanh: DIA DIEM KINH DOANH 1 - CONG TY TNHH VAT
LIEUMOILUYIVIET NAM
Ten dia diem kinh doanh viét bang tieng nuoc ngoai:
2. Dia chi:
S0 35, Duong D18, Khu Hoa Phu 1, Phuong Binh Duong, Thanh pho Ho Chi Minh,
Viet Nam
Dienthoai:13326272255
S6 Fax:
Thu dien t’r: 136033164@qq.com
Website:
3. Thong tin ve nguoi dung dau
4. Thong tin ve doanh nghiep/chi nhanh chu quan:
Ten doanh nghiep/chi nhanh: CONG TY TNHH VAT LIEU MOI LUYI VIET NAM
Ma s6 doanh nghiep/chi nhanh: 3703192163
Dia chi tru so chinh cua doanh nghiep/chi nhanh: Nha xuong so 15
"""


class LocationCertificateTests(unittest.TestCase):
    def test_parses_certificate_fields_and_keeps_missing_contacts_blank(self) -> None:
        result = LocationCertificateService()._parse_certificate(RAPID_OCR_TEXT)
        self.assertEqual("3703192163", result["parent_entity_registration_number"])
        self.assertEqual("CÔNG TY TNHH VẬT LIỆU MỚI LUYI VIỆT NAM", result["parent_entity_name"])
        self.assertEqual("00001", result["location_business_code"])
        self.assertEqual(
            "ĐỊA ĐIỂM KINH DOANH 1 - CÔNG TY TNHH VẬT LIỆU MỚI LUYI VIỆT NAM",
            result["location_business_name"],
        )
        self.assertEqual("13326272255", result["location_phone"])
        self.assertEqual("136033164@qq.com", result["location_email"])
        self.assertEqual("", result["location_fax"])
        self.assertEqual("", result["location_website"])


class TerminationWorkspaceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def test_workspace_has_requested_sections_and_scanned_fields(self) -> None:
        workspace = LocationBusinessTerminationWorkspace(on_back=lambda: None, on_save=lambda _: None)
        try:
            self.assertEqual([key for key, _ in CERTIFICATE_FIELDS], list(workspace.organization_card_fields))
            certificate_group = next(
                group for group in workspace.findChildren(QGroupBox)
                if group.title() == "Thông tin giấy chứng nhận địa điểm kinh doanh nhận diện"
            )
            certificate_grid = certificate_group.layout()

            def field_position(key: str) -> tuple[int, int, int, int]:
                field = workspace.organization_card_fields[key]
                index = next(
                    index for index in range(certificate_grid.count())
                    if certificate_grid.itemAt(index).widget() is field
                )
                return certificate_grid.getItemPosition(index)

            self.assertEqual((1, 1, 1, 3), field_position("parent_entity_name"))
            self.assertEqual((2, 1, 1, 3), field_position("parent_entity_registration_number"))
            self.assertEqual([title for title, _ in TERMINATION_FIELD_GROUPS], list(workspace.form_sections))
            self.assertEqual([
                "Thông tin cơ quan tiếp nhận và thời gian ký",
                "Thông tin người đại diện pháp luật",
                "Thông tin doanh nghiệp",
                "Lý do chấm dứt hoạt động",
                "Thông tin người được ủy quyền",
            ], list(workspace.form_sections))
            self.assertIsInstance(workspace.fields["termination_reason"], QTextEdit)
            self.assertEqual(8, len(workspace.ocr_card_fields))
            self.assertEqual(LOCATION_TERMINATION_CASE_TYPE, workspace.build_case_summary().case_type)
            self.assertTrue(any(isinstance(section, QGroupBox) for section in workspace.form_sections.values()))
            authorized_section = workspace.form_sections["Thông tin người được ủy quyền"]
            export_button = next(
                button for button in authorized_section.findChildren(QPushButton)
                if button.text() == "Xuất hồ sơ"
            )
            self.assertEqual("PrimaryButton", export_button.objectName())
            option = next(p for unit in UNIT_TYPE_SPECS if unit["id"] == "ddkd"
                          for p in unit["procedures"] if p.id == "ddkd_terminate")
            self.assertTrue(option.is_ready)
        finally:
            workspace.deleteLater()

    def test_open_uploaded_pdf_uses_windows_shell(self) -> None:
        with TemporaryDirectory() as directory:
            uploaded_pdf = Path(directory) / "certificate.pdf"
            uploaded_pdf.touch()
            workspace = LocationBusinessTerminationWorkspace(on_back=lambda: None, on_save=lambda _: None)
            try:
                workspace.ocr_people[0].front_path = str(uploaded_pdf)
                with patch("mingo_assistant.ui.views.business_registration_workspace.os.startfile", create=True) as startfile:
                    workspace._open_or_choose_person_image("front")
                startfile.assert_called_once_with(str(uploaded_pdf))
            finally:
                workspace.deleteLater()

    def test_open_uploaded_pdf_reports_missing_file_instead_of_silently_reopening_picker(self) -> None:
        with TemporaryDirectory() as directory:
            missing_pdf = Path(directory) / "missing.pdf"
            workspace = LocationBusinessTerminationWorkspace(on_back=lambda: None, on_save=lambda _: None)
            try:
                workspace.ocr_people[0].front_path = str(missing_pdf)
                with patch("mingo_assistant.ui.views.business_registration_workspace.QMessageBox.warning") as warning, \
                     patch.object(workspace, "_choose_person_image") as choose:
                    workspace._open_or_choose_person_image("front")
                warning.assert_called_once()
                choose.assert_not_called()
            finally:
                workspace.deleteLater()

    def test_pdf_scan_populates_form_without_assigning_legal_representative(self) -> None:
        workspace = LocationBusinessTerminationWorkspace(on_back=lambda: None, on_save=lambda _: None)
        try:
            values = LocationCertificateService()._parse_certificate(RAPID_OCR_TEXT)
            result = OrganizationDocumentResult(values, RAPID_OCR_TEXT, "ocr")
            workspace.organization_document_service.extract_pdf = lambda *_args, **_kwargs: result
            with patch("mingo_assistant.ui.views.business_registration_workspace.QMessageBox.warning"):
                workspace._scan_organization_document("certificate.pdf")
            self.assertEqual("00001", workspace._field_value("location_business_code"))
            self.assertEqual("3703192163", workspace._field_value("parent_entity_registration_number"))
            self.assertEqual("", workspace._field_value("legal_representative_full_name"))
            self.assertEqual("", workspace._field_value("authorized_person_full_name"))
            summary = workspace.build_case_summary()
            self.assertEqual("00001", summary.form_values["location_business_code"])
            self.assertEqual("organization", workspace.ocr_people[0].recognition_type)
            self.assertEqual("", workspace._field_value("business_registration_office"))
            self.assertEqual("27", workspace._field_value("document_day"))
            self.assertEqual("09", workspace._field_value("document_month"))
            self.assertEqual("2026", workspace._field_value("document_year"))
            workspace._set_field("business_registration_office", "Phòng Đăng ký kinh doanh Thành phố Hồ Chí Minh")
            self.assertEqual("Thành phố Hồ Chí Minh", workspace._field_value("document_place"))
            self.assertEqual(
                "Thành phố Hồ Chí Minh, ngày 27 tháng 09 năm 2026",
                workspace._field_value("document_place_date"),
            )
            reopened = LocationBusinessTerminationWorkspace(
                on_back=lambda: None, on_save=lambda _: None, case=summary
            )
            try:
                self.assertEqual("00001", reopened._field_value("location_business_code"))
                self.assertEqual("136033164@qq.com", reopened.organization_card_fields["location_email"].text())
            finally:
                reopened.deleteLater()
        finally:
            workspace.deleteLater()

    def test_save_certificate_button_persists_draft_and_confirms_result(self) -> None:
        with TemporaryDirectory() as directory:
            repository = CaseRepository(Path(directory) / "cases.json")
            saved_cases = []
            def save(case):
                repository.save_case(case)
                saved_cases.append(case)

            workspace = LocationBusinessTerminationWorkspace(
                on_back=lambda: None, on_save=save
            )
            try:
                button = next(
                    button for button in workspace.findChildren(QPushButton)
                    if button.text() == "Lưu dữ liệu giấy chứng nhận"
                )
                with patch("mingo_assistant.ui.views.location_business_termination_workspace.QMessageBox.warning") as warning:
                    button.click()
                warning.assert_called_once()
                self.assertFalse(saved_cases)

                values = LocationCertificateService()._parse_certificate(RAPID_OCR_TEXT)
                workspace._update_organization_result_fields(values)
                workspace._set_field("termination_reason", "Kinh doanh không hiệu quả")
                with patch("mingo_assistant.ui.views.location_business_termination_workspace.QMessageBox.information") as confirmation:
                    button.click()
                confirmation.assert_called_once()
                self.assertEqual(1, len(saved_cases))
                self.assertIs(workspace.case, saved_cases[0])
                self.assertEqual("00001", saved_cases[0].form_values["location_business_code"])
                self.assertEqual("Kinh doanh không hiệu quả", saved_cases[0].form_values["termination_reason"])
                self.assertEqual("Đã lưu dữ liệu giấy chứng nhận", workspace.ocr_status_label.text())

                stored_case = repository.get_case(saved_cases[0].id)
                self.assertIsNotNone(stored_case)
                reopened = LocationBusinessTerminationWorkspace(
                    on_back=lambda: None, on_save=lambda _: None, case=stored_case
                )
                try:
                    self.assertEqual("00001", reopened._field_value("location_business_code"))
                    self.assertEqual("136033164@qq.com", reopened.organization_card_fields["location_email"].text())
                    self.assertEqual("Kinh doanh không hiệu quả", reopened._field_value("termination_reason"))
                finally:
                    reopened.deleteLater()
            finally:
                workspace.deleteLater()

    def test_selected_scanned_representative_supplies_word_signer(self) -> None:
        workspace = LocationBusinessTerminationWorkspace(on_back=lambda: None, on_save=lambda _: None)
        try:
            workspace._set_fields_from_mapping(LocationCertificateService()._parse_certificate(RAPID_OCR_TEXT))
            workspace._set_fields_from_mapping({"document_day": "27", "document_month": "09", "document_year": "2026"})
            workspace.ocr_people.append(OcrPerson("NGUYỄN VĂN A", data={
                "legal_representative_full_name": "NGUYỄN VĂN A",
                "legal_representative_id_number": "012345678901",
                "legal_representative_permanent_address": "12 Đường Mẫu, Phường Bến Thành, Thành phố Hồ Chí Minh",
            }))
            workspace._refresh_person_fill_selector("legal_representative")
            selector = workspace.person_fill_selectors["legal_representative"]
            selector.setCurrentIndex(selector.findData("person:1"))
            values = workspace.form_values()
            workspace._set_fields_from_mapping({
                "legal_representative_position": "Giám đốc",
                "termination_reason": "Kinh doanh không hiệu quả",
                "authorized_person_full_name": "TRẦN THỊ B",
                "authorized_person_id_number": "098765432109",
                "authorized_person_phone": "0901234567",
                "authorized_person_email": "tranthib@example.com",
            })
            values = workspace.form_values()
            self.assertEqual("NGUYỄN VĂN A", values["legal_representative_full_name"])
            self.assertEqual("NGUYỄN VĂN A", values["signer_full_name"])
            self.assertEqual([], workspace._validation_issues(values))
        finally:
            workspace.deleteLater()

    def test_export_allows_a_receiving_office_outside_the_location_province(self) -> None:
        workspace = LocationBusinessTerminationWorkspace(on_back=lambda: None, on_save=lambda _: None)
        try:
            issues = workspace._validation_issues({"location_province": "Thành phố Hải Phòng"})
            self.assertNotIn(
                "Mẫu Word số 1 đang ghi sẵn cơ quan tiếp nhận và nơi ký tại Thành phố Hồ Chí Minh",
                issues,
            )
        finally:
            workspace.deleteLater()


class TerminationNoticeTests(unittest.TestCase):
    def test_fills_only_highlighted_runs_and_preserves_their_formatting(self) -> None:
        template_dir = Path(__file__).resolve().parents[2] / "assets/templates/word/dependent_units/ddkd_terminate"
        template = Document(template_dir / LocationTerminationTemplateGenerator.NOTICE_NAME)
        values = {
            "parent_entity_name": "CÔNG TY TNHH VẬT LIỆU MỚI LUYI VIỆT NAM",
            "parent_entity_registration_number": "3703192163",
            "location_business_name": "ĐỊA ĐIỂM KINH DOANH 1 - CÔNG TY TNHH VẬT LIỆU MỚI LUYI VIỆT NAM",
            "location_business_code": "00001",
            "location_address_detail": "Số 35, Đường D18, Khu Hòa Phú 1",
            "location_ward": "Phường Bình Dương",
            "location_province": "Thành phố Hồ Chí Minh",
            "termination_reason": "Kinh doanh không hiệu quả",
            "document_day": "27", "document_month": "09", "document_year": "2026",
            "legal_representative_full_name": "NGUYỄN VĂN A",
        }
        with TemporaryDirectory() as directory:
            generator = LocationTerminationTemplateGenerator(template_dir, Path(directory))
            result = generator.generate(values, "Hồ sơ chấm dứt", Path(directory))
            self.assertEqual(3, len(result.generated))
            notice_path = next(item.output_path for item in result.generated if item.source_name == LocationTerminationTemplateGenerator.NOTICE_NAME)
            output = Document(notice_path)
            self.assertEqual(values["parent_entity_name"], output.tables[0].cell(0, 0).paragraphs[0].text)
            self.assertIn("3703192163", output.paragraphs[8].text)
            self.assertIn("00001", output.paragraphs[12].text)
            self.assertIn("Số 35, Đường D18", output.paragraphs[14].text)
            self.assertIn(values["termination_reason"], output.paragraphs[23].text)
            self.assertIn("NGUYỄN VĂN A", output.tables[1].cell(0, 1).text)
            self.assertEqual(template.paragraphs[17].text, output.paragraphs[17].text)
            self.assertEqual(template.paragraphs[18].text, output.paragraphs[18].text)
            self.assertEqual(template.paragraphs[24].text, output.paragraphs[24].text)
            original = template.paragraphs[8].runs[1]
            filled = output.paragraphs[8].runs[1]
            self.assertEqual(original.font.highlight_color, filled.font.highlight_color)
            self.assertEqual(original.font.size, filled.font.size)
            self.assertEqual(original.font.name, filled.font.name)
            self.assertEqual(original.bold, filled.bold)
            original_paragraphs = list(generator._iter_paragraphs(template))
            filled_paragraphs = list(generator._iter_paragraphs(output))
            self.assertEqual(len(original_paragraphs), len(filled_paragraphs))
            for before, after in zip(original_paragraphs, filled_paragraphs):
                self.assertEqual(len(before.runs), len(after.runs))
                for original_run, filled_run in zip(before.runs, after.runs):
                    self.assertEqual(original_run.font.highlight_color, filled_run.font.highlight_color)
                    self.assertEqual(original_run.font.size, filled_run.font.size)
                    self.assertEqual(original_run.font.name, filled_run.font.name)
                    self.assertEqual(original_run.bold, filled_run.bold)
                    self.assertEqual(original_run.italic, filled_run.italic)
                    if original_run.font.highlight_color is None:
                        self.assertEqual(original_run.text, filled_run.text)

    def test_fills_authorization_highlights_and_preserves_non_highlighted_address(self) -> None:
        template_dir = Path(__file__).resolve().parents[2] / "assets/templates/word/dependent_units/ddkd_terminate"
        template = Document(template_dir / LocationTerminationTemplateGenerator.AUTHORIZATION_NAME)
        values = {
            "parent_entity_name": "CÔNG TY TNHH VẬT LIỆU MỚI LUYI VIỆT NAM",
            "parent_entity_registration_number": "3703192163",
            "location_business_name": "ĐỊA ĐIỂM KINH DOANH 1 - CÔNG TY TNHH VẬT LIỆU MỚI LUYI VIỆT NAM",
            "business_registration_office": "Phòng Đăng ký kinh doanh Thành phố Hồ Chí Minh",
            "document_place": "Thành phố Hồ Chí Minh",
            "document_day": "27", "document_month": "09", "document_year": "2026",
            "legal_representative_full_name": "NGUYỄN VĂN A",
            "legal_representative_id_number": "012345678901",
            "legal_representative_permanent_address": "12 Đường Mẫu, Phường Bến Thành, Thành phố Hồ Chí Minh",
            "legal_representative_position": "Giám đốc",
            "authorized_person_full_name": "TRẦN THỊ B",
            "authorized_person_id_number": "098765432109",
            "authorized_person_phone": "0901234567",
            "authorized_person_email": "tranthib@example.com",
        }
        with TemporaryDirectory() as directory:
            generator = LocationTerminationTemplateGenerator(template_dir, Path(directory))
            result = generator.generate(values, "Hồ sơ chấm dứt", Path(directory))
            authorization_path = next(
                item.output_path for item in result.generated
                if item.source_name == LocationTerminationTemplateGenerator.AUTHORIZATION_NAME
            )
            output = Document(authorization_path)
            self.assertIn(values["legal_representative_full_name"], output.paragraphs[5].text)
            self.assertIn(values["parent_entity_registration_number"], output.paragraphs[8].text)
            self.assertIn(values["authorized_person_full_name"], output.paragraphs[10].text)
            self.assertIn(values["authorized_person_email"], output.paragraphs[13].text)
            self.assertIn("Phòng Đăng ký kinh doanh\u00A0Thành phố Hồ Chí Minh", output.paragraphs[15].text)
            self.assertIn(values["location_business_name"], output.paragraphs[16].text)
            self.assertIn(values["document_place"], output.paragraphs[33].text)
            self.assertIn(values["legal_representative_full_name"], output.tables[0].cell(0, 1).text)
            self.assertEqual(template.paragraphs[12].text, output.paragraphs[12].text)
            original_paragraphs = list(generator._iter_paragraphs(template))
            filled_paragraphs = list(generator._iter_paragraphs(output))
            self.assertEqual(len(original_paragraphs), len(filled_paragraphs))
            self.assertEqual(20, sum(
                run.font.highlight_color is not None
                for paragraph in filled_paragraphs for run in paragraph.runs
            ))
            for before, after in zip(original_paragraphs, filled_paragraphs):
                self.assertEqual(len(before.runs), len(after.runs))
                for original_run, filled_run in zip(before.runs, after.runs):
                    self.assertEqual(original_run.font.highlight_color, filled_run.font.highlight_color)
                    self.assertEqual(original_run.font.size, filled_run.font.size)
                    self.assertEqual(original_run.font.name, filled_run.font.name)
                    self.assertEqual(original_run.bold, filled_run.bold)
                    self.assertEqual(original_run.italic, filled_run.italic)
                    if original_run.font.highlight_color is None:
                        self.assertEqual(original_run.text, filled_run.text)

    def test_fills_signing_guide_with_the_selected_legal_representative(self) -> None:
        template_dir = Path(__file__).resolve().parents[2] / "assets/templates/word/dependent_units/ddkd_terminate"
        template = Document(template_dir / LocationTerminationTemplateGenerator.SIGNING_GUIDE_NAME)
        values = {
            "legal_representative_full_name": "NGUYỄN VĂN A",
            "parent_entity_name": "CÔNG TY TNHH VẬT LIỆU MỚI LUYI VIỆT NAM",
        }
        with TemporaryDirectory() as directory:
            generator = LocationTerminationTemplateGenerator(template_dir, Path(directory))
            result = generator.generate(values, "Hồ sơ chấm dứt", Path(directory))
            signing_guide_path = next(
                item.output_path for item in result.generated
                if item.source_name == LocationTerminationTemplateGenerator.SIGNING_GUIDE_NAME
            )
            output = Document(signing_guide_path)
            self.assertEqual(3, len(result.generated))
            for row_index in (1, 2):
                self.assertIn("Nguyễn Văn A", output.tables[0].cell(row_index, 3).text)
            self.assertEqual(template.tables[0].cell(1, 1).text, output.tables[0].cell(1, 1).text)
            self.assertEqual(template.tables[0].cell(2, 4).text, output.tables[0].cell(2, 4).text)
            original_paragraphs = list(generator._iter_paragraphs(template))
            filled_paragraphs = list(generator._iter_paragraphs(output))
            self.assertEqual(len(original_paragraphs), len(filled_paragraphs))
            self.assertEqual(2, sum(
                run.font.highlight_color is not None
                for paragraph in filled_paragraphs for run in paragraph.runs
            ))
            for before, after in zip(original_paragraphs, filled_paragraphs):
                self.assertEqual(len(before.runs), len(after.runs))
                for original_run, filled_run in zip(before.runs, after.runs):
                    self.assertEqual(original_run.font.highlight_color, filled_run.font.highlight_color)
                    self.assertEqual(original_run.font.size, filled_run.font.size)
                    self.assertEqual(original_run.font.name, filled_run.font.name)
                    self.assertEqual(original_run.bold, filled_run.bold)
                    self.assertEqual(original_run.italic, filled_run.italic)
                    if original_run.font.highlight_color is None:
                        self.assertEqual(original_run.text, filled_run.text)


if __name__ == "__main__":
    unittest.main()
